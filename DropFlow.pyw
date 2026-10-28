#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
DropFlow - a tiny, fully local drag-and-drop file toolbox for Windows.

Press the global shortcut (default Ctrl+Shift+D) and a floating drop zone appears
under your mouse.  Drop files or folders on it and pick what to do with them.
Everything runs on your machine.  No cloud, no telemetry, no console window.

Quick install (Python 3.10+ / Windows 10 or 11):

    pip install pystray tkinterdnd2 Pillow PyMuPDF python-docx reportlab ^
                python-pptx openpyxl pytesseract arabic-reshaper python-bidi

Optional RAR support (either one is enough):

    pip install rarfile
    or install 7-Zip / WinRAR (DropFlow uses 7z.exe or UnRAR.exe)

Optional helper programs (DropFlow detects all of them automatically):

  * LibreOffice  - best-quality Office (Word/PowerPoint/Excel) -> PDF conversion.
                   https://www.libreoffice.org/download/
  * FFmpeg       - video compression and audio extraction.
                   winget install Gyan.FFmpeg   (or https://ffmpeg.org/)
  * Tesseract    - OCR.  Use the UB Mannheim Windows build and tick the
                   languages you need during setup, then choose them in
                   Settings (for example "ara+eng").
                   https://github.com/UB-Mannheim/tesseract/wiki

Run it by double-clicking DropFlow.pyw.  Settings, config and log live in
%APPDATA%\\DropFlow\\ (config.json, dropflow.log).

Universal media conversion is provided by media_converter.py; keep that file
next to DropFlow.pyw.  Office/document conversion is provided by
office_converter.py; keep that file next to DropFlow.pyw too.

Included tools: PDF merge/compress/split/page ranges/export/OCR, searchable
PDFs, image and GIF conversion, Word/PowerPoint/Excel/text conversion, ZIP/RAR
extraction with password prompts, video compression with
resolution/target-size/GPU options, and audio extraction.

Debugging: run "python DropFlow.pyw --console" to keep a console attached.
"""
from __future__ import annotations

import ctypes
from ctypes import wintypes
import importlib
import importlib.util
import io
import json
import logging
import os
import queue
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from logging.handlers import RotatingFileHandler
from pathlib import Path
from typing import Callable
import tkinter as tk
from tkinter import ttk

# --------------------------------------------------------------------------- #
#  Bootstrap: platform guard, win32 basics, console-free start
# --------------------------------------------------------------------------- #
if sys.platform != "win32":
    raise SystemExit("DropFlow is a Windows-only application.")

if sys.version_info < (3, 10):
    raise SystemExit("DropFlow needs Python 3.10 or newer.")

# pythonw.exe has no stdout/stderr; some libraries print and would crash.
if sys.stdout is None:
    sys.stdout = open(os.devnull, "w", encoding="utf-8")
if sys.stderr is None:
    sys.stderr = open(os.devnull, "w", encoding="utf-8")

CREATE_NO_WINDOW = getattr(subprocess, "CREATE_NO_WINDOW", 0x08000000)
DETACHED_PROCESS = 0x00000008

_user32 = ctypes.WinDLL("user32", use_last_error=True)
_kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
try:
    _shcore = ctypes.WinDLL("shcore", use_last_error=True)
except Exception:
    _shcore = None


def _set_dpi_awareness() -> None:
    try:  # per-monitor v2 on Windows 10+
        _user32.SetProcessDpiAwarenessContext.argtypes = [ctypes.c_void_p]
        _user32.SetProcessDpiAwarenessContext.restype = wintypes.BOOL
        if _user32.SetProcessDpiAwarenessContext(ctypes.c_void_p(-4)):
            return
    except Exception:
        pass
    try:
        ctypes.windll.shcore.SetProcessDpiAwareness(1)
    except Exception:
        try:
            _user32.SetProcessDPIAware()
        except Exception:
            pass


_set_dpi_awareness()


def _msgbox(text: str, title: str = "DropFlow", flags: int = 0x10) -> None:
    try:
        _user32.MessageBoxW.argtypes = [wintypes.HWND, wintypes.LPCWSTR,
                                        wintypes.LPCWSTR, wintypes.UINT]
        _user32.MessageBoxW.restype = ctypes.c_int
        _user32.MessageBoxW(None, text, title, flags)
    except Exception:
        pass


def _check_hard_dependencies() -> None:
    # In a PyInstaller build these are already frozen into the exe, and
    # importlib.util.find_spec() can behave differently inside the archive.
    if getattr(sys, "frozen", False):
        return
    needed = {
        "pystray": "pystray",
        "tkinterdnd2": "tkinterdnd2",
        "PIL": "Pillow",
    }
    missing = [pip for mod, pip in needed.items() if importlib.util.find_spec(mod) is None]
    if importlib.util.find_spec("pymupdf") is None and importlib.util.find_spec("fitz") is None:
        missing.append("PyMuPDF")
    if missing:
        _msgbox(
            "DropFlow needs these Python packages:\n\n    " + ", ".join(missing)
            + "\n\nInstall everything with:\n\n"
              "pip install pystray tkinterdnd2 Pillow PyMuPDF python-docx reportlab "
              "pytesseract arabic-reshaper python-bidi",
            flags=0x10)
        sys.exit(1)


_check_hard_dependencies()

_MUTEX = None


def _ensure_single_instance() -> None:
    global _MUTEX
    try:
        _kernel32.CreateMutexW.argtypes = [ctypes.c_void_p, wintypes.BOOL, wintypes.LPCWSTR]
        _kernel32.CreateMutexW.restype = wintypes.HANDLE
        _MUTEX = _kernel32.CreateMutexW(None, False, "Local\\DropFlow_SingleInstance")
        if ctypes.get_last_error() == 183:  # ERROR_ALREADY_EXISTS
            _msgbox("DropFlow is already running.\n\nLook for its icon in the system tray "
                    "or press your shortcut.", flags=0x40)
            sys.exit(0)
    except Exception:
        pass


def _relaunch_without_console() -> None:
    """If started with python.exe, hand off to pythonw.exe and leave no console."""
    if getattr(sys, "frozen", False):
        return
    if "--console" in sys.argv:
        return
    if os.environ.get("DROPFLOW_RELAUNCHED") == "1":
        return
    exe = Path(sys.executable)
    if exe.name.lower() != "python.exe":
        return
    pythonw = exe.with_name("pythonw.exe")
    child_exe = pythonw if pythonw.is_file() else exe
    script = Path(__file__).resolve()
    env = os.environ.copy()
    env["DROPFLOW_RELAUNCHED"] = "1"
    try:
        subprocess.Popen(
            [str(child_exe), str(script), *sys.argv[1:]],
            cwd=str(script.parent), env=env,
            stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
            creationflags=DETACHED_PROCESS | CREATE_NO_WINDOW, close_fds=True)
    except Exception:
        # Last resort: detach this process from its console and keep running.
        try:
            sys.stdout = open(os.devnull, "w", encoding="utf-8")
            sys.stderr = open(os.devnull, "w", encoding="utf-8")
            _kernel32.FreeConsole()
        except Exception:
            pass
        return
    raise SystemExit(0)


import pystray  # noqa: E402
from PIL import Image, ImageDraw, ImageOps  # noqa: E402
from tkinterdnd2 import DND_FILES, TkinterDnD  # noqa: E402

try:
    import pymupdf as fitz  # noqa: E402
except ImportError:  # older PyMuPDF
    import fitz  # type: ignore  # noqa: E402

try:
    import media_converter as media  # noqa: E402
    _MEDIA_IMPORT_ERROR = ""
except Exception as _media_exc:  # optional module shipped next to DropFlow.pyw
    media = None  # type: ignore[assignment]
    _MEDIA_IMPORT_ERROR = str(_media_exc)

try:
    import office_converter as office  # noqa: E402
    _OFFICE_IMPORT_ERROR = ""
except Exception as _office_exc:
    office = None  # type: ignore[assignment]
    _OFFICE_IMPORT_ERROR = str(_office_exc)


# --------------------------------------------------------------------------- #
#  Constants, theme, config, logging
# --------------------------------------------------------------------------- #
APP_NAME = "DropFlow"
SUFFIX = "_dropflow"
DEFAULT_HOTKEY = "ctrl+shift+d"
IMG_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff", ".gif"}
VIDEO_EXTS = {".mp4", ".mkv", ".mov", ".avi", ".wmv", ".webm", ".m4v",
              ".flv", ".mpg", ".mpeg", ".ts", ".3gp", ".ogv"}
OFFICE_INPUT_EXTS = {".docx", ".pptx", ".ppt", ".xlsx", ".xls",
                     ".txt", ".csv", ".json"}
ARCHIVE_EXTS = {".zip", ".rar"}
ARABIC_RE = re.compile(r"[\u0590-\u08FF\uFB1D-\uFDFF\uFE70-\uFEFF]")
IMAGE_PDF_DPI = 150.0

BG = "#181825"
FG = "#cdd6f4"
ACCENT = "#cba6f7"
ACCENT_HI = "#e2cffb"
SURFACE = "#1f1f31"
HOVER = "#313244"
SUBTLE = "#7f849c"
GREEN = "#a6e3a1"
RED = "#f38ba8"
FONT = "Segoe UI"

SCALE = 1.0


def px(n: float) -> int:
    """Scale a pixel size by the screen DPI factor."""
    return int(round(n * SCALE))


DATA_DIR = Path(os.environ.get("APPDATA", str(Path.home()))) / APP_NAME
CONFIG_PATH = DATA_DIR / "config.json"
RUN_KEY = r"Software\Microsoft\Windows\CurrentVersion\Run"

log = logging.getLogger(APP_NAME)


def _setup_logging() -> None:
    if log.handlers:
        return
    log.setLevel(logging.INFO)
    log.propagate = False
    try:
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        handler = RotatingFileHandler(DATA_DIR / "dropflow.log", maxBytes=512_000,
                                      backupCount=1, encoding="utf-8")
        handler.setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(message)s"))
        log.addHandler(handler)
    except Exception:
        pass
    if "--console" in sys.argv and sys.stderr is not None:
        sh = logging.StreamHandler(sys.stderr)
        sh.setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(message)s"))
        log.addHandler(sh)

    def _excepthook(exc_type, exc_value, exc_tb):
        log.critical("Uncaught exception", exc_info=(exc_type, exc_value, exc_tb))

    def _thread_excepthook(args):
        log.critical("Thread exception", exc_info=(args.exc_type, args.exc_value,
                                                   args.exc_traceback))

    sys.excepthook = _excepthook
    threading.excepthook = _thread_excepthook


def load_config() -> dict:
    cfg = {"hotkey": DEFAULT_HOTKEY, "autostart": False, "ocr_langs": "ara+eng"}
    try:
        data = json.loads(CONFIG_PATH.read_text(encoding="utf-8"))
        if isinstance(data, dict):
            hk = data.get("hotkey")
            if isinstance(hk, str) and hk.strip():
                cfg["hotkey"] = hk.strip().lower()
            cfg["autostart"] = bool(data.get("autostart", False))
            langs = data.get("ocr_langs")
            if isinstance(langs, str) and langs.strip():
                cfg["ocr_langs"] = langs.strip()
    except Exception:
        pass
    return cfg


def save_config(cfg: dict) -> None:
    try:
        CONFIG_PATH.parent.mkdir(parents=True, exist_ok=True)
        tmp = CONFIG_PATH.with_suffix(".json.tmp")
        tmp.write_text(json.dumps(cfg, indent=2), encoding="utf-8")
        os.replace(tmp, CONFIG_PATH)
    except Exception:
        log.exception("Could not save config")


def pretty_hotkey(hk: str) -> str:
    names = {
        "ctrl": "Ctrl", "alt": "Alt", "shift": "Shift", "win": "Win",
        "space": "Space", "enter": "Enter", "esc": "Esc", "tab": "Tab",
        "backspace": "Backspace", "delete": "Delete", "insert": "Insert",
        "home": "Home", "end": "End", "pageup": "Page Up", "pagedown": "Page Down",
        "left": "Left", "up": "Up", "right": "Right", "down": "Down",
        "printscreen": "Print Screen", "pause": "Pause", "capslock": "Caps Lock",
        "numlock": "Num Lock", "scrolllock": "Scroll Lock",
        "minus": "Minus", "plus": "Plus", "equal": "Equal", "comma": "Comma",
        "period": "Period", "slash": "Slash", "backslash": "Backslash",
        "semicolon": "Semicolon", "quote": "Quote", "bracketleft": "[",
        "bracketright": "]", "grave": "`", "multiply": "Num *", "add": "Num +",
        "subtract": "Num -", "decimal": "Num .", "divide": "Num /",
    }
    parts = []
    for raw in hk.split("+"):
        part = raw.strip().lower()
        if not part:
            continue
        if re.fullmatch(r"f(?:[1-9]|1[0-9]|2[0-4])", part):
            parts.append(part.upper())
        elif part in names:
            parts.append(names[part])
        elif re.fullmatch(r"num[0-9]", part):
            parts.append("Num " + part[-1])
        elif len(part) == 1:
            parts.append(part.upper())
        else:
            parts.append(part.capitalize())
    return " + ".join(parts)


def launch_command() -> str:
    if getattr(sys, "frozen", False):
        return f'"{sys.executable}"'
    exe = Path(sys.executable)
    if exe.name.lower() == "python.exe":
        pythonw = exe.with_name("pythonw.exe")
        if pythonw.is_file():
            exe = pythonw
    return f'"{exe}" "{Path(__file__).resolve()}"'


def app_dir() -> Path:
    """Folder the app was launched from (works for source and frozen exe)."""
    if getattr(sys, "frozen", False):
        try:
            return Path(sys.executable).resolve().parent
        except Exception:
            pass
    return Path(__file__).resolve().parent


def set_autostart(enabled: bool) -> None:
    import winreg
    with winreg.CreateKeyEx(winreg.HKEY_CURRENT_USER, RUN_KEY, 0, winreg.KEY_SET_VALUE) as key:
        if enabled:
            winreg.SetValueEx(key, APP_NAME, 0, winreg.REG_SZ, launch_command())
        else:
            try:
                winreg.DeleteValue(key, APP_NAME)
            except FileNotFoundError:
                pass


# --------------------------------------------------------------------------- #
#  Win32 window helpers
# --------------------------------------------------------------------------- #
def force_foreground(target: tk.Misc | int, alt_trick: bool = True) -> None:
    """Bring a Tk window to the front and give it keyboard focus.

    Windows blocks SetForegroundWindow from background processes.  A quick Alt
    tap is the well-known way to release that lock, which makes Esc and other
    keys work immediately when the drop zone appears.

    Do not use the Alt tap on normal title-bar windows (Settings): releasing
    Alt can enter the window's system-menu modal loop and freeze the UI.
    """
    hwnd = None
    is_widget = isinstance(target, tk.Misc)
    try:
        if is_widget:
            target.update_idletasks()
            hwnd = _user32.GetAncestor(target.winfo_id(), 2) or target.winfo_id()
        else:
            hwnd = int(target)
        if not hwnd:
            return
        if alt_trick:
            _user32.keybd_event(0x12, 0, 0, 0)  # Alt down
        try:
            # Deliberately avoid AttachThreadInput: if the current foreground
            # window's thread is busy, attaching to it can deadlock the UI.
            _user32.BringWindowToTop(hwnd)
            _user32.SetForegroundWindow(hwnd)
            _user32.SetFocus(hwnd)
        finally:
            if alt_trick:
                _user32.keybd_event(0x12, 0, 2, 0)  # Alt up
        if is_widget:
            target.focus_force()
    except Exception:
        log.debug("Could not force foreground", exc_info=True)
        if is_widget:
            try:
                target.focus_force()
            except Exception:
                pass


def dark_titlebar(win: tk.Misc) -> None:
    try:
        win.update_idletasks()
        hwnd = _user32.GetAncestor(win.winfo_id(), 2) or win.winfo_id()
        val = ctypes.c_int(1)
        dwm = ctypes.windll.dwmapi
        dwm.DwmSetWindowAttribute.argtypes = [wintypes.HWND, wintypes.DWORD,
                                              ctypes.c_void_p, wintypes.DWORD]
        for attr in (20, 19):
            if dwm.DwmSetWindowAttribute(hwnd, attr, ctypes.byref(val),
                                         ctypes.sizeof(val)) == 0:
                break
    except Exception:
        pass


# --------------------------------------------------------------------------- #
#  Monitor geometry and DPI
# --------------------------------------------------------------------------- #
SM_XVIRTUALSCREEN = 76
SM_YVIRTUALSCREEN = 77
SM_CXVIRTUALSCREEN = 78
SM_CYVIRTUALSCREEN = 79
MONITOR_DEFAULTTONEAREST = 2


class _MONITORINFO(ctypes.Structure):
    _fields_ = [
        ("cbSize", wintypes.DWORD),
        ("rcMonitor", wintypes.RECT),
        ("rcWork", wintypes.RECT),
        ("dwFlags", wintypes.DWORD),
    ]


def virtual_screen_rect() -> tuple[int, int, int, int]:
    """Bounding box of every monitor, including negative positions."""
    try:
        x = int(_user32.GetSystemMetrics(SM_XVIRTUALSCREEN))
        y = int(_user32.GetSystemMetrics(SM_YVIRTUALSCREEN))
        w = int(_user32.GetSystemMetrics(SM_CXVIRTUALSCREEN))
        h = int(_user32.GetSystemMetrics(SM_CYVIRTUALSCREEN))
        if w > 0 and h > 0:
            return x, y, w, h
    except Exception:
        pass
    try:
        return (0, 0, int(_user32.GetSystemMetrics(0)),
                int(_user32.GetSystemMetrics(1)))
    except Exception:
        return 0, 0, 1280, 720


def monitor_dpi_at(x: int, y: int) -> int:
    """Effective DPI of the monitor nearest the given screen point."""
    try:
        hmon = _user32.MonitorFromPoint(wintypes.POINT(int(x), int(y)),
                                        MONITOR_DEFAULTTONEAREST)
        if hmon and _shcore is not None:
            dx, dy = wintypes.UINT(), wintypes.UINT()
            if _shcore.GetDpiForMonitor(hmon, 0, ctypes.byref(dx),
                                        ctypes.byref(dy)) == 0 and dx.value:
                return int(dx.value)
    except Exception:
        pass
    try:
        value = int(_user32.GetDpiForSystem())
        if value:
            return value
    except Exception:
        pass
    return 96


def primary_dpi() -> int:
    try:
        value = int(_user32.GetDpiForSystem())
        if value:
            return value
    except Exception:
        pass
    return monitor_dpi_at(0, 0)


def apply_ui_scale(root: tk.Misc, scale: float) -> None:
    """Set global layout scale and Tk font scaling for the target monitor."""
    global SCALE
    SCALE = min(3.0, max(1.0, float(scale)))
    try:
        root.tk.call("tk", "scaling", SCALE * 96.0 / 72.0)
    except Exception:
        pass


# --------------------------------------------------------------------------- #
#  Generic helpers
# --------------------------------------------------------------------------- #
@dataclass
class Result:
    message: str
    outputs: list = field(default_factory=list)
    preview: str | None = None
    clipboard: str | None = None
    failed: int = 0


@dataclass
class Action:
    label: str
    func: Callable
    # (key, caption, min, max, default, unit, step) -> shows an inline slider first
    ask: tuple | None = None
    # (key, caption, default, hint) -> shows a text entry first
    ask_text: tuple | None = None
    # video-specific options panel (CRF, resolution, target size, GPU)
    video_options: bool = False
    # universal media converter format picker
    media_options: bool = False
    # universal document converter target picker
    office_options: bool = False


class Cancelled(Exception):
    """Raised inside a worker when the user presses Cancel."""


class ArchivePasswordError(Exception):
    """Raised when an archive needs a password or the supplied one is wrong."""


class ArchiveUnsupportedError(Exception):
    """Raised when the stdlib cannot handle an archive's encryption mode."""


@dataclass
class PasswordRequest:
    archive: str
    kind: str
    error: str = ""
    password: str | None = None
    cancelled: bool = False
    event: threading.Event = field(default_factory=threading.Event)


def require(module: str, pip_name: str | None = None):
    try:
        return importlib.import_module(module)
    except ImportError as exc:
        name = pip_name or module
        raise RuntimeError(f"Missing Python package '{name}'.\n"
                           f"Install it with:  pip install {name}") from exc


def get_unique_filename(path, is_dir: bool = False) -> Path:
    """Never overwrite: report_dropflow.pdf -> report_dropflow (1).pdf -> (2) ..."""
    p = Path(path)
    if not p.exists():
        return p
    stem, suffix = (p.name, "") if is_dir else (p.stem, p.suffix)
    n = 1
    while True:
        candidate = p.with_name(f"{stem} ({n}){suffix}")
        if not candidate.exists():
            return candidate
        n += 1


def derive_output(src: Path, ext: str | None = None) -> Path:
    return get_unique_filename(src.with_name(f"{src.stem}{SUFFIX}{ext or src.suffix}"))


def discard(paths) -> None:
    """Best-effort cleanup of partial results after Cancel or a failure."""
    for p in paths:
        try:
            p = Path(p)
            if p.is_dir():
                shutil.rmtree(p, ignore_errors=True)
            else:
                p.unlink(missing_ok=True)
        except Exception:
            pass


def fmt_size(n: float) -> str:
    for unit in ("B", "KB", "MB", "GB"):
        if n < 1024:
            return f"{n:.0f} {unit}" if unit == "B" else f"{n:.1f} {unit}"
        n /= 1024
    return f"{n:.1f} TB"


def natural_key(p: Path):
    return [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", p.name.lower())]


def clean_dropped_paths(tk_root: tk.Misc, raw: str) -> list[Path]:
    """Turn tkdnd's '{C:/My Files/a.pdf} C:/x.png' string into real, existing Paths."""
    try:
        items = list(tk_root.tk.splitlist(raw))
    except tk.TclError:
        items = re.findall(r"\{[^}]*\}|\S+", raw)
    out, seen = [], set()
    for item in items:
        s = item.strip().strip("{}").strip().strip('"').strip()
        if not s:
            continue
        try:
            p = Path(os.path.normpath(s))
        except Exception:
            continue
        key = os.path.normcase(str(p))
        if key in seen or not p.exists():
            continue
        seen.add(key)
        out.append(p)
    return out


def reveal(path: Path) -> None:
    try:
        p = Path(path)
        if p.is_dir():
            os.startfile(str(p))
        else:
            subprocess.Popen(["explorer", "/select,", str(p)],
                             creationflags=CREATE_NO_WINDOW)
    except Exception:
        try:
            os.startfile(str(Path(path).parent))
        except Exception:
            log.exception("Could not open folder")


def open_path(path: Path) -> None:
    try:
        os.startfile(str(path))
    except Exception:
        log.exception("Could not open %s", path)


def _tail(text: str, limit: int = 600) -> str:
    compact = " ".join((text or "").split())
    return compact[-limit:]


def log_path() -> Path:
    return DATA_DIR / "dropflow.log"


def log_tail(max_chars: int = 6000) -> str:
    try:
        text = log_path().read_text(encoding="utf-8", errors="replace")
        return text[-max_chars:]
    except Exception:
        return "(log file is not available yet)"


def _batch_summary(done: int, errors: list[str], noun: str = "file") -> str:
    suffix = "" if done == 1 else "s"
    text = f"{done} {noun}{suffix} done"
    if errors:
        text += f", {len(errors)} failed \u2014 see log"
    return text


def _record_failure(errors: list[str], label: str, exc: BaseException) -> None:
    errors.append(f"{label}: {exc}")
    log.exception("Batch item failed: %s", label)


def _get_bidi():
    try:
        from bidi.algorithm import get_display
    except ImportError:
        from bidi import get_display  # python-bidi >= 0.5
    return get_display


def visual_text(text: str) -> str:
    """Reshape + reorder Arabic so Tk (which has no bidi engine) shows it correctly."""
    try:
        import arabic_reshaper
        get_display = _get_bidi()
    except Exception:
        return text
    out = []
    for line in text.splitlines():
        out.append(get_display(arabic_reshaper.reshape(line)) if ARABIC_RE.search(line) else line)
    return "\n".join(out)


# --------------------------------------------------------------------------- #
#  Job context: progress + cooperative cancellation
# --------------------------------------------------------------------------- #
_PROGRESS_RE = re.compile(r"out_time_(?:us|ms)=(\d+)")


def _kill_tree(proc: subprocess.Popen) -> None:
    if proc.poll() is not None:
        return
    try:
        subprocess.run(["taskkill", "/PID", str(proc.pid), "/T", "/F"],
                       stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL,
                       stderr=subprocess.DEVNULL, creationflags=CREATE_NO_WINDOW,
                       timeout=8)
    except Exception:
        try:
            proc.kill()
        except Exception:
            pass


class JobContext:
    """Carries progress messages and the cancel signal into worker actions."""

    def __init__(self, progress: Callable[[str], None], cfg: dict | None = None,
                 password_cb: Callable[[PasswordRequest], None] | None = None):
        self._progress = progress
        self.cfg = cfg or {}
        self._password_cb = password_cb
        self.passwords: dict[str, str] = {}
        self._cancel = threading.Event()
        self._proc: subprocess.Popen | None = None
        self._lock = threading.Lock()

    @property
    def cancelled(self) -> bool:
        return self._cancel.is_set()

    @property
    def cancel_event(self) -> threading.Event:
        return self._cancel

    def check(self) -> None:
        if self._cancel.is_set():
            raise Cancelled()

    def progress(self, text: str) -> None:
        self.check()
        self._emit(text)

    def _emit(self, text: str) -> None:
        try:
            self._progress(text)
        except Exception:
            pass

    def ask_password(self, archive: str, kind: str, error: str = "") -> str | None:
        """Ask the Tk UI for an archive password and wait for the answer."""
        if self._password_cb is None:
            return None
        request = PasswordRequest(archive=archive, kind=kind, error=error)
        try:
            self._password_cb(request)
        except Exception:
            log.exception("Could not open the password dialog")
            return None
        while not request.event.wait(0.1):
            if self._cancel.is_set():
                request.cancelled = True
                request.event.set()
                return None
        if request.cancelled or self._cancel.is_set():
            return None
        return request.password

    def cancel(self) -> None:
        # Only set the flag here: this runs on the Tk thread, so it must never
        # block.  The worker watching run() kills the helper process promptly.
        self._cancel.set()

    def run(self, cmd: list[str], *, timeout: float | None = None,
            duration: float | None = None, label: str = "Working") -> subprocess.CompletedProcess:
        """Run a helper process with no console, live progress and Cancel support."""
        self.check()
        proc = subprocess.Popen(
            cmd, stdin=subprocess.DEVNULL, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
            text=True, encoding="utf-8", errors="replace", bufsize=1,
            creationflags=CREATE_NO_WINDOW)
        out_lines: list[str] = []
        err_lines: list[str] = []
        last_pct = [-1]

        def reader(pipe, sink, parse_progress: bool):
            try:
                for line in iter(pipe.readline, ""):
                    sink.append(line)
                    if parse_progress and duration and duration > 0:
                        m = _PROGRESS_RE.search(line)
                        if m:
                            pct = min(99, max(0, int(int(m.group(1)) / (duration * 1_000_000) * 100)))
                            if pct != last_pct[0]:
                                last_pct[0] = pct
                                self._emit(f"{label}\u2026 {pct}%")
            except Exception:
                pass
            finally:
                try:
                    pipe.close()
                except Exception:
                    pass

        t_out = threading.Thread(target=reader, args=(proc.stdout, out_lines, True), daemon=True)
        t_err = threading.Thread(target=reader, args=(proc.stderr, err_lines, False), daemon=True)
        t_out.start()
        t_err.start()
        with self._lock:
            self._proc = proc
        deadline = time.monotonic() + timeout if timeout else None
        rc = None
        try:
            while True:
                rc = proc.poll()
                if rc is not None:
                    break
                if self._cancel.is_set():
                    _kill_tree(proc)
                    try:
                        proc.wait(timeout=8)
                    except Exception:
                        pass
                    raise Cancelled()
                if deadline is not None and time.monotonic() > deadline:
                    _kill_tree(proc)
                    try:
                        proc.wait(timeout=8)
                    except Exception:
                        pass
                    raise TimeoutError(f"{label} took too long and was stopped.")
                time.sleep(0.08)
        finally:
            with self._lock:
                self._proc = None
        t_out.join(2)
        t_err.join(2)
        if self._cancel.is_set():
            raise Cancelled()
        return subprocess.CompletedProcess(cmd, rc, "".join(out_lines), "".join(err_lines))


# --------------------------------------------------------------------------- #
#  Optional helper program discovery
# --------------------------------------------------------------------------- #
_SOFFICE: str | None | bool = False
_FFMPEG: str | None | bool = False
_FFPROBE: str | None | bool = False
_FFMPEG_SCORE = -1
_FFMPEG_ENCODERS: set[str] | None = None
_TESSERACT: str | None | bool = False


def _run_hidden(cmd: list[str], timeout: float = 20) -> subprocess.CompletedProcess:
    return subprocess.run(cmd, stdin=subprocess.DEVNULL, stdout=subprocess.PIPE,
                          stderr=subprocess.STDOUT, text=True, encoding="utf-8",
                          errors="replace", timeout=timeout,
                          creationflags=CREATE_NO_WINDOW)


def _find_tesseract() -> str | None:
    global _TESSERACT
    if _TESSERACT is not False:
        return _TESSERACT
    cands = []
    if os.environ.get("TESSERACT_CMD"):
        cands.append(os.environ["TESSERACT_CMD"])
    found = shutil.which("tesseract")
    if found:
        cands.append(found)
    cands += [
        r"C:\Program Files\Tesseract-OCR\tesseract.exe",
        r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
        str(Path(os.environ.get("LOCALAPPDATA", "")) / "Programs" / "Tesseract-OCR" / "tesseract.exe"),
        str(Path(os.environ.get("LOCALAPPDATA", "")) / "Tesseract-OCR" / "tesseract.exe"),
        str(app_dir() / "Tesseract-OCR" / "tesseract.exe"),
        str(app_dir() / "tesseract.exe"),
    ]
    _TESSERACT = None
    for c in cands:
        try:
            if c and Path(c).is_file():
                _TESSERACT = c
                break
        except Exception:
            continue
    return _TESSERACT


def _find_soffice() -> str | None:
    global _SOFFICE
    if _SOFFICE is not False:
        return _SOFFICE
    cands = []
    if os.environ.get("SOFFICE_PATH"):
        cands.append(os.environ["SOFFICE_PATH"])
    for name in ("soffice", "soffice.com"):
        found = shutil.which(name)
        if found:
            cands.append(found)
    cands += [
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        str(Path(os.environ.get("LOCALAPPDATA", "")) / "Programs" / "LibreOffice" / "program" / "soffice.exe"),
        str(app_dir() / "LibreOffice" / "program" / "soffice.exe"),
    ]
    _SOFFICE = None
    for c in cands:
        try:
            if c and Path(c).is_file():
                _SOFFICE = c
                break
        except Exception:
            continue
    return _SOFFICE


def _ffmpeg_candidates() -> list[str]:
    out: list[str] = []
    seen: set[str] = set()

    def add(value) -> None:
        if not value:
            return
        try:
            p = Path(value)
            if p.is_file():
                key = os.path.normcase(str(p))
                if key not in seen:
                    seen.add(key)
                    out.append(str(p))
        except Exception:
            pass

    add(os.environ.get("FFMPEG_PATH"))
    try:
        import imageio_ffmpeg  # optional, gives a bundled modern build
        add(imageio_ffmpeg.get_ffmpeg_exe())
    except Exception:
        pass
    add(shutil.which("ffmpeg"))
    add(str(app_dir() / "ffmpeg.exe"))
    add(str(app_dir() / "bin" / "ffmpeg.exe"))
    add(str(app_dir() / "ffmpeg" / "bin" / "ffmpeg.exe"))
    add(str(Path(os.environ.get("LOCALAPPDATA", "")) / "Microsoft" / "WinGet" / "Links" / "ffmpeg.exe"))
    add(r"C:\Program Files\ffmpeg\bin\ffmpeg.exe")
    add(r"C:\Program Files (x86)\ffmpeg\bin\ffmpeg.exe")
    add(r"C:\ffmpeg\bin\ffmpeg.exe")
    add(r"C:\ProgramData\chocolatey\bin\ffmpeg.exe")
    return out


def _ffmpeg_version_score(path: str) -> int:
    try:
        res = _run_hidden([path, "-version"], timeout=10)
        text = (res.stdout or "") + (res.stderr or "")
        m = re.search(r"ffmpeg version\s+(?:n-?)?(\d+)\.(\d+)", text, re.I)
        if m:
            return int(m.group(1)) * 100 + int(m.group(2))
        if "ffmpeg version" in text:
            return 1
    except Exception:
        return -1
    return -1


def _find_ffmpeg() -> str | None:
    global _FFMPEG, _FFMPEG_SCORE, _FFPROBE
    if _FFMPEG is not False:
        return _FFMPEG
    best, best_score = None, -1
    for cand in _ffmpeg_candidates():
        score = _ffmpeg_version_score(cand)
        if score > best_score:
            best, best_score = cand, score
    _FFMPEG = best
    _FFMPEG_SCORE = best_score
    _FFPROBE = None
    if best:
        sibling = Path(best).with_name("ffprobe.exe")
        if sibling.is_file():
            _FFPROBE = str(sibling)
        else:
            found = shutil.which("ffprobe")
            if found:
                _FFPROBE = found
    return _FFMPEG


def _find_ffprobe() -> str | None:
    _find_ffmpeg()
    return _FFPROBE


def _ffmpeg_encoders() -> set[str]:
    global _FFMPEG_ENCODERS
    if _FFMPEG_ENCODERS is not None:
        return _FFMPEG_ENCODERS
    encoders: set[str] = set()
    exe = _find_ffmpeg()
    if exe:
        try:
            res = _run_hidden([exe, "-encoders"], timeout=15)
            for line in (res.stdout or "").splitlines():
                parts = line.split()
                if len(parts) >= 2 and len(parts[0]) >= 3:
                    encoders.add(parts[1])
        except Exception:
            pass
    _FFMPEG_ENCODERS = encoders
    return encoders


def _ffmpeg_audio_choice(prefer_mp3: bool = True) -> tuple[str, str]:
    """Return (encoder, extension) for an audio track that can go into MP4/MP3."""
    enc = _ffmpeg_encoders()
    if prefer_mp3 and "libmp3lame" in enc:
        return "libmp3lame", ".mp3"
    if "aac" in enc:
        return "aac", ".m4a"
    if "libmp3lame" in enc:
        return "libmp3lame", ".mp3"
    return "copy", ".m4a"


def _ffprobe_duration(path: Path) -> float | None:
    exe = _find_ffprobe()
    if not exe:
        return None
    try:
        res = _run_hidden([
            exe, "-v", "error", "-show_entries", "format=duration",
            "-of", "default=noprint_wrappers=1:nokey=1", str(path),
        ], timeout=20)
        value = (res.stdout or "").strip().splitlines()[0]
        return float(value)
    except Exception:
        return None


def _ffprobe_audio_info(path: Path) -> tuple[str | None, int | None]:
    """Return (codec_name, bit_rate) for the first audio stream, if any."""
    exe = _find_ffprobe()
    if not exe:
        return None, None
    try:
        res = _run_hidden([
            exe, "-v", "error", "-select_streams", "a:0",
            "-show_entries", "stream=codec_name,bit_rate",
            "-of", "default=noprint_wrappers=1", str(path),
        ], timeout=20)
        codec, bitrate = None, None
        for line in (res.stdout or "").splitlines():
            if line.startswith("codec_name="):
                codec = line.split("=", 1)[1].strip().lower() or None
            elif line.startswith("bit_rate="):
                raw = line.split("=", 1)[1].strip()
                if raw.isdigit():
                    bitrate = int(raw)
        return codec, bitrate
    except Exception:
        return None, None


# --------------------------------------------------------------------------- #
#  Image helpers
# --------------------------------------------------------------------------- #
def load_image(path: Path) -> Image.Image:
    im = Image.open(path)
    im = ImageOps.exif_transpose(im)
    im.load()
    return im


def flatten_rgb(im: Image.Image) -> Image.Image:
    if im.mode in ("RGBA", "LA") or (im.mode == "P" and "transparency" in im.info):
        rgba = im.convert("RGBA")
        bg = Image.new("RGB", rgba.size, (255, 255, 255))
        bg.paste(rgba, mask=rgba.split()[-1])
        return bg
    return im.convert("RGB")


def _rgb_or_rgba(im: Image.Image) -> Image.Image:
    has_alpha = "A" in im.getbands() or (im.mode == "P" and "transparency" in im.info)
    return im.convert("RGBA" if has_alpha else "RGB")


def _png_safe(im: Image.Image) -> Image.Image:
    return im if im.mode in ("1", "L", "LA", "P", "RGB", "RGBA", "I;16") else im.convert("RGBA")


_RESAMPLE = getattr(Image, "Resampling", Image).LANCZOS
_QUANTIZE = getattr(Image, "Quantize", Image).FASTOCTREE


def _image_header(path: Path) -> tuple[int, int, str, int]:
    """Read only the image header: size, format and EXIF orientation."""
    with Image.open(path) as im:
        w, h = im.size
        fmt = (im.format or "").upper()
        try:
            orientation = int(im.getexif().get(274, 1) or 1)
        except Exception:
            orientation = 1
    if orientation not in (1, 2, 3, 4, 5, 6, 7, 8):
        orientation = 1
    return w, h, fmt, orientation


def _pil_image_bytes(im: Image.Image) -> tuple[bytes, str]:
    buf = io.BytesIO()
    has_alpha = ("A" in im.getbands() or
                 (im.mode == "P" and "transparency" in im.info))
    if has_alpha:
        im.convert("RGBA").save(buf, "PNG", optimize=True)
        return buf.getvalue(), "PNG"
    im.convert("RGB").save(buf, "JPEG", quality=92, optimize=True, progressive=True)
    return buf.getvalue(), "JPEG"


def _new_image_page(doc, data: bytes, width_pt: float, height_pt: float, rotation: int = 0):
    page = doc.new_page(width=width_pt, height=height_pt)
    try:
        page.insert_image(page.rect, stream=data)
    except Exception:
        try:
            doc.delete_page(page.number)
        except Exception:
            pass
        raise
    if rotation:
        page.set_rotation(rotation)
    return page


def _add_image_page(doc, path: Path, ctx: JobContext) -> None:
    """Append one PDF page for an image, embedding JPEG/PNG bytes directly when possible."""
    ctx.check()
    w, h, fmt, orientation = _image_header(path)
    rotate = {3: 180, 6: 90, 8: 270}.get(orientation, 0)
    direct = fmt in ("JPEG", "PNG") and orientation in (1, 3, 6, 8)
    data = b""
    if direct:
        data = path.read_bytes()
        pw = w * 72.0 / IMAGE_PDF_DPI
        ph = h * 72.0 / IMAGE_PDF_DPI
        try:
            _new_image_page(doc, data, pw, ph, rotate)
            return
        except Exception:
            log.info("Direct embed failed for %s; re-encoding", path.name)
        # _new_image_page already removed its failed page.

    # Fallback: decode/rotate/convert with Pillow, then embed a clean stream.
    with Image.open(path) as im:
        im = ImageOps.exif_transpose(im)
        disp_w, disp_h = im.size
        data, _ = _pil_image_bytes(im)
    _new_image_page(doc, data, disp_w * 72.0 / IMAGE_PDF_DPI,
                    disp_h * 72.0 / IMAGE_PDF_DPI)


# --------------------------------------------------------------------------- #
#  Actions: PDF
# --------------------------------------------------------------------------- #
def _open_pdf(path: Path):
    doc = fitz.open(str(path))
    if doc.needs_pass:
        doc.close()
        raise RuntimeError(f"'{path.name}' is password-protected.")
    return doc


def _pdf_files(paths) -> list[Path]:
    return [Path(p) for p in paths if Path(p).is_file() and Path(p).suffix.lower() == ".pdf"]


def act_merge_pdfs(paths, opts, ctx: JobContext) -> Result:
    pdfs = _pdf_files(paths)
    if len(pdfs) < 2:
        raise RuntimeError("Drop at least two PDF files to merge.")
    out = get_unique_filename(pdfs[0].with_name(f"{pdfs[0].stem}_merged{SUFFIX}.pdf"))
    merged = fitz.open()
    errors: list[str] = []
    pages = 0
    ok = False
    try:
        for i, p in enumerate(pdfs, 1):
            ctx.check()
            ctx.progress(f"Merging {i}/{len(pdfs)}: {p.name}")
            try:
                src = _open_pdf(p)
                try:
                    merged.insert_pdf(src)
                    pages += src.page_count
                finally:
                    src.close()
            except Cancelled:
                raise
            except Exception as exc:
                _record_failure(errors, p.name, exc)
        if pages == 0:
            raise RuntimeError("None of the selected PDFs could be read.")
        merged.set_metadata({"title": out.stem, "producer": APP_NAME})
        ctx.progress("Saving merged PDF\u2026")
        merged.save(str(out), garbage=3, deflate=True)
        ok = True
    except Cancelled:
        discard([out])
        raise
    finally:
        merged.close()
        if not ok:
            out.unlink(missing_ok=True)
    good = len(pdfs) - len(errors)
    note = f" ({len(errors)} failed \u2014 see log)" if errors else ""
    return Result(f"Merged {good} PDF(s) into {out.name} ({pages} page(s), "
                  f"{fmt_size(out.stat().st_size)}){note}", [out], failed=len(errors))


def act_compress_pdfs(paths, opts, ctx: JobContext) -> Result:
    pdfs = _pdf_files(paths)
    if not pdfs:
        raise RuntimeError("No PDF files to compress.")
    q = int(opts.get("quality", 60))
    dpi_target = int(72 + q * 0.9)
    outs, before, after, skipped = [], 0, 0, 0
    errors: list[str] = []
    try:
        for i, p in enumerate(pdfs, 1):
            ctx.check()
            ctx.progress(f"Compressing {i}/{len(pdfs)}: {p.name}")
            out = None
            try:
                out = derive_output(p, ".pdf")
                doc = _open_pdf(p)
                try:
                    if hasattr(doc, "rewrite_images"):
                        try:
                            doc.rewrite_images(dpi_threshold=dpi_target + 30,
                                               dpi_target=dpi_target, quality=q)
                        except Exception:
                            log.exception("rewrite_images failed (continuing with lossless pass)")
                    doc.save(str(out), garbage=4, deflate=True, deflate_images=True,
                             deflate_fonts=True, clean=True)
                finally:
                    doc.close()
                ctx.check()
                old, new = p.stat().st_size, out.stat().st_size
                if new >= old:  # nothing gained, don't litter the folder
                    out.unlink(missing_ok=True)
                    skipped += 1
                    continue
                outs.append(out)
                before += old
                after += new
            except Cancelled:
                if out is not None:
                    out.unlink(missing_ok=True)
                raise
            except Exception as exc:
                if out is not None:
                    out.unlink(missing_ok=True)
                _record_failure(errors, p.name, exc)
    except Cancelled:
        discard(outs)
        raise
    if not outs:
        if errors:
            raise RuntimeError("No PDFs could be compressed:\n" + "\n".join(errors[:5]))
        return Result("These PDFs are already as small as they can get.")
    saved = 100 - after * 100 / before
    note = f" ({len(errors)} failed \u2014 see log)" if errors else ""
    extra = f" ({skipped} left unchanged)" if skipped else ""
    return Result(f"Compressed {len(outs)} PDF(s): {fmt_size(before)} -> {fmt_size(after)} "
                  f"(-{saved:.0f}%){extra}{note}", outs, failed=len(errors))


def act_split_pdf(paths, opts, ctx: JobContext) -> Result:
    pdfs = _pdf_files(paths)
    if not pdfs:
        raise RuntimeError("No PDF files to split.")
    per = max(1, int(opts.get("pages", 1)))
    folders, count = [], 0
    errors: list[str] = []
    try:
        for p in pdfs:
            ctx.check()
            ctx.progress(f"Splitting {p.name}\u2026")
            folder = None
            try:
                doc = _open_pdf(p)
                try:
                    folder = get_unique_filename(p.with_name(f"{p.stem}_split{SUFFIX}"),
                                                 is_dir=True)
                    folder.mkdir(parents=True)
                    total = doc.page_count
                    if total == 0:
                        raise RuntimeError("PDF has no pages.")
                    width = max(3, len(str(total)))
                    for start in range(0, total, per):
                        ctx.check()
                        end = min(start + per, total)
                        if per == 1:
                            name = f"{p.stem}_page_{start + 1:0{width}d}{SUFFIX}.pdf"
                        else:
                            name = (f"{p.stem}_pages_{start + 1:0{width}d}-"
                                    f"{end:0{width}d}{SUFFIX}.pdf")
                        out = get_unique_filename(folder / name)
                        part = fitz.open()
                        try:
                            part.insert_pdf(doc, from_page=start, to_page=end - 1)
                            part.save(str(out), garbage=3, deflate=True)
                        finally:
                            part.close()
                        count += 1
                        ctx.progress(f"Splitting {p.name}: {end}/{total} pages")
                    folders.append(folder)
                finally:
                    doc.close()
            except Cancelled:
                if folder is not None:
                    discard([folder])
                raise
            except Exception as exc:
                if folder is not None:
                    discard([folder])
                _record_failure(errors, p.name, exc)
    except Cancelled:
        discard(folders)
        raise
    if not folders:
        if errors:
            raise RuntimeError("No PDFs could be split:\n" + "\n".join(errors[:5]))
        return Result("No PDFs were split.")
    note = f" ({len(errors)} failed \u2014 see log)" if errors else ""
    return Result(f"Split {len(folders)} PDF(s) into {count} file(s) in "
                  f"{len(folders)} folder(s){note}", folders, failed=len(errors))


def _parse_page_ranges(spec, page_count: int) -> list[tuple[int, int]]:
    """Parse '1-3, 5, 7-9' (also 5-, -5, all) into 1-based inclusive ranges."""
    text = str(spec or "").strip().lower()
    text = text.replace("\u2013", "-").replace("\u2014", "-")
    text = re.sub(r"\s*-\s*", "-", text)
    if not text:
        raise ValueError("Enter page numbers, e.g. 1-3, 5, 7-9 or all.")
    if text in ("all", "*"):
        if page_count < 1:
            raise ValueError("This PDF has no pages.")
        return [(1, page_count)]
    ranges: list[tuple[int, int]] = []
    for part in re.split(r"[,;\s]+", text):
        if not part:
            continue
        if part.isdigit():
            start = end = int(part)
        else:
            match = re.fullmatch(r"(\d+)?\s*-\s*(\d+)?", part)
            if not match or (not match.group(1) and not match.group(2)):
                raise ValueError(f"Invalid page range: '{part}'")
            start = int(match.group(1)) if match.group(1) else 1
            end = int(match.group(2)) if match.group(2) else page_count
        if start < 1 or end < 1 or start > page_count or end > page_count:
            raise ValueError(f"Range '{part}' is outside 1-{page_count}.")
        if start > end:
            raise ValueError(f"Range '{part}' starts after it ends.")
        ranges.append((start, end))
    if not ranges:
        raise ValueError("Enter at least one page or range.")
    return ranges


def _range_label(ranges: list[tuple[int, int]]) -> str:
    parts = [str(a) if a == b else f"{a}-{b}" for a, b in ranges]
    label = re.sub(r"[^0-9A-Za-z_-]+", "_", "_".join(parts)).strip("_")
    return (label[:40] or "range")


def act_extract_pages(paths, opts, ctx: JobContext) -> Result:
    """Extract the selected pages of each PDF into one PDF per source."""
    pdfs = _pdf_files(paths)
    if not pdfs:
        raise RuntimeError("No PDF files to extract pages from.")
    spec = opts.get("pages", "")
    outs, errors = [], []
    try:
        for p in pdfs:
            ctx.check()
            ctx.progress(f"Extracting pages from {p.name}\u2026")
            out = None
            try:
                doc = _open_pdf(p)
                try:
                    ranges = _parse_page_ranges(spec, doc.page_count)
                    label = _range_label(ranges)
                    out = get_unique_filename(
                        p.with_name(f"{p.stem}_pages_{label}{SUFFIX}.pdf"))
                    dst = fitz.open()
                    try:
                        for start, end in ranges:
                            ctx.check()
                            dst.insert_pdf(doc, from_page=start - 1, to_page=end - 1)
                        dst.set_metadata({"title": out.stem, "producer": APP_NAME})
                        dst.save(str(out), garbage=3, deflate=True)
                    finally:
                        dst.close()
                finally:
                    doc.close()
                outs.append(out)
            except Cancelled:
                if out is not None:
                    out.unlink(missing_ok=True)
                raise
            except Exception as exc:
                if out is not None:
                    out.unlink(missing_ok=True)
                _record_failure(errors, p.name, exc)
    except Cancelled:
        discard(outs)
        raise
    if not outs:
        if errors:
            raise RuntimeError("No pages could be extracted:\n" + "\n".join(errors[:5]))
        return Result("No pages were extracted.")
    note = f" ({len(errors)} failed \u2014 see log)" if errors else ""
    return Result(f"Extracted page ranges from {len(outs)} PDF(s){note}", outs,
                  failed=len(errors))


def act_split_ranges(paths, opts, ctx: JobContext) -> Result:
    """Split each PDF into one file per explicit page range."""
    pdfs = _pdf_files(paths)
    if not pdfs:
        raise RuntimeError("No PDF files to split at ranges.")
    spec = opts.get("pages", "")
    folders, count, errors = [], 0, []
    try:
        for p in pdfs:
            ctx.check()
            ctx.progress(f"Splitting {p.name} at page ranges\u2026")
            folder = None
            try:
                doc = _open_pdf(p)
                try:
                    ranges = _parse_page_ranges(spec, doc.page_count)
                    folder = get_unique_filename(
                        p.with_name(f"{p.stem}_ranges{SUFFIX}"), is_dir=True)
                    folder.mkdir(parents=True)
                    width = max(3, len(str(doc.page_count)))
                    for start, end in ranges:
                        ctx.check()
                        if start == end:
                            name = (f"{p.stem}_page_{start:0{width}d}"
                                    f"{SUFFIX}.pdf")
                        else:
                            name = (f"{p.stem}_pages_{start:0{width}d}-"
                                    f"{end:0{width}d}{SUFFIX}.pdf")
                        out = get_unique_filename(folder / name)
                        part = fitz.open()
                        try:
                            part.insert_pdf(doc, from_page=start - 1, to_page=end - 1)
                            part.save(str(out), garbage=3, deflate=True)
                        finally:
                            part.close()
                        count += 1
                    folders.append(folder)
                finally:
                    doc.close()
            except Cancelled:
                if folder is not None:
                    discard([folder])
                raise
            except Exception as exc:
                if folder is not None:
                    discard([folder])
                _record_failure(errors, p.name, exc)
    except Cancelled:
        discard(folders)
        raise
    if not folders:
        if errors:
            raise RuntimeError("No PDFs could be split:\n" + "\n".join(errors[:5]))
        return Result("No PDFs were split.")
    note = f" ({len(errors)} failed \u2014 see log)" if errors else ""
    return Result(f"Split {len(folders)} PDF(s) into {count} range file(s) in "
                  f"{len(folders)} folder(s){note}", folders, failed=len(errors))


def make_pdf_to_images(fmt: str) -> Callable:
    ext = ".png" if fmt == "PNG" else ".jpg"

    def run(paths, opts, ctx: JobContext) -> Result:
        pdfs = _pdf_files(paths)
        if not pdfs:
            raise RuntimeError("No PDF files to export.")
        if fmt == "PNG":
            dpi = max(36, int(opts.get("dpi", 150)))
            quality = 95
        else:
            dpi, quality = 150, max(40, int(opts.get("quality", 85)))
        folders, count = [], 0
        errors: list[str] = []
        try:
            for p in pdfs:
                ctx.check()
                ctx.progress(f"Opening {p.name}\u2026")
                folder = None
                try:
                    doc = _open_pdf(p)
                    try:
                        folder = get_unique_filename(
                            p.with_name(f"{p.stem}_pages{SUFFIX}"), is_dir=True)
                        folder.mkdir(parents=True)
                        total = doc.page_count
                        if total == 0:
                            raise RuntimeError("PDF has no pages.")
                        width = max(3, len(str(total)))
                        for i in range(total):
                            ctx.check()
                            page = doc.load_page(i)
                            pix = page.get_pixmap(dpi=dpi, alpha=False)
                            out = folder / f"page_{i + 1:0{width}d}{ext}"
                            if fmt == "PNG":
                                pix.save(str(out))
                            else:
                                pix.save(str(out), jpg_quality=quality)
                            del pix
                            count += 1
                            ctx.progress(f"Rendering {p.name}: page {i + 1}/{total}")
                        folders.append(folder)
                    finally:
                        doc.close()
                except Cancelled:
                    if folder is not None:
                        discard([folder])
                    raise
                except Exception as exc:
                    if folder is not None:
                        discard([folder])
                    _record_failure(errors, p.name, exc)
        except Cancelled:
            discard(folders)
            raise
        if not folders:
            if errors:
                raise RuntimeError("No PDFs could be exported:\n" + "\n".join(errors[:5]))
            return Result("No pages were exported.")
        note = f" ({len(errors)} failed \u2014 see log)" if errors else ""
        return Result(f"Exported {count} page(s) as {fmt} images into "
                      f"{len(folders)} folder(s){note}", folders, failed=len(errors))

    return run


_TESS_LANGS: list[str] | None = None


def _available_ocr_languages() -> list[str]:
    """Installed Tesseract language codes, cached for this session."""
    global _TESS_LANGS
    if _TESS_LANGS is not None:
        return _TESS_LANGS
    cmd = _find_tesseract()
    if not cmd:
        _TESS_LANGS = []
        return _TESS_LANGS
    try:
        pytesseract = require("pytesseract")
        pytesseract.pytesseract.tesseract_cmd = cmd
        _TESS_LANGS = sorted({
            lang for lang in pytesseract.get_languages(config="")
            if lang != "osd" and "/" not in lang
        })
    except Exception:
        _TESS_LANGS = []
    return _TESS_LANGS


def _ocr_setup(preferred=None):
    pytesseract = require("pytesseract")
    cmd = _find_tesseract()
    if not cmd:
        raise RuntimeError(
            "Tesseract OCR engine was not found.\n\n"
            "Install the UB Mannheim Windows build, tick the language data you "
            "need during setup, then restart DropFlow.")
    pytesseract.pytesseract.tesseract_cmd = cmd
    available = set(_available_ocr_languages())
    if not available:
        try:
            available = set(pytesseract.get_languages(config=""))
        except Exception:
            available = set()

    if isinstance(preferred, (list, tuple)):
        wanted = [str(x).strip().lower() for x in preferred if str(x).strip()]
    else:
        wanted = [x.lower() for x in re.split(r"[+,;\s]+", str(preferred or "")) if x]
    if not wanted:
        wanted = ["ara", "eng"]
    langs = [lang for lang in wanted if lang in available]
    missing = [lang for lang in wanted if lang not in available]
    if not langs:
        fallback = [lang for lang in ("ara", "eng") if lang in available]
        if fallback:
            langs = fallback
        elif available:
            langs = [sorted(available)[0]]
        else:
            raise RuntimeError(
                "Tesseract has no language data installed.\n\n"
                "Run the Tesseract installer again and tick the languages you need.")
    return pytesseract, langs, missing


def act_ocr(paths, opts, ctx: JobContext) -> Result:
    images = [Path(p) for p in paths if Path(p).is_file() and Path(p).suffix.lower() in IMG_EXTS]
    if not images:
        raise RuntimeError("No images to read.")
    pytesseract, langs, missing = _ocr_setup(ctx.cfg.get("ocr_langs"))
    lang = "+".join(langs)
    outs, previews, errors = [], [], []
    try:
        for i, p in enumerate(images, 1):
            ctx.check()
            ctx.progress(f"Reading {i}/{len(images)}: {p.name}")
            try:
                src_img = load_image(p)
                try:
                    img = src_img.convert("L")
                finally:
                    src_img.close()
                try:
                    if max(img.size) < 1600:
                        img = img.resize((img.width * 2, img.height * 2), _RESAMPLE)
                    img = ImageOps.autocontrast(img)
                    text = pytesseract.image_to_string(img, lang=lang).strip()
                finally:
                    img.close()
                ctx.check()
                if not text:
                    continue
                out = derive_output(p, ".txt")
                out.write_text(text, encoding="utf-8-sig")
                outs.append(out)
                previews.append(text)
            except Cancelled:
                raise
            except Exception as exc:
                _record_failure(errors, p.name, exc)
    except Cancelled:
        discard(outs)
        raise
    if not outs:
        if errors:
            raise RuntimeError("No images could be read:\n" + "\n".join(errors[:5]))
        return Result("No text was detected in these images.")
    note = f" (language data missing: {', '.join(missing)})" if missing else ""
    if errors:
        note += f" ({len(errors)} failed \u2014 see log)"
    first = previews[0] if len(previews) == 1 else "\n\n".join(previews[:3])
    return Result(f"Extracted text from {len(outs)} image(s) -> .txt{note}", outs,
                  preview=visual_text(first), clipboard=first if len(outs) == 1 else None,
                  failed=len(errors))


def act_ocr_pdf(paths, opts, ctx: JobContext) -> Result:
    pdfs = _pdf_files(paths)
    if not pdfs:
        raise RuntimeError("No PDF files to read.")
    dpi = max(100, int(opts.get("dpi", 250)))
    preferred = ctx.cfg.get("ocr_langs")
    pytesseract = None
    langs: list[str] = []
    missing: list[str] = []
    lang = ""
    outs, previews, errors = [], [], []
    failed_files = 0
    native_pages = ocr_pages = 0
    try:
        for p in pdfs:
            ctx.check()
            file_failed = False
            ocr_error = None
            try:
                doc = _open_pdf(p)
                try:
                    total = doc.page_count
                    chunks = []
                    for i in range(total):
                        ctx.check()
                        ctx.progress(f"Reading {p.name}: page {i + 1}/{total}")
                        page = doc.load_page(i)
                        text = page.get_text("text").strip()
                        if len(text) >= 15:
                            native_pages += 1
                        else:
                            if pytesseract is None and ocr_error is None:
                                try:
                                    pytesseract, langs, missing = _ocr_setup(preferred)
                                    lang = "+".join(langs)
                                except Exception as exc:
                                    ocr_error = exc
                                    file_failed = True
                                    errors.append(f"{p.name}: {exc}")
                                    log.exception("OCR setup failed for %s", p.name)
                            if ocr_error is not None:
                                continue
                            pix = page.get_pixmap(dpi=dpi, colorspace=fitz.csGRAY, alpha=False)
                            try:
                                img = Image.frombytes("L", (pix.width, pix.height), pix.samples)
                                try:
                                    img = ImageOps.autocontrast(img)
                                    text = pytesseract.image_to_string(img, lang=lang).strip()
                                finally:
                                    img.close()
                            finally:
                                del pix
                            ocr_pages += 1
                        if text:
                            chunks.append(f"===== Page {i + 1} =====\n{text}")
                    if chunks:
                        out = derive_output(p, ".txt")
                        content = "\n\n".join(chunks)
                        out.write_text(content, encoding="utf-8-sig")
                        outs.append(out)
                        previews.append(content)
                finally:
                    doc.close()
            except Cancelled:
                raise
            except Exception as exc:
                file_failed = True
                errors.append(f"{p.name}: {exc}")
                log.exception("OCR failed for %s", p.name)
            if file_failed:
                failed_files += 1
    except Cancelled:
        discard(outs)
        raise
    if not outs:
        if errors:
            raise RuntimeError("No text could be extracted:\n" + "\n".join(errors[:5]))
        return Result("No text was detected in these PDFs.")
    note = f" (language data missing: {', '.join(missing)})" if missing else ""
    if failed_files:
        note += f" ({failed_files} PDF(s) had errors \u2014 see log)"
    first = previews[0] if len(previews) == 1 else "\n\n".join(previews[:2])
    return Result(
        f"Extracted {len(outs)} PDF(s) to text ({native_pages} page(s) from the text layer, "
        f"{ocr_pages} OCRed){note}", outs,
        preview=visual_text(first), clipboard=first if len(outs) == 1 else None,
        failed=failed_files)


def _tesseract_pdf_bytes(img: Image.Image, pytesseract, lang: str, dpi: float) -> bytes:
    """Return a one-page PDF with the image and Tesseract's invisible text layer."""
    rgb = img if img.mode == "RGB" else img.convert("RGB")
    try:
        rgb.info["dpi"] = (float(dpi), float(dpi))
        return pytesseract.image_to_pdf_or_hocr(rgb, extension="pdf", lang=lang)
    finally:
        if rgb is not img:
            rgb.close()


def _insert_tesseract_page(dst, data: bytes, width: float, height: float) -> None:
    """Place a Tesseract page into dst at the exact target size, preserving text."""
    src = fitz.open(stream=data, filetype="pdf")
    page = None
    try:
        page = dst.new_page(width=max(1.0, width), height=max(1.0, height))
        page.show_pdf_page(page.rect, src, 0)
    except Exception:
        if page is not None:
            try:
                dst.delete_page(page.number)
            except Exception:
                pass
        raise
    finally:
        src.close()


def _ocr_language_note(langs: list[str], missing: list[str]) -> str:
    if missing:
        return f" (language data missing: {', '.join(missing)})"
    return ""


def act_ocr_to_pdf(paths, opts, ctx: JobContext) -> Result:
    """OCR images into one combined, searchable PDF."""
    images = [Path(p) for p in paths if Path(p).is_file() and Path(p).suffix.lower() in IMG_EXTS]
    if not images:
        raise RuntimeError("No images to OCR into a PDF.")
    pytesseract, langs, missing = _ocr_setup(ctx.cfg.get("ocr_langs"))
    lang = "+".join(langs)
    out = get_unique_filename(images[0].with_name(f"{images[0].stem}_searchable{SUFFIX}.pdf"))
    doc = fitz.open()
    errors: list[str] = []
    ocred = 0
    failed_items = 0
    ok = False
    try:
        for i, p in enumerate(images, 1):
            ctx.check()
            ctx.progress(f"OCR {i}/{len(images)}: {p.name}")
            try:
                src_img = load_image(p)
                try:
                    rgb = flatten_rgb(src_img)
                finally:
                    src_img.close()
                try:
                    page_w = rgb.width * 72.0 / IMAGE_PDF_DPI
                    page_h = rgb.height * 72.0 / IMAGE_PDF_DPI
                    data = _tesseract_pdf_bytes(rgb, pytesseract, lang, IMAGE_PDF_DPI)
                finally:
                    rgb.close()
                _insert_tesseract_page(doc, data, page_w, page_h)
                ocred += 1
            except Cancelled:
                raise
            except Exception as exc:
                _record_failure(errors, p.name, exc)
                failed_items += 1
                # Keep the page in the document even if OCR failed.
                try:
                    _add_image_page(doc, p, ctx)
                except Exception as exc2:
                    errors.append(f"{p.name} (image page): {exc2}")
        if doc.page_count == 0:
            raise RuntimeError("None of the images could be added to the PDF.")
        pages = doc.page_count
        doc.set_metadata({"title": out.stem, "producer": APP_NAME})
        doc.save(str(out), garbage=3, deflate=True)
        ok = True
    except Cancelled:
        discard([out])
        raise
    finally:
        doc.close()
        if not ok:
            out.unlink(missing_ok=True)

    text = ""
    try:
        with fitz.open(str(out)) as check_doc:
            text = "\n".join(page.get_text("text") for page in check_doc).strip()
    except Exception:
        pass
    note = _ocr_language_note(langs, missing)
    if errors:
        note += f" ({len(errors)} failed \u2014 see log)"
    return Result(f"Created searchable PDF with {pages} page(s), "
                  f"{ocred} OCRed{note}", [out],
                  preview=visual_text(text[:3000]) if text else None,
                  clipboard=text or None, failed=failed_items)


def act_ocr_pdf_to_pdf(paths, opts, ctx: JobContext) -> Result:
    """Make scanned PDFs searchable; native text pages are copied as-is."""
    pdfs = _pdf_files(paths)
    if not pdfs:
        raise RuntimeError("No PDF files to OCR.")
    dpi = max(100, int(opts.get("dpi", 250)))
    preferred = ctx.cfg.get("ocr_langs")
    pytesseract = None
    langs: list[str] = []
    missing: list[str] = []
    lang = ""
    ocr_error = None
    outs, errors = [], []
    failed_files = 0
    ocred_pages = native_pages = 0
    try:
        for p in pdfs:
            ctx.check()
            ctx.progress(f"Building searchable PDF: {p.name}")
            out = None
            file_failed = False
            try:
                doc = _open_pdf(p)
                try:
                    out = get_unique_filename(
                        p.with_name(f"{p.stem}_searchable{SUFFIX}.pdf"))
                    dst = fitz.open()
                    try:
                        total = doc.page_count
                        if total == 0:
                            raise RuntimeError("PDF has no pages.")
                        for i in range(total):
                            ctx.check()
                            ctx.progress(f"OCR {p.name}: page {i + 1}/{total}")
                            page = doc.load_page(i)
                            native = page.get_text("text").strip()
                            if len(native) >= 15:
                                dst.insert_pdf(doc, from_page=i, to_page=i)
                                native_pages += 1
                                continue
                            if pytesseract is None and ocr_error is None:
                                try:
                                    pytesseract, langs, missing = _ocr_setup(preferred)
                                    lang = "+".join(langs)
                                except Exception as exc:
                                    ocr_error = exc
                                    file_failed = True
                                    errors.append(f"{p.name}: {exc}")
                                    log.exception("OCR setup failed for %s", p.name)
                            if ocr_error is not None:
                                file_failed = True
                                dst.insert_pdf(doc, from_page=i, to_page=i)
                                continue
                            try:
                                pix = page.get_pixmap(dpi=dpi, alpha=False)
                                try:
                                    img = Image.frombytes(
                                        "RGB", (pix.width, pix.height), pix.samples)
                                    try:
                                        data = _tesseract_pdf_bytes(
                                            img, pytesseract, lang, dpi)
                                    finally:
                                        img.close()
                                finally:
                                    del pix
                                _insert_tesseract_page(dst, data,
                                                       page.rect.width, page.rect.height)
                                ocred_pages += 1
                            except Cancelled:
                                raise
                            except Exception as exc:
                                file_failed = True
                                errors.append(f"{p.name} page {i + 1}: {exc}")
                                log.exception("OCR failed for %s page %s", p.name, i + 1)
                                dst.insert_pdf(doc, from_page=i, to_page=i)
                        dst.set_metadata({"title": out.stem, "producer": APP_NAME})
                        dst.save(str(out), garbage=3, deflate=True)
                    finally:
                        dst.close()
                finally:
                    doc.close()
                outs.append(out)
            except Cancelled:
                if out is not None:
                    out.unlink(missing_ok=True)
                raise
            except Exception as exc:
                if out is not None:
                    out.unlink(missing_ok=True)
                _record_failure(errors, p.name, exc)
                file_failed = True
            if file_failed:
                failed_files += 1
    except Cancelled:
        discard(outs)
        raise
    if not outs:
        if errors:
            raise RuntimeError("No searchable PDFs could be created:\n"
                               + "\n".join(errors[:5]))
        return Result("No searchable PDFs were created.")
    note = _ocr_language_note(langs, missing)
    if failed_files:
        note += f" ({failed_files} PDF(s) had errors \u2014 see log)"
    text = ""
    try:
        with fitz.open(str(outs[0])) as check_doc:
            text = "\n".join(page.get_text("text") for page in check_doc).strip()
    except Exception:
        pass
    return Result(
        f"Created {len(outs)} searchable PDF(s) ({native_pages} native page(s), "
        f"{ocred_pages} OCRed){note}", outs,
        preview=visual_text(text[:3000]) if text else None,
        clipboard=text or None, failed=failed_files)
def act_images_to_pdf(paths, opts, ctx: JobContext) -> Result:
    images = [Path(p) for p in paths if Path(p).is_file() and Path(p).suffix.lower() in IMG_EXTS]
    if not images:
        raise RuntimeError("No images to convert.")
    tag = "_combined" if len(images) > 1 else ""
    out = get_unique_filename(images[0].with_name(f"{images[0].stem}{tag}{SUFFIX}.pdf"))
    doc = fitz.open()
    errors: list[str] = []
    pages = 0
    ok = False
    try:
        for i, p in enumerate(images, 1):
            ctx.check()
            ctx.progress(f"Adding {i}/{len(images)}: {p.name}")
            try:
                _add_image_page(doc, p, ctx)
                pages += 1
            except Cancelled:
                raise
            except Exception as exc:
                _record_failure(errors, p.name, exc)
        if pages == 0:
            raise RuntimeError("None of the images could be added.")
        doc.set_metadata({"title": out.stem, "producer": APP_NAME})
        ctx.progress("Writing PDF\u2026")
        doc.save(str(out), garbage=3, deflate=True)
        ok = True
    except Cancelled:
        discard([out])
        raise
    finally:
        doc.close()
        if not ok:
            out.unlink(missing_ok=True)
    note = f" ({len(errors)} failed \u2014 see log)" if errors else ""
    return Result(f"Created {out.name} with {pages} page(s) "
                  f"({fmt_size(out.stat().st_size)}){note}", [out], failed=len(errors))


def act_mixed_to_pdf(paths, opts, ctx: JobContext) -> Result:
    items = [Path(p) for p in paths if Path(p).is_file()
             and (Path(p).suffix.lower() == ".pdf" or Path(p).suffix.lower() in IMG_EXTS)]
    if not items:
        raise RuntimeError("Drop images and/or PDFs to combine.")
    if all(p.suffix.lower() == ".pdf" for p in items):
        return act_merge_pdfs(items, opts, ctx)
    out = get_unique_filename(items[0].with_name(f"{items[0].stem}_combined{SUFFIX}.pdf"))
    doc = fitz.open()
    pdf_count = img_count = 0
    errors: list[str] = []
    ok = False
    try:
        for i, p in enumerate(items, 1):
            ctx.check()
            ctx.progress(f"Adding {i}/{len(items)}: {p.name}")
            try:
                if p.suffix.lower() == ".pdf":
                    src = _open_pdf(p)
                    try:
                        doc.insert_pdf(src)
                        pdf_count += 1
                    finally:
                        src.close()
                elif p.suffix.lower() in IMG_EXTS:
                    _add_image_page(doc, p, ctx)
                    img_count += 1
                else:
                    raise RuntimeError(f"Unsupported file in the mix: {p.name}")
            except Cancelled:
                raise
            except Exception as exc:
                _record_failure(errors, p.name, exc)
        if doc.page_count == 0:
            raise RuntimeError("None of the selected files could be combined.")
        doc.set_metadata({"title": out.stem, "producer": APP_NAME})
        ctx.progress("Writing combined PDF\u2026")
        doc.save(str(out), garbage=3, deflate=True)
        ok = True
    except Cancelled:
        discard([out])
        raise
    finally:
        doc.close()
        if not ok:
            out.unlink(missing_ok=True)
    note = f" ({len(errors)} failed \u2014 see log)" if errors else ""
    return Result(f"Combined {pdf_count} PDF(s) and {img_count} image(s) into {out.name} "
                  f"({fmt_size(out.stat().st_size)}){note}", [out], failed=len(errors))


def act_images_to_gif(paths, opts, ctx: JobContext) -> Result:
    images = [Path(p) for p in paths if Path(p).is_file() and Path(p).suffix.lower() in IMG_EXTS]
    if len(images) < 2:
        raise RuntimeError("Drop at least two images to build an animated GIF.")
    duration = int(opts.get("duration", 500))
    out: Path | None = None
    frames: list[Image.Image] = []
    errors: list[str] = []
    size = None
    try:
        for i, p in enumerate(images, 1):
            ctx.check()
            ctx.progress(f"Building frame {i}/{len(images)}")
            try:
                src_im = load_image(p)
                try:
                    im = flatten_rgb(src_im)
                finally:
                    src_im.close()
                try:
                    if size is None:
                        im.thumbnail((800, 800))
                        size = im.size
                    else:
                        im.thumbnail(size)
                    canvas = Image.new("RGB", size, (255, 255, 255))
                    canvas.paste(im, ((size[0] - im.width) // 2,
                                      (size[1] - im.height) // 2))
                    frames.append(canvas)
                finally:
                    im.close()
            except Cancelled:
                raise
            except Exception as exc:
                _record_failure(errors, p.name, exc)
        if len(frames) < 2:
            if errors:
                raise RuntimeError("Not enough usable images to build a GIF:\n"
                                   + "\n".join(errors[:5]))
            raise RuntimeError("Drop at least two usable images to build a GIF.")
        out = get_unique_filename(images[0].with_name(
            f"{images[0].stem}_animation{SUFFIX}.gif"))
        ctx.progress("Encoding GIF\u2026")
        try:
            frames[0].save(str(out), save_all=True, append_images=frames[1:],
                           duration=duration, loop=0, optimize=False)
        except BaseException:
            out.unlink(missing_ok=True)
            raise
    except Cancelled:
        if out is not None:
            out.unlink(missing_ok=True)
        raise
    finally:
        for frame in frames:
            frame.close()
    note = f" ({len(errors)} failed \u2014 see log)" if errors else ""
    return Result(f"Animated GIF with {len(frames)} frames "
                  f"({fmt_size(out.stat().st_size)}){note}", [out], failed=len(errors))


_SAME_FORMAT = {"PNG": {".png"}, "JPG": {".jpg", ".jpeg"}, "WEBP": {".webp"}}


def make_convert(target: str) -> Callable:
    ext = {"PNG": ".png", "JPG": ".jpg", "WEBP": ".webp"}[target]

    def run(paths, opts, ctx: JobContext) -> Result:
        images = [Path(p) for p in paths if Path(p).is_file() and Path(p).suffix.lower() in IMG_EXTS]
        if not images:
            raise RuntimeError("No images to convert.")
        outs, skipped, errors = [], 0, []
        try:
            for i, p in enumerate(images, 1):
                if p.suffix.lower() in _SAME_FORMAT[target]:
                    skipped += 1
                    continue
                ctx.check()
                ctx.progress(f"Converting {i}/{len(images)}: {p.name}")
                out = None
                im = None
                try:
                    out = derive_output(p, ext)
                    im = load_image(p)
                    if target == "JPG":
                        jpg = flatten_rgb(im)
                        try:
                            jpg.save(str(out), "JPEG", quality=92, optimize=True)
                        finally:
                            jpg.close()
                    elif target == "PNG":
                        _png_safe(im).save(str(out), "PNG", optimize=True)
                    else:
                        webp = _rgb_or_rgba(im)
                        try:
                            webp.save(str(out), "WEBP", quality=90, method=6)
                        finally:
                            if webp is not im:
                                webp.close()
                    outs.append(out)
                except Cancelled:
                    if out is not None:
                        out.unlink(missing_ok=True)
                    raise
                except Exception as exc:
                    if out is not None:
                        out.unlink(missing_ok=True)
                    _record_failure(errors, p.name, exc)
                finally:
                    if im is not None:
                        im.close()
        except Cancelled:
            discard(outs)
            raise
        if not outs:
            if errors:
                raise RuntimeError(f"No images could be converted to {target}:\n"
                                   + "\n".join(errors[:5]))
            raise RuntimeError(f"Nothing to do: everything is already {target}.")
        extra = f" ({skipped} skipped, already {target})" if skipped else ""
        if errors:
            extra += f" ({len(errors)} failed \u2014 see log)"
        return Result(f"Converted {len(outs)} image(s) to {target}{extra}", outs,
                      failed=len(errors))

    return run


def act_compress_image(paths, opts, ctx: JobContext) -> Result:
    images = [Path(p) for p in paths if Path(p).is_file() and Path(p).suffix.lower() in IMG_EXTS]
    if not images:
        raise RuntimeError("No images to compress.")
    q = int(opts.get("quality", 70))
    outs, before, after, skipped = [], 0, 0, 0
    errors: list[str] = []
    try:
        for i, src in enumerate(images, 1):
            ctx.check()
            ctx.progress(f"Compressing {i}/{len(images)}: {src.name}")
            out = None
            im = None
            try:
                out = derive_output(src, ".png")
                im = load_image(src)
                ext = src.suffix.lower()
                if ext in (".jpg", ".jpeg", ".bmp", ".tif", ".tiff"):
                    out = derive_output(src, ".jpg")
                    jpg = flatten_rgb(im)
                    try:
                        jpg.save(str(out), "JPEG", quality=q, optimize=True, progressive=True)
                    finally:
                        jpg.close()
                elif ext == ".webp":
                    out = derive_output(src, ".webp")
                    webp = _rgb_or_rgba(im)
                    try:
                        webp.save(str(out), "WEBP", quality=q, method=6)
                    finally:
                        if webp is not im:
                            webp.close()
                else:  # PNG and other lossless formats: "quality" means palette size
                    out = derive_output(src, ".png")
                    if q >= 100:
                        _png_safe(im).save(str(out), "PNG", optimize=True)
                    else:
                        colors = max(16, round(256 * q / 100))
                        quant = im.convert("RGBA").quantize(colors=colors, method=_QUANTIZE)
                        try:
                            quant.save(str(out), "PNG", optimize=True)
                        finally:
                            quant.close()
                old, new = src.stat().st_size, out.stat().st_size
                if new >= old:
                    out.unlink(missing_ok=True)
                    skipped += 1
                    continue
                outs.append(out)
                before += old
                after += new
            except Cancelled:
                if out is not None:
                    out.unlink(missing_ok=True)
                raise
            except Exception as exc:
                if out is not None:
                    out.unlink(missing_ok=True)
                _record_failure(errors, src.name, exc)
            finally:
                if im is not None:
                    im.close()
    except Cancelled:
        discard(outs)
        raise
    if not outs:
        if errors:
            raise RuntimeError("No images could be compressed:\n" + "\n".join(errors[:5]))
        return Result("No smaller files were produced; these images are already well optimized.")
    saved = 100 - after * 100 / before
    extra = f" ({skipped} left unchanged)" if skipped else ""
    if errors:
        extra += f" ({len(errors)} failed \u2014 see log)"
    return Result(f"Compressed {len(outs)} image(s): {fmt_size(before)} -> {fmt_size(after)} "
                  f"(-{saved:.0f}%){extra}", outs, failed=len(errors))


# --------------------------------------------------------------------------- #
#  Actions: Word documents
# --------------------------------------------------------------------------- #
_PDF_FONT: str | None = None


def _pdf_font() -> str:
    """Use Arial from Windows (full Arabic coverage) when available."""
    global _PDF_FONT
    if _PDF_FONT:
        return _PDF_FONT
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    fonts = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts"
    files = {"DF": "arial.ttf", "DF-B": "arialbd.ttf", "DF-I": "ariali.ttf", "DF-BI": "arialbi.ttf"}
    try:
        if all((fonts / f).exists() for f in files.values()):
            for name, fname in files.items():
                pdfmetrics.registerFont(TTFont(name, str(fonts / fname)))
            pdfmetrics.registerFontFamily("DF", normal="DF", bold="DF-B",
                                          italic="DF-I", boldItalic="DF-BI")
            _PDF_FONT = "DF"
    except Exception:
        log.exception("Could not register Arial")
    _PDF_FONT = _PDF_FONT or "Helvetica"
    return _PDF_FONT


def _render_docx_to_pdf(src: Path, dst: Path) -> None:
    """Built-in fallback renderer (no LibreOffice / no MS Office required)."""
    require("docx", "python-docx")
    require("reportlab")
    from docx import Document
    from docx.table import Table as DocxTable
    from docx.text.paragraph import Paragraph as DocxParagraph
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from reportlab.platypus import (Image as RLImage, PageBreak, Paragraph,
                                    SimpleDocTemplate, Spacer, Table, TableStyle)
    from xml.sax.saxutils import escape

    font = _pdf_font()
    bold_font = "DF-B" if font == "DF" else "Helvetica-Bold"
    doc = Document(str(src))
    try:
        sec = doc.sections[0]
        pw, ph = sec.page_width.pt, sec.page_height.pt
        lm, rm, tm, bm = (sec.left_margin.pt, sec.right_margin.pt,
                          sec.top_margin.pt, sec.bottom_margin.pt)
    except Exception:
        pw, ph = A4
        lm = rm = tm = bm = 56.0
    avail = pw - lm - rm
    aligns = {0: TA_LEFT, 1: TA_CENTER, 2: TA_RIGHT, 3: TA_JUSTIFY}
    heading_sizes = {"title": 22, "subtitle": 14, "heading 1": 18, "heading 2": 15,
                     "heading 3": 13, "heading 4": 12}
    counter = {"n": 0}

    def mk_style(size, bold, align, indent=0):
        return ParagraphStyle("df", fontName=bold_font if bold else font, fontSize=size,
                              leading=size * 1.3, alignment=align, leftIndent=indent)

    def inline_markup(par) -> str:
        parts = []
        try:
            items = list(par.iter_inner_content())
        except Exception:
            items = list(par.runs)
        for it in items:
            runs = getattr(it, "runs", None)
            if runs is None:
                runs = [it]
            for r in runs:
                if not getattr(r, "text", ""):
                    continue
                s = escape(r.text).replace("\t", "    ").replace("\n", "<br/>")
                if r.bold:
                    s = f"<b>{s}</b>"
                if r.italic:
                    s = f"<i>{s}</i>"
                if r.underline:
                    s = f"<u>{s}</u>"
                parts.append(s)
        return "".join(parts)

    def rtl_flowables(text, size, bold, align, width):
        try:
            import arabic_reshaper
            get_display = _get_bidi()
        except Exception:
            style = mk_style(size, bold, align)
            return [Paragraph(escape(seg), style) for seg in text.split("\n")]
        fname = bold_font if bold else font
        style = mk_style(size, bold, align)
        out = []
        for segment in text.replace("\t", " ").split("\n"):
            shaped = arabic_reshaper.reshape(segment)
            lines, cur = [], []
            for word in shaped.split(" "):
                if cur and stringWidth(" ".join(cur + [word]), fname, size) > width * 0.97:
                    lines.append(" ".join(cur))
                    cur = [word]
                else:
                    cur.append(word)
            lines.append(" ".join(cur))
            out += [Paragraph(escape(get_display(line)), style) for line in lines]
        return out

    def para_images(par, width):
        out = []
        try:
            for run in par.runs:
                for rid in run._element.xpath(".//a:blip/@r:embed"):
                    blob = doc.part.related_parts[rid].blob
                    with Image.open(io.BytesIO(blob)) as im:
                        im.load()
                        w_pt, h_pt = im.width * 0.75, im.height * 0.75
                        if im.format in ("JPEG", "PNG"):
                            data = blob
                        else:
                            buf = io.BytesIO()
                            im.convert("RGBA").save(buf, "PNG")
                            data = buf.getvalue()
                    k = min(1.0, width / w_pt, (ph - tm - bm) * 0.9 / h_pt)
                    out.append(RLImage(io.BytesIO(data), width=w_pt * k, height=h_pt * k))
        except Exception:
            log.exception("Skipping an image that could not be rendered")
        return out

    def convert_paragraph(par, width, size=11.0, in_table=False):
        flow = []
        text = par.text or ""
        style_name = (par.style.name if par.style is not None else "").lower()
        if not in_table and par.paragraph_format.page_break_before:
            flow.append(PageBreak())
        hsize = heading_sizes.get(style_name)
        bold = hsize is not None
        size = hsize or size
        pPr = par._p.pPr
        is_list = style_name.startswith("list") or (pPr is not None and pPr.numPr is not None)
        prefix = ""
        if is_list:
            if "number" in style_name:
                counter["n"] += 1
                prefix = f"{counter['n']}. "
            else:
                prefix = "\u2022 "
        else:
            counter["n"] = 0
        raw_align = par.alignment
        align = aligns.get(getattr(raw_align, "value", raw_align)) if raw_align is not None else None

        images_in_par = para_images(par, width)
        if text.strip():
            if ARABIC_RE.search(text):
                flow += rtl_flowables(text, size, bold,
                                      TA_CENTER if align == TA_CENTER else TA_RIGHT, width)
            else:
                markup = inline_markup(par) or escape(text)
                indent = 14 if is_list else 0
                flow.append(Paragraph(prefix + markup,
                                      mk_style(size, bold,
                                               align if align is not None else TA_LEFT, indent)))
        elif not images_in_par and not in_table:
            flow.append(Spacer(1, size * 1.2))
        flow += images_in_par
        flow.append(Spacer(1, 2 if in_table else (8 if bold else 4)))
        if not in_table and par._p.xpath('.//w:br[@w:type="page"]'):
            flow.append(PageBreak())
        return flow

    def convert_table(tbl):
        ncols = max((len(r.cells) for r in tbl.rows), default=0)
        if not ncols:
            return []
        cw = avail / ncols
        data = []
        for r in tbl.rows:
            row, seen = [], set()
            for c in r.cells:
                if id(c._tc) in seen:  # merged cell repeats
                    row.append("")
                    continue
                seen.add(id(c._tc))
                cell_flow = []
                for p in c.paragraphs:
                    cell_flow += convert_paragraph(p, cw - 10, size=9.5, in_table=True)
                row.append(cell_flow or "")
            row += [""] * (ncols - len(row))
            data.append(row)
        t = Table(data, colWidths=[cw] * ncols)
        t.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#999999")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        return [t, Spacer(1, 8)]

    story = []
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            story += convert_paragraph(DocxParagraph(child, doc._body), avail)
        elif tag == "tbl":
            try:
                story += convert_table(DocxTable(child, doc._body))
            except Exception:
                log.exception("Table could not be rendered")
                story.append(Paragraph("[table could not be rendered]",
                                       mk_style(9, False, TA_LEFT)))
    if not story:
        story.append(Spacer(1, 1))
    SimpleDocTemplate(str(dst), pagesize=(pw, ph), leftMargin=lm, rightMargin=rm,
                      topMargin=tm, bottomMargin=bm, title=src.stem).build(story)


def _new_temp_dir(prefix: str) -> Path:
    """Create a private-ish work folder using inherited Windows ACLs.

    tempfile.mkdtemp() applies a restrictive mode that can confuse Windows
    sandboxes; a plain mkdir under %TEMP% is friendlier and just as unique.
    """
    base = Path(tempfile.gettempdir())
    for _ in range(20):
        candidate = base / f"{prefix}{os.getpid()}_{time.time_ns()}"
        try:
            candidate.mkdir()
            return candidate
        except FileExistsError:
            continue
    raise RuntimeError("Could not create a temporary folder.")


def _soffice_convert(soffice: str, src: Path, out: Path, ctx: JobContext) -> bool:
    tmp = None
    try:
        tmp = _new_temp_dir("dropflow_lo_")
        outdir = tmp / "out"
        outdir.mkdir()
        profile = tmp / "profile"
        uri = profile.resolve().as_uri()
        ctx.progress(f"Converting with LibreOffice: {src.name}")
        cmd = [
            soffice, "--headless", "--nologo", "--nodefault", "--nofirststartwizard",
            "--nolockcheck", f"-env:UserInstallation={uri}",
            "--convert-to", "pdf", "--outdir", str(outdir), str(src),
        ]
        res = ctx.run(cmd, timeout=300, label="LibreOffice")
        pdfs = sorted(outdir.glob("*.pdf"))
        if not pdfs:
            log.warning("LibreOffice produced no PDF for %s (rc=%s): %s",
                        src, res.returncode, _tail(res.stderr, 800))
            return False
        shutil.move(str(pdfs[0]), str(out))
        return True
    except Cancelled:
        raise
    except Exception:
        log.exception("LibreOffice conversion failed for %s", src)
        return False
    finally:
        if tmp is not None:
            shutil.rmtree(tmp, ignore_errors=True)


def act_docx_to_pdf(paths, opts, ctx: JobContext) -> Result:
    docs = [Path(p) for p in paths if Path(p).is_file() and Path(p).suffix.lower() == ".docx"]
    if not docs:
        raise RuntimeError("No Word documents to convert.")
    soffice = _find_soffice()
    outs, used_lo, used_builtin = [], 0, 0
    errors: list[str] = []
    try:
        for i, p in enumerate(docs, 1):
            ctx.check()
            ctx.progress(f"Converting {i}/{len(docs)}: {p.name}")
            out = derive_output(p, ".pdf")
            try:
                ok = False
                if soffice:
                    try:
                        ok = _soffice_convert(soffice, p, out, ctx)
                    except Cancelled:
                        out.unlink(missing_ok=True)
                        raise
                if not ok:
                    ctx.check()
                    out.unlink(missing_ok=True)
                    _render_docx_to_pdf(p, out)
                    used_builtin += 1
                else:
                    used_lo += 1
                outs.append(out)
            except Cancelled:
                out.unlink(missing_ok=True)
                raise
            except Exception as exc:
                out.unlink(missing_ok=True)
                _record_failure(errors, p.name, exc)
    except Cancelled:
        discard(outs)
        raise
    if not outs:
        if errors:
            raise RuntimeError("No Word documents could be converted:\n"
                               + "\n".join(errors[:5]))
        return Result("No Word documents were converted.")
    details = []
    if used_lo:
        details.append(f"{used_lo} with LibreOffice")
    if used_builtin:
        details.append(f"{used_builtin} with the built-in renderer")
    msg = f"Converted {len(outs)} document(s) to PDF ({', '.join(details)})."
    if errors:
        msg += f" {len(errors)} failed \u2014 see log."
    if used_builtin:
        if soffice:
            msg += " LibreOffice could not convert at least one file; the built-in renderer was used."
        else:
            msg += " Install LibreOffice for pixel-perfect layouts."
    return Result(msg, outs, failed=len(errors))


def act_docx_to_txt(paths, opts, ctx: JobContext) -> Result:
    require("docx", "python-docx")
    from docx import Document
    from docx.table import Table as DocxTable
    from docx.text.paragraph import Paragraph as DocxParagraph
    docs = [Path(p) for p in paths if Path(p).is_file() and Path(p).suffix.lower() == ".docx"]
    if not docs:
        raise RuntimeError("No Word documents to read.")
    outs = []
    errors: list[str] = []
    try:
        for i, p in enumerate(docs, 1):
            ctx.check()
            ctx.progress(f"Extracting {i}/{len(docs)}: {p.name}")
            try:
                doc = Document(str(p))
                lines = []
                for child in doc.element.body.iterchildren():
                    tag = child.tag.rsplit("}", 1)[-1]
                    if tag == "p":
                        lines.append(DocxParagraph(child, doc._body).text)
                    elif tag == "tbl":
                        for row in DocxTable(child, doc._body).rows:
                            seen, cells = set(), []
                            for c in row.cells:
                                if id(c._tc) in seen:
                                    continue
                                seen.add(id(c._tc))
                                cells.append(c.text.replace("\n", " "))
                            lines.append("\t".join(cells))
                out = derive_output(p, ".txt")
                out.write_text("\n".join(lines), encoding="utf-8-sig")
                outs.append(out)
            except Cancelled:
                raise
            except Exception as exc:
                _record_failure(errors, p.name, exc)
    except Cancelled:
        discard(outs)
        raise
    if not outs:
        if errors:
            raise RuntimeError("No documents could be read:\n" + "\n".join(errors[:5]))
        return Result("No text was extracted.")
    note = f" ({len(errors)} failed \u2014 see log)" if errors else ""
    return Result(f"Extracted text from {len(outs)} document(s){note}", outs,
                  failed=len(errors))


def _require_office():
    if office is None:
        detail = f"\n\nImport error: {_OFFICE_IMPORT_ERROR}" if _OFFICE_IMPORT_ERROR else ""
        raise RuntimeError(
            "office_converter.py was not found next to DropFlow.pyw.\n\n"
            "Place office_converter.py in the same folder as DropFlow.pyw and "
            f"restart DropFlow.{detail}")


def _new_office_converter(ctx: JobContext):
    _require_office()
    return office.OfficeConverter(
        soffice_path=_find_soffice(),
        dpi=150,
        ocr_lang=str(ctx.cfg.get("ocr_langs", "eng") or "eng"),
    )


def _office_progress(ctx: JobContext, text: str) -> None:
    try:
        ctx.progress(text)
    except Cancelled:
        pass


def act_convert_office(paths, opts, ctx: JobContext) -> Result:
    _require_office()
    docs = [Path(p) for p in paths
            if Path(p).is_file() and Path(p).suffix.lower() in OFFICE_INPUT_EXTS]
    if not docs:
        raise RuntimeError("No Office, text, or data files to convert.")
    target = office.OfficeConverter.normalize_target(str(opts.get("target", "pdf")))
    converter = _new_office_converter(ctx)
    outs = []
    errors: list[str] = []
    methods: dict[str, int] = {}
    try:
        for i, p in enumerate(docs, 1):
            ctx.check()
            ctx.progress(f"Converting {i}/{len(docs)}: {p.name} \u2192 {target.upper()}\u2026")
            out = None
            try:
                if target in office.DIR_TARGETS:
                    folder = "slides" if target in ("png", "jpg") else "csv"
                    out = get_unique_filename(
                        p.with_name(f"{p.stem}_{folder}"), is_dir=True)
                    out.mkdir(parents=True)
                else:
                    out = derive_output(p, f".{target}")
                callbacks = office.OfficeCallbacks(
                    on_start=lambda src, t, _p=p: _office_progress(
                        ctx, f"Converting {_p.name} \u2192 {t.upper()}\u2026"),
                    on_progress=lambda pct, msg, _p=p: _office_progress(
                        ctx, f"{_p.name}: {msg}"),
                )
                res = converter.convert(p, target, out, callbacks=callbacks,
                                        cancel_event=ctx.cancel_event)
                outs.extend(res.outputs)
                methods[res.method] = methods.get(res.method, 0) + 1
            except office.OfficeCancelled as exc:
                if out is not None:
                    discard([out])
                raise Cancelled() from exc
            except Exception as exc:
                if ctx.cancelled:
                    if out is not None:
                        discard([out])
                    raise Cancelled() from exc
                if out is not None:
                    discard([out])
                _record_failure(errors, p.name, exc)
    except Cancelled:
        discard(outs)
        raise
    if not outs:
        if errors:
            raise RuntimeError("No documents could be converted:\n" +
                               "\n".join(errors[:5]))
        return Result("No documents were converted.")
    note = f" ({len(errors)} failed \u2014 see log)" if errors else ""
    method_note = ""
    if methods:
        method_note = " (" + ", ".join(f"{v} {k}" for k, v in methods.items()) + ")"
    return Result(f"Converted {len(docs) - len(errors)} document(s) to "
                  f"{target.upper()}{method_note}{note}", outs, failed=len(errors))


def act_images_pdf_to_pptx(paths, opts, ctx: JobContext) -> Result:
    _require_office()
    sources = [Path(p) for p in paths
               if Path(p).is_file()
               and Path(p).suffix.lower() in (IMG_EXTS | {".pdf"})]
    if not sources:
        raise RuntimeError("No images or PDFs to build a presentation from.")
    converter = _new_office_converter(ctx)
    out = derive_output(sources[0], ".pptx")
    callbacks = office.OfficeCallbacks(
        on_start=lambda src, t: _office_progress(
            ctx, f"Building presentation from {len(sources)} file(s)\u2026"),
        on_progress=lambda pct, msg: _office_progress(ctx, msg),
    )
    try:
        res = converter.to_pptx_many(
            sources, out, callbacks=callbacks, cancel_event=ctx.cancel_event)
    except office.OfficeCancelled as exc:
        discard([out])
        raise Cancelled() from exc
    except Exception:
        discard([out])
        if ctx.cancelled:
            raise Cancelled()
        raise
    return Result(f"Created a PowerPoint presentation from {len(sources)} "
                  f"file(s) \u2014 {len(res.outputs)} deck", res.outputs)


def act_data_to_xlsx(paths, opts, ctx: JobContext) -> Result:
    _require_office()
    sources = [Path(p) for p in paths
               if Path(p).is_file()
               and Path(p).suffix.lower() in (office.DATA_EXTS | office.PDF_EXTS)]
    if not sources:
        raise RuntimeError("No CSV, JSON, or PDF data to put into Excel.")
    converter = _new_office_converter(ctx)
    outs = []
    errors: list[str] = []
    try:
        for i, p in enumerate(sources, 1):
            ctx.check()
            ctx.progress(f"Building workbook {i}/{len(sources)}: {p.name}\u2026")
            out = derive_output(p, ".xlsx")
            try:
                converter.convert(p, "xlsx", out, cancel_event=ctx.cancel_event)
                outs.append(out)
            except office.OfficeCancelled as exc:
                discard([out])
                raise Cancelled() from exc
            except Exception as exc:
                if ctx.cancelled:
                    discard([out])
                    raise Cancelled() from exc
                discard([out])
                _record_failure(errors, p.name, exc)
    except Cancelled:
        discard(outs)
        raise
    if not outs:
        if errors:
            raise RuntimeError("No workbooks could be created:\n" +
                               "\n".join(errors[:5]))
        return Result("No workbooks were created.")
    note = f" ({len(errors)} failed \u2014 see log)" if errors else ""
    return Result(f"Created {len(outs)} Excel workbook(s){note}", outs,
                  failed=len(errors))


# --------------------------------------------------------------------------- #
#  Actions: video and audio (FFmpeg)
# --------------------------------------------------------------------------- #
def _require_ffmpeg() -> str:
    exe = _find_ffmpeg()
    if not exe:
        raise RuntimeError(
            "FFmpeg was not found.\n\n"
            "Install it with:  winget install Gyan.FFmpeg\n"
            "or download it from https://ffmpeg.org/ and put ffmpeg.exe on PATH "
            "(or set the FFMPEG_PATH environment variable), then restart DropFlow.")
    return exe


def _video_files(paths) -> list[Path]:
    return [Path(p) for p in paths if Path(p).is_file() and Path(p).suffix.lower() in VIDEO_EXTS]


def _ffmpeg_has_encoder(name: str) -> bool:
    return name in _ffmpeg_encoders()


def _ffprobe_video_size(path: Path) -> tuple[int | None, int | None]:
    exe = _find_ffprobe()
    if not exe:
        return None, None
    try:
        res = _run_hidden([
            exe, "-v", "error", "-select_streams", "v:0",
            "-show_entries", "stream=width,height",
            "-of", "default=noprint_wrappers=1", str(path),
        ], timeout=20)
        width = height = None
        for line in (res.stdout or "").splitlines():
            if line.startswith("width="):
                value = line.split("=", 1)[1].strip()
                if value.isdigit():
                    width = int(value)
            elif line.startswith("height="):
                value = line.split("=", 1)[1].strip()
                if value.isdigit():
                    height = int(value)
        return width, height
    except Exception:
        return None, None


_RESOLUTION_TARGETS = {"2160p": 2160, "4k": 2160, "1440p": 1440,
                       "1080p": 1080, "720p": 720, "480p": 480}


def _video_filter(width: int | None, height: int | None, resolution: str) -> str:
    target = _RESOLUTION_TARGETS.get(str(resolution or "").strip().lower(), 0)
    even = "scale=trunc(iw/2)*2:trunc(ih/2)*2"
    if not target or not width or not height:
        return even
    if width >= height:
        if height <= target:
            return even
        return f"scale=-2:{target}"
    if width <= target:
        return even
    return f"scale={target}:-2"


def _parse_target_mb(value) -> float:
    if value is None:
        return 0.0
    text = str(value).strip().lower().replace("mb", "").replace(",", ".").strip()
    if not text:
        return 0.0
    try:
        number = float(text)
    except ValueError as exc:
        raise RuntimeError("Target size must be a number of megabytes, e.g. 50.") from exc
    if number < 0:
        raise RuntimeError("Target size cannot be negative.")
    return number


def _video_audio_args(src_codec: str | None, src_bitrate: int | None,
                      audio_codec: str,
                      reencode_kbps: int | None = None) -> tuple[list[str], float]:
    if reencode_kbps is not None:
        if audio_codec == "copy":
            return ["-c:a", "copy"], (src_bitrate or 128_000) / 1000.0
        args = []
        if audio_codec == "aac":
            args += ["-strict", "-2"]
        args += ["-c:a", audio_codec, "-b:a", f"{max(32, reencode_kbps)}k"]
        return args, float(max(32, reencode_kbps))
    if audio_codec == "copy":
        return ["-c:a", "copy"], (src_bitrate or 128_000) / 1000.0
    if src_codec in ("aac", "mp3") and (src_bitrate is None or src_bitrate <= 140_000):
        return ["-c:a", "copy"], (src_bitrate or 128_000) / 1000.0
    args = []
    if audio_codec == "aac":
        args += ["-strict", "-2"]
    args += ["-c:a", audio_codec, "-b:a", "128k"]
    return args, 128.0


def _cpu_video_args(crf: int, video_kbps: int | None = None,
                    pass_no: int | None = None, passlog: str | None = None) -> list[str]:
    if video_kbps:
        args = ["-c:v", "libx264", "-preset", "medium",
                "-b:v", f"{video_kbps}k", "-maxrate", f"{int(video_kbps * 1.5)}k",
                "-bufsize", f"{video_kbps * 2}k", "-pix_fmt", "yuv420p"]
        if pass_no:
            args += ["-pass", str(pass_no)]
            if passlog:
                args += ["-passlogfile", passlog]
        return args
    return ["-c:v", "libx264", "-crf", str(crf), "-preset", "medium",
            "-pix_fmt", "yuv420p"]


def _gpu_video_args(crf: int, video_kbps: int | None = None) -> list[str]:
    args = ["-c:v", "h264_nvenc", "-preset", "p5"]
    if video_kbps:
        args += ["-b:v", f"{video_kbps}k", "-maxrate", f"{int(video_kbps * 1.5)}k",
                 "-bufsize", f"{video_kbps * 2}k"]
    else:
        args += ["-cq", str(crf), "-b:v", "0"]
    return args + ["-pix_fmt", "yuv420p"]


def act_compress_video(paths, opts, ctx: JobContext) -> Result:
    videos = _video_files(paths)
    if not videos:
        raise RuntimeError("No video files to compress.")
    ffmpeg = _require_ffmpeg()
    crf = max(14, min(35, int(opts.get("quality", 23))))
    target_mb = _parse_target_mb(opts.get("target_mb"))
    resolution = str(opts.get("resolution", "Keep original"))
    audio_codec, _ = _ffmpeg_audio_choice()
    gpu_available = _ffmpeg_has_encoder("h264_nvenc")
    want_gpu = bool(opts.get("gpu")) and gpu_available
    notes: list[str] = []
    if opts.get("gpu") and not gpu_available:
        notes.append("GPU encoding is not available in this FFmpeg build; CPU was used.")
    outs, before, after, skipped = [], 0, 0, 0
    errors: list[str] = []
    try:
        for i, p in enumerate(videos, 1):
            ctx.check()
            ctx.progress(f"Compressing {i}/{len(videos)}: {p.name}")
            out = None
            pass_dir = None
            try:
                out = derive_output(p, ".mp4")
                duration = _ffprobe_duration(p)
                width, height = _ffprobe_video_size(p)
                src_codec, src_bitrate = _ffprobe_audio_info(p)
                vf = _video_filter(width, height, resolution)
                video_kbps = None
                if target_mb > 0:
                    if not duration:
                        raise RuntimeError(
                            "Target size needs ffprobe duration information, which was "
                            "not available for this file.")
                    total_kbps = (target_mb * 1024 * 1024 * 8) / duration / 1000.0
                    if src_codec is None:
                        audio_args, audio_kbps = [], 0.0
                    else:
                        audio_budget = max(48, min(128, int(total_kbps * 0.20)))
                        audio_args, audio_kbps = _video_audio_args(
                            src_codec, src_bitrate, audio_codec,
                            reencode_kbps=audio_budget)
                    video_kbps = max(50, int(total_kbps - audio_kbps - 8))
                else:
                    audio_args, audio_kbps = _video_audio_args(
                        src_codec, src_bitrate, audio_codec)
                passlog = None
                if target_mb > 0:
                    pass_dir = _new_temp_dir("dropflow_pass_")
                    passlog = str(pass_dir / "pass")

                def make_cmd(use_gpu: bool, pass_no: int | None = None) -> list[str]:
                    cmd = [ffmpeg, "-loglevel", "error", "-y", "-i", str(p),
                           "-map", "0:v:0", "-map", "0:a:0?", "-vf", vf]
                    if use_gpu:
                        cmd += _gpu_video_args(crf, video_kbps)
                    else:
                        cmd += _cpu_video_args(crf, video_kbps, pass_no, passlog)
                    if pass_no == 1:
                        cmd += ["-an", "-f", "null", os.devnull,
                                "-progress", "pipe:1", "-nostats"]
                    else:
                        cmd += audio_args
                        cmd += ["-movflags", "+faststart",
                                "-progress", "pipe:1", "-nostats", str(out)]
                    return cmd

                def check_result(res, phase: str, expect_output: bool = True) -> None:
                    if res.returncode != 0 or (expect_output and (
                            not out.exists() or out.stat().st_size == 0)):
                        detail = _tail((res.stderr or "") + " " + (res.stdout or ""))
                        raise RuntimeError(f"FFmpeg {phase} failed. {detail}")

                def encode(use_gpu: bool) -> None:
                    if target_mb > 0 and not use_gpu:
                        ctx.progress(f"Compressing {p.name}: pass 1/2\u2026")
                        check_result(ctx.run(make_cmd(False, 1), duration=duration,
                                             label="Compressing 1/2"), "pass 1",
                                     expect_output=False)
                        ctx.progress(f"Compressing {p.name}: pass 2/2\u2026")
                        check_result(ctx.run(make_cmd(False, 2), duration=duration,
                                             label="Compressing 2/2"), "pass 2")
                    else:
                        label = "Compressing (GPU)" if use_gpu else "Compressing"
                        check_result(ctx.run(make_cmd(use_gpu), duration=duration,
                                             label=label), "encode")

                if want_gpu:
                    try:
                        encode(True)
                    except Cancelled:
                        raise
                    except Exception as gpu_exc:
                        log.warning("GPU encode failed for %s: %s", p.name, gpu_exc)
                        out.unlink(missing_ok=True)
                        encode(False)
                        notes.append(f"{p.name}: GPU failed, CPU used")
                else:
                    encode(False)

                old, new = p.stat().st_size, out.stat().st_size
                if new >= old:
                    out.unlink(missing_ok=True)
                    skipped += 1
                    continue
                outs.append(out)
                before += old
                after += new
                if target_mb > 0 and new > target_mb * 1024 * 1024 * 1.05:
                    notes.append(f"{p.name}: target {target_mb:g} MB, got "
                                 f"{new / (1024 * 1024):.1f} MB")
            except Cancelled:
                if out is not None:
                    out.unlink(missing_ok=True)
                raise
            except Exception as exc:
                if out is not None:
                    out.unlink(missing_ok=True)
                _record_failure(errors, p.name, exc)
            finally:
                if pass_dir is not None:
                    discard([pass_dir])
    except Cancelled:
        discard(outs)
        raise
    if not outs:
        if errors:
            raise RuntimeError("No videos could be compressed:\n" + "\n".join(errors[:5]))
        return Result("No smaller videos were produced; these files are already compressed.")
    saved = 100 - after * 100 / before
    extra = f" ({skipped} left unchanged)" if skipped else ""
    if errors:
        extra += f" ({len(errors)} failed \u2014 see log)"
    note = ""
    if notes:
        note = " " + " ".join(notes[:3])
    return Result(f"Compressed {len(outs)} video(s): {fmt_size(before)} -> {fmt_size(after)} "
                  f"(-{saved:.0f}%){extra}{note}", outs, failed=len(errors))


def act_extract_audio(paths, opts, ctx: JobContext) -> Result:
    videos = _video_files(paths)
    if not videos:
        raise RuntimeError("No video files to extract audio from.")
    ffmpeg = _require_ffmpeg()
    bitrate = max(64, min(320, int(opts.get("quality", 192))))
    codec, ext = _ffmpeg_audio_choice(prefer_mp3=True)
    if codec == "copy":
        raise RuntimeError("This FFmpeg build has no MP3 or AAC encoder, so audio cannot be "
                           "extracted. Install a full FFmpeg build and restart DropFlow.")
    kind = "MP3" if ext == ".mp3" else "M4A"
    outs = []
    errors: list[str] = []
    try:
        for i, p in enumerate(videos, 1):
            ctx.check()
            ctx.progress(f"Extracting audio {i}/{len(videos)}: {p.name}")
            out = None
            try:
                out = derive_output(p, ext)
                duration = _ffprobe_duration(p)
                cmd = [
                    ffmpeg, "-loglevel", "error", "-y", "-i", str(p),
                    "-map", "0:a:0", "-vn",
                ]
                if codec == "aac":
                    cmd += ["-strict", "-2"]
                cmd += ["-c:a", codec, "-b:a", f"{bitrate}k",
                        "-progress", "pipe:1", "-nostats", str(out)]
                res = ctx.run(cmd, duration=duration, label="Extracting audio")
                if res.returncode != 0 or not out.exists() or out.stat().st_size == 0:
                    out.unlink(missing_ok=True)
                    detail = _tail((res.stderr or "") + " " + (res.stdout or ""))
                    if "matches no streams" in detail or "does not contain any stream" in detail:
                        raise RuntimeError(f"'{p.name}' has no audio track to extract.")
                    raise RuntimeError(f"No audio track could be extracted from '{p.name}'. {detail}")
                outs.append(out)
            except Cancelled:
                if out is not None:
                    out.unlink(missing_ok=True)
                raise
            except Exception as exc:
                if out is not None:
                    out.unlink(missing_ok=True)
                _record_failure(errors, p.name, exc)
    except Cancelled:
        discard(outs)
        raise
    if not outs:
        if errors:
            raise RuntimeError("No audio could be extracted:\n" + "\n".join(errors[:5]))
        return Result("No audio was extracted.")
    total = sum(o.stat().st_size for o in outs)
    note = f" ({len(errors)} failed \u2014 see log)" if errors else ""
    return Result(f"Extracted {kind} audio from {len(outs)} video(s) "
                  f"({fmt_size(total)}){note}", outs, failed=len(errors))


def act_convert_media(paths, opts, ctx: JobContext) -> Result:
    """Universal media conversion through media_converter.py."""
    if media is None:
        detail = f"\n\nImport error: {_MEDIA_IMPORT_ERROR}" if _MEDIA_IMPORT_ERROR else ""
        raise RuntimeError(
            "media_converter.py was not found next to DropFlow.pyw.\n\n"
            "Place media_converter.py in the same folder as DropFlow.pyw and "
            f"restart DropFlow.{detail}")
    videos = _video_files(paths)
    if not videos:
        raise RuntimeError("No video files to convert.")
    ffmpeg = _require_ffmpeg()
    ffprobe = _find_ffprobe()
    try:
        converter = media.MediaConverter(ffmpeg_path=ffmpeg, ffprobe_path=ffprobe)
    except Exception as exc:
        raise RuntimeError(str(exc)) from exc

    target = media.MediaConverter.normalize_format(str(opts.get("format", "mp4")))
    outs = []
    errors: list[str] = []
    methods = {"remux": 0, "reencode": 0, "audio": 0}

    def media_progress(text: str) -> None:
        try:
            ctx.progress(text)
        except Cancelled:
            pass
    try:
        for i, p in enumerate(videos, 1):
            ctx.check()
            ctx.progress(f"Converting {i}/{len(videos)}: {p.name} \u2192 {target.upper()}\u2026")
            out = None
            try:
                out = derive_output(p, f".{target}")
                callbacks = media.ConversionCallbacks(
                    on_start=lambda src, fmt, _p=p: media_progress(
                        f"Converting {_p.name} \u2192 {fmt.upper()}\u2026"),
                    on_progress=lambda pct, msg, _p=p: media_progress(
                        f"{_p.name}: {msg}"),
                )
                res = converter.convert(
                    p, target, out,
                    callbacks=callbacks,
                    cancel_event=ctx.cancel_event,
                )
                outs.append(res.output)
                methods[res.method] = methods.get(res.method, 0) + 1
                ctx.progress(f"Converted {p.name} \u2192 {target.upper()} "
                             f"({res.method})")
            except media.MediaCancelled as exc:
                if out is not None:
                    out.unlink(missing_ok=True)
                raise Cancelled() from exc
            except Exception as exc:
                if ctx.cancelled:
                    if out is not None:
                        out.unlink(missing_ok=True)
                    raise Cancelled() from exc
                if out is not None:
                    out.unlink(missing_ok=True)
                _record_failure(errors, p.name, exc)
    except Cancelled:
        discard(outs)
        raise
    if not outs:
        if errors:
            raise RuntimeError("No videos could be converted:\n" + "\n".join(errors[:5]))
        return Result("No videos were converted.")
    note = f" ({len(errors)} failed \u2014 see log)" if errors else ""
    detail = []
    if methods.get("remux"):
        detail.append(f"{methods['remux']} stream copy")
    if methods.get("reencode"):
        detail.append(f"{methods['reencode']} re-encoded")
    if methods.get("audio"):
        detail.append(f"{methods['audio']} audio")
    suffix = f" ({', '.join(detail)})" if detail else ""
    return Result(f"Converted {len(outs)} video(s) to {target.upper()}{suffix}{note}",
                  outs, failed=len(errors))


# --------------------------------------------------------------------------- #
#  Actions: archives (ZIP + RAR)
# --------------------------------------------------------------------------- #
_RAR_TOOL: tuple[str, str] | None | bool = False

_RAR_MISSING_MESSAGE = (
    "RAR extraction needs an extraction backend.\n\n"
    "Install one of these, then restart DropFlow:\n"
    "  \u2022 7-Zip: https://www.7-zip.org/\n"
    "  \u2022 WinRAR's UnRAR.exe: https://www.rarlab.com/rar_add.htm\n"
    "  \u2022 pip install rarfile  (plus UnRAR or 7-Zip)\n\n"
    "DropFlow looks for unrar.exe, 7z.exe and bsdtar on PATH, plus the "
    "standard WinRAR and 7-Zip install folders."
)


def _find_rar_tool() -> tuple[str, str] | None:
    """Return (kind, path) for UnRAR, 7-Zip or bsdtar, in that order."""
    global _RAR_TOOL
    if _RAR_TOOL is not False:
        return _RAR_TOOL
    cands: list[tuple[str, str]] = []
    seen: set[str] = set()

    def add(kind: str, value) -> None:
        if not value:
            return
        try:
            path = Path(value)
            if path.is_file():
                key = os.path.normcase(str(path))
                if key not in seen:
                    seen.add(key)
                    cands.append((kind, str(path)))
        except Exception:
            pass

    for env_name, kind in (("UNRAR_TOOL", "unrar"), ("SEVENZIP_TOOL", "7z"),
                           ("SEVENZIP_PATH", "7z"), ("7Z_PATH", "7z")):
        add(kind, os.environ.get(env_name))
    for name in ("unrar", "UnRAR", "unrar.exe"):
        add("unrar", shutil.which(name))
    for name in ("7z", "7za", "7zz", "7z.exe"):
        add("7z", shutil.which(name))
    add("bsdtar", shutil.which("bsdtar"))
    base = app_dir()
    add("unrar", base / "UnRAR.exe")
    add("unrar", base / "unrar.exe")
    add("7z", base / "7z.exe")
    add("7z", base / "7za.exe")
    add("7z", base / "7zz.exe")
    add("bsdtar", base / "bsdtar.exe")
    for p in (
        r"C:\Program Files\WinRAR\UnRAR.exe",
        r"C:\Program Files (x86)\WinRAR\UnRAR.exe",
        r"C:\Program Files\7-Zip\7z.exe",
        r"C:\Program Files (x86)\7-Zip\7z.exe",
        str(Path(os.environ.get("LOCALAPPDATA", "")) / "Programs" / "7-Zip" / "7z.exe"),
    ):
        kind = "unrar" if "unrar" in p.lower() else ("7z" if "7-zip" in p.lower() else "bsdtar")
        add(kind, p)
    priority = {"unrar": 0, "7z": 1, "bsdtar": 2}
    cands.sort(key=lambda item: priority.get(item[0], 9))
    _RAR_TOOL = cands[0] if cands else None
    return _RAR_TOOL


def _configure_rarfile(rarfile, tool: tuple[str, str] | None) -> None:
    if not tool:
        return
    kind, path = tool
    try:
        if kind == "unrar":
            rarfile.UNRAR_TOOL = path
        elif kind == "7z":
            if hasattr(rarfile, "SEVENZIP_TOOL"):
                rarfile.SEVENZIP_TOOL = path
        elif kind == "bsdtar":
            if hasattr(rarfile, "BSDTAR_TOOL"):
                rarfile.BSDTAR_TOOL = path
    except Exception:
        pass


def _safe_archive_name(raw) -> str:
    name = str(raw).replace("\\", "/")
    return "/".join(re.sub(r'[<>:"|?*\x00-\x1f]', "_", part) for part in name.split("/"))


def _archive_name_unsafe(raw) -> bool:
    text = str(raw).replace("\\", "/")
    if text.startswith("/") or re.match(r"^[A-Za-z]:", text):
        return True
    name = _safe_archive_name(text)
    if not name:
        return True
    return any(part == ".." for part in name.split("/") if part)


def _archive_target(dest: Path, root: Path, raw):
    name = _safe_archive_name(raw)
    if not name or name in (".", ".."):
        return None, None
    target = (dest / name).resolve()
    if target != root and root not in target.parents:
        log.warning("Blocked unsafe archive path: %s", raw)
        return None, None
    return target, name


def _zip_member_name(m: zipfile.ZipInfo) -> str:
    name = m.filename
    if not (m.flag_bits & 0x800):  # no UTF-8 flag: stdlib guessed cp437
        try:
            name = name.encode("cp437").decode("utf-8")
        except (UnicodeEncodeError, UnicodeDecodeError):
            pass
    return _safe_archive_name(name)


def _prompt_archive_password(ctx: JobContext, p: Path, kind: str,
                             error: str = "") -> str:
    if ctx.cancelled:
        raise Cancelled()
    log.info("Password prompt for %s (%s)%s", p.name, kind,
             " after a failed attempt" if error else "")
    password = ctx.ask_password(p.name, kind, error)
    if password is None:
        log.info("Password entry cancelled for %s", p.name)
        if ctx.cancelled:
            raise Cancelled()
        raise RuntimeError("Password entry was cancelled.")
    return password


def _find_7z_tool() -> str | None:
    tool = _find_rar_tool()
    if tool and tool[0] == "7z":
        return tool[1]
    for name in ("7z", "7za", "7zz", "7z.exe"):
        found = shutil.which(name)
        if found:
            return found
    for candidate in (
        app_dir() / "7z.exe",
        app_dir() / "7za.exe",
        app_dir() / "7zz.exe",
        Path(r"C:\Program Files\7-Zip\7z.exe"),
        Path(r"C:\Program Files (x86)\7-Zip\7z.exe"),
        Path(os.environ.get("LOCALAPPDATA", "")) / "Programs" / "7-Zip" / "7z.exe",
    ):
        try:
            if candidate.is_file():
                return str(candidate)
        except Exception:
            pass
    return None


def _extract_zip_to(p: Path, dest: Path, ctx: JobContext,
                    password: str | None = None) -> None:
    root = dest.resolve()
    try:
        with zipfile.ZipFile(p) as zf:
            members = zf.infolist()
            encrypted = any(m.flag_bits & 0x1 for m in members)
            if encrypted and not password:
                raise ArchivePasswordError("Password required")
            pwd_bytes = password.encode("utf-8") if password else None
            for j, m in enumerate(members, 1):
                ctx.check()
                target, name = _archive_target(dest, root, _zip_member_name(m))
                if target is None:
                    continue
                if m.is_dir() or name.endswith("/"):
                    target.mkdir(parents=True, exist_ok=True)
                    continue
                target.parent.mkdir(parents=True, exist_ok=True)
                try:
                    with zf.open(m, pwd=pwd_bytes if m.flag_bits & 0x1 else None) as s:
                        with open(target, "wb") as d:
                            shutil.copyfileobj(s, d)
                except NotImplementedError as exc:
                    raise ArchiveUnsupportedError(str(exc)) from exc
                except RuntimeError as exc:
                    if "password" in str(exc).lower():
                        raise ArchivePasswordError("Wrong password") from exc
                    raise
                if j % 25 == 0:
                    ctx.progress(f"Extracting {p.name}: {j}/{len(members)}")
    except ArchivePasswordError:
        raise
    except NotImplementedError as exc:
        raise ArchiveUnsupportedError(str(exc)) from exc
    except RuntimeError as exc:
        if "password" in str(exc).lower():
            raise ArchivePasswordError("Password required") from exc
        raise


def _extract_zip_external_7z(p: Path, dest: Path, ctx: JobContext,
                             password: str | None) -> None:
    exe = _find_7z_tool()
    if not exe:
        raise RuntimeError(
            "This ZIP uses AES encryption, which Python's zip reader cannot open.\n"
            "Install 7-Zip so DropFlow can extract it.")
    pwd_args = [f"-p{password}"] if password else []
    listing = ctx.run([exe, "l", "-slt", *pwd_args, str(p)], label="Reading ZIP")
    detail = _tail((listing.stderr or "") + " " + (listing.stdout or ""))
    if listing.returncode != 0:
        low = detail.lower()
        if "password" in low or "encrypted" in low or (password and "wrong" in low):
            raise ArchivePasswordError("Wrong password" if password else "Password required")
        raise RuntimeError(f"7-Zip failed to read '{p.name}'. {detail}")
    output = listing.stdout or ""
    if "Encrypted = +" in output and not password:
        raise ArchivePasswordError("Password required")
    lines = output.splitlines()
    start = next((i for i, line in enumerate(lines)
                  if line.strip().startswith("----------")), -1)
    if start >= 0:
        names = [line[7:].strip() for line in lines[start + 1:]
                 if line.startswith("Path = ")]
        if any(_archive_name_unsafe(name) for name in names):
            raise RuntimeError(f"'{p.name}' contains unsafe paths and was not extracted.")
    res = ctx.run([exe, "x", "-y", f"-o{dest}", *pwd_args, str(p)],
                  label="Extracting ZIP")
    if res.returncode != 0:
        detail = _tail((res.stderr or "") + " " + (res.stdout or ""))
        low = detail.lower()
        if ("password" in low or "encrypted" in low or (password and "wrong" in low)
                or (password and "data error" in low)):
            raise ArchivePasswordError("Wrong password" if password else "Password required")
        raise RuntimeError(f"7-Zip failed for '{p.name}'. {detail}")


def _extract_zip_with_password(p: Path, dest: Path, ctx: JobContext) -> None:
    password = ctx.passwords.get(str(p))
    for attempt in range(4):
        ctx.check()
        try:
            try:
                _extract_zip_to(p, dest, ctx, password)
            except ArchiveUnsupportedError:
                _extract_zip_external_7z(p, dest, ctx, password)
            if password:
                ctx.passwords[str(p)] = password
            return
        except ArchivePasswordError:
            if password:
                log.warning("Password rejected for %s", p.name)
            if attempt >= 3:
                raise RuntimeError(
                    f"Could not unlock '{p.name}' after 3 password attempts.")
            password = _prompt_archive_password(
                ctx, p, "ZIP", "Wrong password. Try again." if password else "")


def _extract_rar_with_rarfile(p: Path, dest: Path, ctx: JobContext,
                              password: str | None = None) -> bool:
    try:
        import rarfile
    except ImportError:
        return False
    _configure_rarfile(rarfile, _find_rar_tool())
    rf = None
    try:
        rf = rarfile.RarFile(str(p))
        if password:
            if hasattr(rf, "setpassword"):
                rf.setpassword(password)
            else:
                log.warning("Installed rarfile does not support passwords; using external tool")
                return False
        if rf.needs_password() and not password:
            raise ArchivePasswordError("Password required")
        members = rf.infolist()
        root = dest.resolve()
        for j, info in enumerate(members, 1):
            ctx.check()
            target, name = _archive_target(dest, root, info.filename)
            if target is None:
                continue
            is_dir = getattr(info, "is_dir", lambda: False)()
            if is_dir or name.endswith("/"):
                target.mkdir(parents=True, exist_ok=True)
                continue
            target.parent.mkdir(parents=True, exist_ok=True)
            try:
                with rf.open(info) as s, open(target, "wb") as d:
                    shutil.copyfileobj(s, d)
            except Exception as exc:
                if type(exc).__name__ in ("PasswordRequired", "RarWrongPassword"):
                    raise ArchivePasswordError("Wrong password") from exc
                if "password" in str(exc).lower():
                    raise ArchivePasswordError("Wrong password") from exc
                raise
            if j % 25 == 0:
                ctx.progress(f"Extracting {p.name}: {j}/{len(members)}")
        return True
    except ArchivePasswordError:
        raise
    except RuntimeError as exc:
        if "password" in str(exc).lower():
            raise ArchivePasswordError("Wrong password") from exc
        log.warning("rarfile could not extract %s (%s); trying an external tool",
                    p.name, exc)
        return False
    except Exception as exc:
        if type(exc).__name__ in ("PasswordRequired", "RarWrongPassword"):
            raise ArchivePasswordError("Wrong password") from exc
        log.warning("rarfile could not extract %s (%s); trying an external tool",
                    p.name, exc)
        return False
    finally:
        if rf is not None:
            try:
                rf.close()
            except Exception:
                pass


def _rar_failure(p: Path, tool: str, detail: str) -> RuntimeError:
    low = (detail or "").lower()
    if "password" in low or "encrypted" in low or "passphrase" in low:
        return RuntimeError(f"'{p.name}' is password-protected.")
    return RuntimeError(f"{tool} failed for '{p.name}'. {detail}")


def _extract_rar_external(p: Path, dest: Path, ctx: JobContext,
                          password: str | None = None) -> None:
    tool = _find_rar_tool()
    if not tool:
        raise RuntimeError(_RAR_MISSING_MESSAGE)
    kind, exe = tool
    if kind == "unrar":
        pwd_arg = f"-p{password}" if password else "-p-"
        listing = ctx.run([exe, "lb", pwd_arg, str(p)], label="Reading RAR")
        if listing.returncode == 11:
            raise ArchivePasswordError("Wrong password" if password else "Password required")
        if listing.returncode != 0:
            detail = _tail((listing.stderr or "") + " " + (listing.stdout or ""))
            raise _rar_failure(p, "UnRAR", detail)
        names = [line.strip() for line in (listing.stdout or "").splitlines() if line.strip()]
        if any(_archive_name_unsafe(name) for name in names):
            raise RuntimeError(f"'{p.name}' contains unsafe paths and was not extracted.")
        res = ctx.run([exe, "x", "-o+", "-y", "-idq", pwd_arg,
                       str(p), str(dest) + os.sep], label="Extracting RAR")
        if res.returncode == 11:
            raise ArchivePasswordError("Wrong password" if password else "Password required")
        if res.returncode != 0:
            detail = _tail((res.stderr or "") + " " + (res.stdout or ""))
            raise _rar_failure(p, "UnRAR", detail)
        return

    if kind == "7z":
        pwd_args = [f"-p{password}"] if password else []
        listing = ctx.run([exe, "l", "-slt", *pwd_args, str(p)], label="Reading RAR")
        detail = _tail((listing.stderr or "") + " " + (listing.stdout or ""))
        if listing.returncode != 0:
            low = detail.lower()
            if "password" in low or "encrypted" in low or (password and "wrong" in low):
                raise ArchivePasswordError(
                    "Wrong password" if password else "Password required")
            raise _rar_failure(p, "7-Zip", detail)
        output = listing.stdout or ""
        if "Encrypted = +" in output and not password:
            raise ArchivePasswordError("Password required")
        lines = output.splitlines()
        start = next((i for i, line in enumerate(lines)
                      if line.strip().startswith("----------")), -1)
        if start >= 0:
            names = [line[7:].strip() for line in lines[start + 1:]
                     if line.startswith("Path = ")]
            if any(_archive_name_unsafe(name) for name in names):
                raise RuntimeError(f"'{p.name}' contains unsafe paths and was not extracted.")
        res = ctx.run([exe, "x", "-y", f"-o{dest}", *pwd_args, str(p)],
                      label="Extracting RAR")
        if res.returncode != 0:
            detail = _tail((res.stderr or "") + " " + (res.stdout or ""))
            low = detail.lower()
            if ("password" in low or "encrypted" in low or (password and "wrong" in low)
                    or (password and "data error" in low)):
                raise ArchivePasswordError(
                    "Wrong password" if password else "Password required")
            raise _rar_failure(p, "7-Zip", detail)
        return

    if kind == "bsdtar":
        if password:
            raise RuntimeError("Password-protected RAR needs UnRAR or 7-Zip; "
                               "bsdtar does not support it.")
        listing = ctx.run([exe, "-tf", str(p)], label="Reading RAR")
        if listing.returncode != 0:
            detail = _tail((listing.stderr or "") + " " + (listing.stdout or ""))
            low = detail.lower()
            if "password" in low or "encrypted" in low:
                raise ArchivePasswordError("Password required")
            raise _rar_failure(p, "bsdtar", detail)
        names = [line.strip() for line in (listing.stdout or "").splitlines() if line.strip()]
        if any(_archive_name_unsafe(name) for name in names):
            raise RuntimeError(f"'{p.name}' contains unsafe paths and was not extracted.")
        res = ctx.run([exe, "-xf", str(p), "-C", str(dest)], label="Extracting RAR")
        if res.returncode != 0:
            detail = _tail((res.stderr or "") + " " + (res.stdout or ""))
            raise _rar_failure(p, "bsdtar", detail)
        return

    raise RuntimeError(_RAR_MISSING_MESSAGE)


def _extract_rar_to(p: Path, dest: Path, ctx: JobContext) -> None:
    password = ctx.passwords.get(str(p))
    for attempt in range(4):
        ctx.check()
        try:
            if _extract_rar_with_rarfile(p, dest, ctx, password):
                if password:
                    ctx.passwords[str(p)] = password
                return
            _extract_rar_external(p, dest, ctx, password)
            if password:
                ctx.passwords[str(p)] = password
            return
        except ArchivePasswordError:
            if password:
                log.warning("Password rejected for %s", p.name)
            if attempt >= 3:
                raise RuntimeError(
                    f"Could not unlock '{p.name}' after 3 password attempts.")
            password = _prompt_archive_password(
                ctx, p, "RAR", "Wrong password. Try again." if password else "")


def _extract_archives(paths, ctx: JobContext) -> Result:
    archives = [Path(p) for p in paths
                if Path(p).is_file() and Path(p).suffix.lower() in ARCHIVE_EXTS]
    if not archives:
        raise RuntimeError("No ZIP or RAR archives to extract.")
    outs, errors = [], []
    try:
        for p in archives:
            ctx.check()
            ctx.progress(f"Extracting {p.name}\u2026")
            dest = None
            try:
                dest = get_unique_filename(p.with_name(f"{p.stem}{SUFFIX}"), is_dir=True)
                dest.mkdir(parents=True)
                if p.suffix.lower() == ".zip":
                    _extract_zip_with_password(p, dest, ctx)
                elif p.suffix.lower() == ".rar":
                    _extract_rar_to(p, dest, ctx)
                else:
                    raise RuntimeError(f"Unsupported archive: {p.name}")
                outs.append(dest)
            except Cancelled:
                if dest is not None:
                    discard([dest])
                raise
            except Exception as exc:
                if dest is not None:
                    discard([dest])
                _record_failure(errors, p.name, exc)
    except Cancelled:
        discard(outs)
        raise
    if not outs:
        if errors:
            raise RuntimeError("No archives could be extracted:\n"
                               + "\n".join(errors[:5]))
        return Result("No archives were extracted.")
    note = f" ({len(errors)} failed \u2014 see log)" if errors else ""
    return Result(f"Extracted {len(outs)} archive(s) into clean folder(s){note}",
                  outs, failed=len(errors))


def act_extract_zip(paths, opts, ctx: JobContext) -> Result:
    zips = [p for p in paths if Path(p).is_file() and Path(p).suffix.lower() == ".zip"]
    if not zips:
        raise RuntimeError("No ZIP files to extract.")
    return _extract_archives(zips, ctx)


def act_extract_rar(paths, opts, ctx: JobContext) -> Result:
    rars = [p for p in paths if Path(p).is_file() and Path(p).suffix.lower() == ".rar"]
    if not rars:
        raise RuntimeError("No RAR files to extract.")
    return _extract_archives(rars, ctx)


def act_extract_archive(paths, opts, ctx: JobContext) -> Result:
    return _extract_archives(paths, ctx)


def act_zip(paths, opts, ctx: JobContext) -> Result:
    paths = [Path(p) for p in paths]
    if not paths:
        raise RuntimeError("Nothing to archive.")
    first = paths[0]
    if len(paths) == 1:
        base = first.name if first.is_dir() else first.stem
    else:
        base = first.parent.name or "archive"
    out = get_unique_filename(first.parent / f"{base}{SUFFIX}.zip")
    out_abs = os.path.normcase(os.path.abspath(str(out)))
    count = 0
    errors: list[str] = []
    try:
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as zf:
            for p in paths:
                if p.is_dir():
                    for root_dir, dirs, files in os.walk(p):
                        ctx.check()
                        rel_root = Path(root_dir).relative_to(p.parent)
                        if not dirs and not files:
                            try:
                                zf.write(root_dir, rel_root)
                            except Exception as exc:
                                _record_failure(errors, str(root_dir), exc)
                        for f in files:
                            ctx.check()
                            fp = Path(root_dir) / f
                            if os.path.normcase(os.path.abspath(str(fp))) == out_abs:
                                continue
                            try:
                                zf.write(fp, rel_root / f)
                                count += 1
                            except Exception as exc:
                                _record_failure(errors, str(fp), exc)
                            if count % 50 == 0:
                                ctx.progress(f"Archiving\u2026 {count} files")
                else:
                    ctx.check()
                    try:
                        zf.write(p, p.name)
                        count += 1
                    except Exception as exc:
                        _record_failure(errors, str(p), exc)
    except Cancelled:
        out.unlink(missing_ok=True)
        raise
    except Exception:
        out.unlink(missing_ok=True)
        raise
    if count == 0:
        out.unlink(missing_ok=True)
        if errors:
            raise RuntimeError("Nothing could be archived:\n" + "\n".join(errors[:5]))
        return Result("Nothing to archive.")
    note = f" ({len(errors)} failed \u2014 see log)" if errors else ""
    return Result(f"Created {out.name}: {count} file(s), "
                  f"{fmt_size(out.stat().st_size)}{note}", [out], failed=len(errors))


# --------------------------------------------------------------------------- #
#  Context-aware action menu
# --------------------------------------------------------------------------- #
def describe(paths: list[Path]) -> str:
    names = {".pdf": "PDF", ".docx": "Word doc", ".zip": "ZIP", ".rar": "RAR",
             ".pptx": "PowerPoint", ".ppt": "PowerPoint", ".xlsx": "Excel",
             ".xls": "Excel", ".txt": "text file", ".csv": "CSV", ".json": "JSON"}
    c: Counter = Counter()
    for p in paths:
        if p.is_dir():
            c["folder"] += 1
        elif p.suffix.lower() in IMG_EXTS:
            c["image"] += 1
        elif p.suffix.lower() in VIDEO_EXTS:
            c["video"] += 1
        else:
            c[names.get(p.suffix.lower(), "file")] += 1
    return "  \u00b7  ".join(f"{v} {k}{'s' if v > 1 else ''}" for k, v in c.items())


def build_actions(paths: list[Path]) -> tuple[str, list[Action]]:
    all_files = all(p.is_file() for p in paths)
    many = len(paths) > 1
    acts: list[Action] = []

    if not all_files:  # folders (possibly mixed with files)
        acts.append(Action("Compress to ZIP", act_zip))
        return describe(paths), acts

    pdfs = [p for p in paths if p.suffix.lower() == ".pdf"]
    images = [p for p in paths if p.suffix.lower() in IMG_EXTS]
    videos = [p for p in paths if p.suffix.lower() in VIDEO_EXTS]
    docs = [p for p in paths if p.suffix.lower() == ".docx"]
    office_docs = [p for p in paths if p.suffix.lower() in OFFICE_INPUT_EXTS]
    zips = [p for p in paths if p.suffix.lower() == ".zip"]
    rars = [p for p in paths if p.suffix.lower() == ".rar"]
    archives = zips + rars

    if pdfs and len(pdfs) == len(paths):
        if many:
            acts.append(Action("Merge into one PDF", act_merge_pdfs))
        acts += [
            Action("Compress PDF" + ("s" if many else ""), act_compress_pdfs,
                   ("quality", "Image quality (lower = smaller file)", 10, 95, 60, "%", 5)),
            Action("Split every N pages", act_split_pdf,
                   ("pages", "Pages per file", 1, 50, 1, "", 1)),
            Action("Extract page range\u2026", act_extract_pages,
                   ask_text=("pages", "Pages to extract", "1-3",
                             "Examples: 1-3, 5, 7-9, all.")),
            Action("Split at page ranges\u2026", act_split_ranges,
                   ask_text=("pages", "Page ranges", "1-3, 5, 7-9",
                             "Creates one PDF per range.")),
            Action("Export pages as PNG", make_pdf_to_images("PNG"),
                   ("dpi", "Resolution", 72, 300, 150, " dpi", 10)),
            Action("Export pages as JPG", make_pdf_to_images("JPG"),
                   ("quality", "JPEG quality", 40, 100, 85, "%", 5)),
            Action("Create searchable PDF (OCR)", act_ocr_pdf_to_pdf,
                   ("dpi", "OCR render resolution", 150, 400, 250, " dpi", 25)),
            Action("Extract text (.txt, OCR)", act_ocr_pdf,
                   ("dpi", "OCR render resolution", 150, 400, 250, " dpi", 25)),
            Action("Create PowerPoint\u2026", act_images_pdf_to_pptx),
            Action("Create Excel workbook\u2026", act_data_to_xlsx),
        ]
    elif images and len(images) == len(paths):
        if many:
            acts += [
                Action("Combine into one PDF", act_images_to_pdf),
                Action("Create animated GIF", act_images_to_gif,
                       ("duration", "Time per frame", 100, 2000, 500, " ms", 50)),
            ]
            exts = {p.suffix.lower() for p in images}
            for fmt in ("PNG", "JPG", "WEBP"):
                if exts - _SAME_FORMAT[fmt]:
                    acts.append(Action(f"Convert all to {fmt}", make_convert(fmt)))
        else:
            ext = paths[0].suffix.lower()
            for fmt in ("PNG", "JPG", "WEBP"):
                if ext not in _SAME_FORMAT[fmt]:
                    acts.append(Action(f"Convert to {fmt}", make_convert(fmt)))
            acts.append(Action("Convert to PDF", act_images_to_pdf))
        acts += [
            Action("Compress image" + ("s" if many else ""), act_compress_image,
                   ("quality", "Quality (lower = smaller file)", 10, 100, 70, "%", 1)),
            Action("Create PowerPoint\u2026", act_images_pdf_to_pptx),
            Action("Create searchable PDF (OCR)", act_ocr_to_pdf),
            Action("Extract text (.txt, OCR)", act_ocr),
        ]
    elif videos and len(videos) == len(paths):
        acts += [
            Action("Compress video" + ("s" if many else ""), act_compress_video,
                   video_options=True),
            Action("Universal media converter\u2026", act_convert_media,
                   media_options=True),
            Action("Extract audio (MP3)", act_extract_audio,
                   ("quality", "Audio bitrate", 96, 320, 192, " kbps", 32)),
        ]
    elif docs and len(docs) == len(paths):
        acts += [
            Action("Convert to PDF", act_docx_to_pdf),
            Action("Extract text to .txt", act_docx_to_txt),
            Action("Universal document converter\u2026", act_convert_office,
                   office_options=True),
        ]
    elif office_docs and len(office_docs) == len(paths):
        acts.append(Action("Universal document converter\u2026", act_convert_office,
                           office_options=True))
        if all(p.suffix.lower() in (".csv", ".json") for p in office_docs):
            acts.append(Action("Create Excel workbook\u2026", act_data_to_xlsx))
    elif zips and len(zips) == len(paths):
        acts.append(Action("Extract to a new folder", act_extract_zip))
    elif rars and len(rars) == len(paths):
        acts.append(Action("Extract RAR to a new folder", act_extract_rar))
    elif archives and len(archives) == len(paths):
        acts.append(Action("Extract archives to new folders", act_extract_archive))
    elif pdfs and images and len(pdfs) + len(images) == len(paths):
        acts.append(Action("Combine all into one PDF", act_mixed_to_pdf))
        acts.append(Action("Create PowerPoint\u2026", act_images_pdf_to_pptx))
        acts.append(Action("Compress PDF" + ("s" if len(pdfs) > 1 else ""),
                           act_compress_pdfs,
                           ("quality", "Image quality (lower = smaller file)", 10, 95, 60, "%", 5)))

    if office_docs and len(office_docs) < len(paths):
        acts.append(Action("Universal document converter\u2026", act_convert_office,
                           office_options=True))
    if videos and len(videos) < len(paths):
        acts.append(Action("Universal media converter\u2026", act_convert_media,
                           media_options=True))
    if archives and len(archives) < len(paths):
        acts.append(Action("Extract archive" + ("s" if len(archives) > 1 else ""),
                           act_extract_archive))
    if not (archives and len(archives) == len(paths)):
        acts.append(Action("Compress to ZIP", act_zip))
    return describe(paths), acts


# --------------------------------------------------------------------------- #
#  Global hotkey: native RegisterHotKey (no keyboard hooks)
# --------------------------------------------------------------------------- #
WM_HOTKEY = 0x0312
WM_APP = 0x8000
MOD_ALT = 0x0001
MOD_CONTROL = 0x0002
MOD_SHIFT = 0x0004
MOD_WIN = 0x0008
MOD_NOREPEAT = 0x4000

_MOD_ALIASES = {
    "ctrl": MOD_CONTROL, "control": MOD_CONTROL,
    "alt": MOD_ALT, "shift": MOD_SHIFT,
    "win": MOD_WIN, "windows": MOD_WIN, "super": MOD_WIN, "meta": MOD_WIN,
}
_KEY_VK = {
    "space": 0x20, "enter": 0x0D, "return": 0x0D, "esc": 0x1B, "escape": 0x1B,
    "tab": 0x09, "backspace": 0x08, "delete": 0x2E, "del": 0x2E, "insert": 0x2D,
    "home": 0x24, "end": 0x23, "pageup": 0x21, "pagedown": 0x22,
    "left": 0x25, "up": 0x26, "right": 0x27, "down": 0x28,
    "printscreen": 0x2C, "pause": 0x13, "capslock": 0x14,
    "numlock": 0x90, "scrolllock": 0x91,
    "minus": 0xBD, "plus": 0xBB, "equal": 0xBB, "comma": 0xBC,
    "period": 0xBE, "slash": 0xBF, "backslash": 0xDC, "semicolon": 0xBA,
    "quote": 0xDE, "bracketleft": 0xDB, "bracketright": 0xDD, "grave": 0xC0,
    "multiply": 0x6A, "add": 0x6B, "subtract": 0x6D, "decimal": 0x6E, "divide": 0x6F,
}
for _n in range(10):
    _KEY_VK[f"num{_n}"] = 0x60 + _n

_HOTKEY_ERRORS = {
    1409: "That shortcut is already used by Windows or another application.",
    1419: "That shortcut is reserved by Windows.",
}


def parse_hotkey(hotkey: str) -> tuple[int, int, str]:
    parts = [p.strip().lower() for p in (hotkey or "").split("+") if p.strip()]
    mods, key = 0, None
    for part in parts:
        if part in _MOD_ALIASES:
            mods |= _MOD_ALIASES[part]
        elif key is None:
            key = part
        else:
            return 0, 0, "Use only one normal key in the shortcut."
    if not mods:
        return 0, 0, "Use at least one modifier (Ctrl, Alt, Shift or Win)."
    if key is None:
        return 0, 0, "Add a normal key to the shortcut."
    if re.fullmatch(r"f(?:[1-9]|1[0-9]|2[0-4])", key):
        vk = 0x70 + int(key[1:]) - 1
    elif len(key) == 1 and key.isalnum():
        vk = ord(key.upper())
    else:
        vk = _KEY_VK.get(key)
    if not vk:
        return 0, 0, f"'{key}' cannot be used as a shortcut key."
    return mods, vk, ""


WNDPROC = ctypes.WINFUNCTYPE(ctypes.c_ssize_t, wintypes.HWND, ctypes.c_uint,
                             wintypes.WPARAM, wintypes.LPARAM)


class WNDCLASSW(ctypes.Structure):
    _fields_ = [
        ("style", ctypes.c_uint),
        ("lpfnWndProc", WNDPROC),
        ("cbClsExtra", ctypes.c_int),
        ("cbWndExtra", ctypes.c_int),
        ("hInstance", wintypes.HINSTANCE),
        ("hIcon", wintypes.HANDLE),
        ("hCursor", wintypes.HANDLE),
        ("hbrBackground", wintypes.HANDLE),
        ("lpszMenuName", wintypes.LPCWSTR),
        ("lpszClassName", wintypes.LPCWSTR),
    ]


class HotkeyManager:
    """Owns a hidden message-only window and one system-wide hotkey."""

    HOTKEY_ID = 0xD0F1

    def __init__(self, callback: Callable[[], None]):
        self._callback = callback
        self._thread: threading.Thread | None = None
        self._hwnd = None
        self._ready = threading.Event()
        self._done = threading.Event()
        self._lock = threading.Lock()
        self._cmd: tuple | None = None
        self._result: tuple[bool, str] = (False, "")
        self._current: str | None = None
        self._wndproc_ref = None

    def start(self) -> None:
        if self._thread and self._thread.is_alive():
            return
        self._thread = threading.Thread(target=self._run, name="DropFlowHotkey", daemon=True)
        self._thread.start()
        if not self._ready.wait(5):
            log.error("Hotkey thread did not start in time")

    def register(self, hotkey: str) -> tuple[bool, str]:
        return self._post(("register", hotkey))

    def unregister(self) -> tuple[bool, str]:
        return self._post(("unregister",))

    def stop(self) -> None:
        if self._thread and self._thread.is_alive():
            if self._hwnd:
                self._post(("stop",))
            try:
                self._thread.join(2)
            except Exception:
                pass

    def _post(self, cmd: tuple) -> tuple[bool, str]:
        if not self._hwnd:
            return False, "The hotkey engine is not running."
        with self._lock:
            self._cmd = cmd
            self._done.clear()
        if not _user32.PostMessageW(self._hwnd, WM_APP, 0, 0):
            return False, "Could not reach the hotkey engine."
        if not self._done.wait(3):
            return False, "The hotkey engine did not respond."
        with self._lock:
            return self._result

    def _run(self) -> None:
        try:
            self._create_window_and_loop()
        except Exception:
            log.exception("Hotkey thread failed")
        finally:
            self._ready.set()
            self._hwnd = None

    def _create_window_and_loop(self) -> None:
        hinst = _kernel32.GetModuleHandleW(None)
        class_name = f"DropFlowHotkey_{os.getpid()}"
        wndproc = WNDPROC(self._wndproc)
        self._wndproc_ref = wndproc  # keep alive

        wc = WNDCLASSW()
        wc.style = 0
        wc.lpfnWndProc = wndproc
        wc.hInstance = hinst
        wc.lpszClassName = class_name
        atom = _user32.RegisterClassW(ctypes.byref(wc))
        if not atom:
            err = ctypes.get_last_error()
            if err != 1410:  # class already exists
                log.error("RegisterClassW failed: %s", err)
        hwnd = _user32.CreateWindowExW(
            0, class_name, "DropFlowHotkey", 0, 0, 0, 0, 0,
            ctypes.c_void_p(-3),  # HWND_MESSAGE
            None, hinst, None)
        if not hwnd:
            log.error("Could not create the hotkey message window")
            return
        self._hwnd = hwnd
        self._ready.set()

        msg = wintypes.MSG()
        while True:
            ret = _user32.GetMessageW(ctypes.byref(msg), None, 0, 0)
            if ret <= 0:
                break
            _user32.TranslateMessage(ctypes.byref(msg))
            _user32.DispatchMessageW(ctypes.byref(msg))
        try:
            _user32.UnregisterHotKey(hwnd, self.HOTKEY_ID)
            _user32.DestroyWindow(hwnd)
            _user32.UnregisterClassW(class_name, hinst)
        except Exception:
            pass

    def _wndproc(self, hwnd, msg, wparam, lparam):
        try:
            if msg == WM_HOTKEY:
                if wparam == self.HOTKEY_ID:
                    try:
                        self._callback()
                    except Exception:
                        log.exception("Hotkey callback failed")
                return 0
            if msg == WM_APP:
                with self._lock:
                    cmd = self._cmd
                result = (False, "Unknown hotkey command.")
                try:
                    if cmd and cmd[0] == "register":
                        result = self._register_now(cmd[1])
                    elif cmd and cmd[0] == "unregister":
                        result = self._unregister_now()
                    elif cmd and cmd[0] == "stop":
                        result = self._unregister_now()
                        _user32.PostQuitMessage(0)
                except Exception as exc:
                    log.exception("Hotkey command failed")
                    result = (False, str(exc))
                with self._lock:
                    self._result = result
                    self._done.set()
                return 0
            if msg == 0x0010:  # WM_CLOSE
                _user32.PostQuitMessage(0)
                return 0
        except Exception:
            log.exception("Hotkey window procedure failed")
        return _user32.DefWindowProcW(hwnd, msg, wparam, lparam)

    def _register_now(self, hotkey: str) -> tuple[bool, str]:
        self._unregister_now()
        mods, vk, err = parse_hotkey(hotkey)
        if err:
            return False, err
        if not _user32.RegisterHotKey(self._hwnd, self.HOTKEY_ID,
                                      mods | MOD_NOREPEAT, vk):
            code = ctypes.get_last_error()
            return False, _HOTKEY_ERRORS.get(code, f"Windows refused it (error {code}).")
        self._current = hotkey
        return True, ""

    def _unregister_now(self) -> tuple[bool, str]:
        if self._hwnd:
            _user32.UnregisterHotKey(self._hwnd, self.HOTKEY_ID)
        self._current = None
        return True, ""


# Prototypes used by the helpers above.  Setting restypes matters on 64-bit
# Windows: without them ctypes truncates window handles to 32 bits.
_user32.GetSystemMetrics.argtypes = [ctypes.c_int]
_user32.GetSystemMetrics.restype = ctypes.c_int
_user32.MonitorFromPoint.argtypes = [wintypes.POINT, wintypes.DWORD]
_user32.MonitorFromPoint.restype = wintypes.HANDLE
_user32.GetMonitorInfoW.argtypes = [wintypes.HANDLE, ctypes.POINTER(_MONITORINFO)]
_user32.GetMonitorInfoW.restype = wintypes.BOOL
_user32.GetDpiForSystem.argtypes = []
_user32.GetDpiForSystem.restype = wintypes.UINT
if _shcore is not None:
    try:
        _shcore.GetDpiForMonitor.argtypes = [wintypes.HANDLE, ctypes.c_int,
                                             ctypes.POINTER(wintypes.UINT),
                                             ctypes.POINTER(wintypes.UINT)]
        _shcore.GetDpiForMonitor.restype = ctypes.c_long
    except Exception:
        pass
_user32.GetAncestor.argtypes = [wintypes.HWND, ctypes.c_uint]
_user32.GetAncestor.restype = wintypes.HWND
_user32.GetForegroundWindow.restype = wintypes.HWND
_user32.SetForegroundWindow.argtypes = [wintypes.HWND]
_user32.SetForegroundWindow.restype = wintypes.BOOL
_user32.BringWindowToTop.argtypes = [wintypes.HWND]
_user32.BringWindowToTop.restype = wintypes.BOOL
_user32.SetFocus.argtypes = [wintypes.HWND]
_user32.SetFocus.restype = wintypes.HWND
_user32.GetWindowThreadProcessId.argtypes = [wintypes.HWND, ctypes.POINTER(wintypes.DWORD)]
_user32.GetWindowThreadProcessId.restype = wintypes.DWORD
_user32.AttachThreadInput.argtypes = [wintypes.DWORD, wintypes.DWORD, wintypes.BOOL]
_user32.AttachThreadInput.restype = wintypes.BOOL
_user32.keybd_event.argtypes = [wintypes.BYTE, wintypes.BYTE, wintypes.DWORD, ctypes.c_size_t]
_user32.PostMessageW.argtypes = [wintypes.HWND, ctypes.c_uint, wintypes.WPARAM, wintypes.LPARAM]
_user32.PostMessageW.restype = wintypes.BOOL
_user32.RegisterHotKey.argtypes = [wintypes.HWND, ctypes.c_int, ctypes.c_uint, ctypes.c_uint]
_user32.RegisterHotKey.restype = wintypes.BOOL
_user32.UnregisterHotKey.argtypes = [wintypes.HWND, ctypes.c_int]
_user32.UnregisterHotKey.restype = wintypes.BOOL
_user32.RegisterClassW.argtypes = [ctypes.POINTER(WNDCLASSW)]
_user32.RegisterClassW.restype = ctypes.c_ushort
_user32.UnregisterClassW.argtypes = [wintypes.LPCWSTR, wintypes.HINSTANCE]
_user32.UnregisterClassW.restype = wintypes.BOOL
_user32.CreateWindowExW.argtypes = [
    wintypes.DWORD, wintypes.LPCWSTR, wintypes.LPCWSTR, wintypes.DWORD,
    ctypes.c_int, ctypes.c_int, ctypes.c_int, ctypes.c_int,
    wintypes.HWND, wintypes.HANDLE, wintypes.HINSTANCE, ctypes.c_void_p]
_user32.CreateWindowExW.restype = wintypes.HWND
_user32.DestroyWindow.argtypes = [wintypes.HWND]
_user32.DestroyWindow.restype = wintypes.BOOL
_user32.DefWindowProcW.argtypes = [wintypes.HWND, ctypes.c_uint,
                                   wintypes.WPARAM, wintypes.LPARAM]
_user32.DefWindowProcW.restype = ctypes.c_ssize_t
_user32.GetMessageW.argtypes = [ctypes.POINTER(wintypes.MSG), wintypes.HWND,
                                wintypes.UINT, wintypes.UINT]
_user32.GetMessageW.restype = wintypes.BOOL
_user32.TranslateMessage.argtypes = [ctypes.POINTER(wintypes.MSG)]
_user32.DispatchMessageW.argtypes = [ctypes.POINTER(wintypes.MSG)]
_user32.DispatchMessageW.restype = ctypes.c_ssize_t
_user32.PostQuitMessage.argtypes = [ctypes.c_int]
_user32.GetAsyncKeyState.argtypes = [ctypes.c_int]
_user32.GetAsyncKeyState.restype = ctypes.c_short
_kernel32.GetCurrentThreadId.restype = wintypes.DWORD
_kernel32.GetModuleHandleW.argtypes = [wintypes.LPCWSTR]
_kernel32.GetModuleHandleW.restype = wintypes.HINSTANCE


_MODIFIER_KEYSYMS = {
    "Control_L", "Control_R", "Alt_L", "Alt_R", "Shift_L", "Shift_R",
    "Super_L", "Super_R", "Meta_L", "Meta_R", "Win_L", "Win_R",
}
_KEYSYM_TOKEN = {
    "space": "space", "Return": "enter", "KP_Enter": "enter",
    "Escape": "esc", "BackSpace": "backspace", "Tab": "tab",
    "ISO_Left_Tab": "tab", "Delete": "delete", "Insert": "insert",
    "Home": "home", "End": "end", "Prior": "pageup", "Next": "pagedown",
    "Left": "left", "Up": "up", "Right": "right", "Down": "down",
    "Print": "printscreen", "Pause": "pause", "Caps_Lock": "capslock",
    "Num_Lock": "numlock", "Scroll_Lock": "scrolllock",
    "minus": "minus", "plus": "plus", "equal": "equal", "comma": "comma",
    "period": "period", "slash": "slash", "backslash": "backslash",
    "semicolon": "semicolon", "apostrophe": "quote", "quotedbl": "quote",
    "bracketleft": "bracketleft", "bracketright": "bracketright",
    "grave": "grave", "dead_grave": "grave",
    "KP_Add": "add", "KP_Subtract": "subtract", "KP_Multiply": "multiply",
    "KP_Divide": "divide", "KP_Decimal": "decimal",
}
for _n in range(10):
    _KEYSYM_TOKEN[f"KP_{_n}"] = f"num{_n}"


def _key_down(vk: int) -> bool:
    try:
        return bool(_user32.GetAsyncKeyState(vk) & 0x8000)
    except Exception:
        return False


def _mods_down() -> list[str]:
    mods = []
    if _key_down(0x11):
        mods.append("ctrl")
    if _key_down(0x12):
        mods.append("alt")
    if _key_down(0x10):
        mods.append("shift")
    if _key_down(0x5B) or _key_down(0x5C):
        mods.append("win")
    return mods


def _keysym_token(keysym: str, event) -> str | None:
    if keysym in _KEYSYM_TOKEN:
        return _KEYSYM_TOKEN[keysym]
    if re.fullmatch(r"F(?:[1-9]|1[0-9]|2[0-4])", keysym):
        return keysym.lower()
    if len(keysym) == 1 and keysym.isalnum():
        return keysym.lower()
    ch = getattr(event, "char", "") or ""
    if len(ch) == 1 and ch.isalnum():
        return ch.lower()
    return None


# --------------------------------------------------------------------------- #
#  UI: floating drop window
# --------------------------------------------------------------------------- #
class DropWindow:
    MIN_W = 340

    def __init__(self, app: "App", cx: int, cy: int):
        self.app, self.cx, self.cy = app, cx, cy
        self.paths: list[Path] = []
        self.actions: list[Action] = []
        self.alive, self.closing, self.busy = True, False, False
        self.state = "zone"
        self._timer = None
        self._job: JobContext | None = None
        self.status: tk.Label | None = None
        self.cancel_btn: tk.Label | None = None
        self.password_dialog: "PasswordDialog | None" = None

        t = self.top = tk.Toplevel(app.root)
        t.withdraw()
        t.overrideredirect(True)
        t.configure(bg=ACCENT)  # 1px accent border
        t.attributes("-topmost", True)
        t.attributes("-alpha", 0.0)
        self.body = tk.Frame(t, bg=BG)
        self.body.pack(fill="both", expand=True, padx=1, pady=1)
        t.bind("<Escape>", self._on_escape)
        t.bind("<Motion>", self._on_motion)
        self._bind_dnd(t)
        self._bind_dnd(self.body)
        self.show_drop_zone()
        t.deiconify()
        self._fit()
        self._fade(0.0, 1.0, 0.16)
        t.after(30, self._refocus)

    # ----- plumbing ------------------------------------------------------- #
    def _fade(self, current, target, step, then=None):
        if not self.alive:
            return
        nxt = min(current + step, target) if target > current else max(current - step, target)
        try:
            self.top.attributes("-alpha", nxt)
        except tk.TclError:
            return
        if nxt == target:
            if then:
                then()
        else:
            self.top.after(14, self._fade, nxt, target, step, then)

    def _fit(self, min_w: int | None = None):
        t = self.top
        t.update_idletasks()
        w = max(t.winfo_reqwidth(), px(min_w or self.MIN_W))
        h = t.winfo_reqheight()
        vx, vy, vw, vh = virtual_screen_rect()
        x = max(vx + 8, min(self.cx - w // 2, vx + vw - w - 8))
        y = max(vy + 8, min(self.cy - h // 2, vy + vh - h - 56))
        t.geometry(f"{w}x{h}+{x}+{y}")

    def _refocus(self):
        if not self.alive:
            return
        try:
            self.top.lift()
            force_foreground(self.top)
        except tk.TclError:
            pass

    def _bind_dnd(self, widget) -> None:
        if getattr(widget, "_df_dnd", False):
            return
        try:
            widget.drop_target_register(DND_FILES)
            widget.dnd_bind("<<Drop>>", self._on_drop)
            widget.dnd_bind("<<DropEnter>>", self._on_enter)
            widget.dnd_bind("<<DropLeave>>", self._on_leave)
            widget._df_dnd = True
        except Exception:
            log.debug("DND registration failed for %s", widget, exc_info=True)

    def _bind_dnd_tree(self) -> None:
        def walk(widget):
            self._bind_dnd(widget)
            for child in widget.winfo_children():
                walk(child)
        walk(self.body)

    def _clear(self):
        self._cancel_timer()
        self._close_password_dialog()
        try:
            self.top.unbind("<Return>")
            self.top.unbind("<KP_Enter>")
        except tk.TclError:
            pass
        for w in self.body.winfo_children():
            w.destroy()
        self.status = None
        self.cancel_btn = None

    def _close_password_dialog(self):
        dialog = self.password_dialog
        self.password_dialog = None
        if dialog is not None:
            try:
                dialog.cancel()
            except Exception:
                pass

    def show_password_dialog(self, request: PasswordRequest):
        if not self.alive or not self.busy or self.closing:
            request.cancelled = True
            request.event.set()
            return
        self._close_password_dialog()
        try:
            self.password_dialog = PasswordDialog(self, request)
        except Exception:
            log.exception("Could not create the password dialog")
            request.cancelled = True
            request.event.set()
            self.password_dialog = None

    def _cancel_timer(self):
        if self._timer is not None:
            try:
                self.top.after_cancel(self._timer)
            except Exception:
                pass
            self._timer = None

    def _arm_idle(self, ms: int = 30000):
        self._cancel_timer()
        if self.state == "zone" and self.alive:
            self._timer = self.top.after(ms, self.close)

    def _on_motion(self, _event=None):
        if self.state == "zone":
            self._arm_idle()
        elif self.state == "done":
            self._cancel_timer()

    def _on_escape(self, _event=None):
        if self.state == "params":
            self.show_menu(self.paths)
        elif self.state == "working":
            self.request_cancel()
        else:
            self.close()
        return "break"

    def close(self, *_):
        if self.closing or not self.alive:
            return
        if self.busy:
            self.request_cancel()
            return
        self.closing = True
        self._cancel_timer()
        try:
            cur = float(self.top.attributes("-alpha"))
        except tk.TclError:
            cur = 1.0
        self._fade(cur, 0.0, 0.2, then=self._destroy)

    def _destroy(self):
        if not self.alive:
            return
        self.alive = False
        self._close_password_dialog()
        if self._job is not None:
            self._job.cancel()
        try:
            self.top.destroy()
        except tk.TclError:
            pass
        if self.app.win is self:
            self.app.win = None

    def shutdown(self):
        """Called when the whole app quits; kills any running helper process."""
        self._close_password_dialog()
        if self._job is not None:
            self._job.cancel()
        self.alive = False
        try:
            self.top.destroy()
        except Exception:
            pass

    # ----- small widgets -------------------------------------------------- #
    def _button(self, parent, text, command, primary=False, center=False):
        bg, fg = (ACCENT, BG) if primary else (SURFACE, FG)
        hbg, hfg = (ACCENT_HI, BG) if primary else (HOVER, ACCENT)
        b = tk.Label(parent, text=text, bg=bg, fg=fg, cursor="hand2",
                     font=(FONT, 10, "bold" if primary else "normal"),
                     anchor="center" if center else "w", padx=px(12), pady=px(7))

        def paint(b_bg, b_fg):
            try:
                b.config(bg=b_bg, fg=b_fg)
            except tk.TclError:
                pass

        b.bind("<Enter>", lambda e: paint(hbg, hfg))
        b.bind("<Leave>", lambda e: paint(bg, fg))
        b.bind("<Button-1>", lambda e: self.top.after(1, command))
        return b

    def _x_button(self, parent):
        x = tk.Label(parent, text="\u2715", fg=SUBTLE, bg=BG, cursor="hand2", font=(FONT, 10))
        x.bind("<Enter>", lambda e: x.config(fg=RED))
        x.bind("<Leave>", lambda e: x.config(fg=SUBTLE))
        x.bind("<Button-1>", lambda e: self.top.after(1, self.close))
        return x

    # ----- views ---------------------------------------------------------- #
    def show_drop_zone(self):
        self._clear()
        self.state = "zone"
        W, H = px(310), px(176)
        c = tk.Canvas(self.body, width=W, height=H, bg=BG, highlightthickness=0)
        c.pack(padx=px(10), pady=(px(10), px(4)))
        self.zone = c
        self.rect = c.create_rectangle(px(4), px(4), W - px(4), H - px(4),
                                       outline=ACCENT, dash=(6, 4), width=2)
        mid = W // 2
        c.create_line(mid, px(34), mid, px(76), width=4, fill=ACCENT, arrow=tk.LAST,
                      arrowshape=(px(14), px(16), px(7)))
        c.create_line(mid - px(28), px(66), mid - px(28), px(90), mid + px(28), px(90),
                      mid + px(28), px(66), width=3, fill=ACCENT, joinstyle=tk.ROUND)
        c.create_text(mid, px(118), text="Drop files or folders here", fill=FG,
                      font=(FONT, 11, "bold"))
        c.create_text(mid, px(142), text="PDF  \u00b7  images  \u00b7  Word  \u00b7  video  \u00b7  ZIP",
                      fill=SUBTLE, font=(FONT, 9))
        tk.Label(self.body, text=f"{pretty_hotkey(self.app.cfg['hotkey'])} toggles  \u00b7  Esc closes",
                 fg=SUBTLE, bg=BG, font=(FONT, 8)).pack(pady=(0, px(8)))
        self._x_button(self.body).place(relx=1.0, x=-px(8), y=px(6), anchor="ne")
        self._bind_dnd_tree()
        self._fit()
        self._arm_idle()

    def _hover(self, on: bool):
        try:
            self.zone.itemconfig(self.rect, outline=ACCENT_HI if on else ACCENT,
                                 width=3 if on else 2)
            self.zone.configure(bg=HOVER if on else BG)
        except (tk.TclError, AttributeError):
            pass

    def _on_enter(self, event):
        self._cancel_timer()
        if self.state == "zone":
            self._hover(True)
        return event.action

    def _on_leave(self, event):
        if self.state == "zone":
            self._hover(False)
            self._arm_idle()
        return event.action

    def _on_drop(self, event):
        if self.busy or self.closing:
            return event.action
        paths = clean_dropped_paths(self.app.root, event.data)
        if not paths:
            self._hover(False)
            return event.action
        self.top.after(10, lambda: self._apply_drop(paths))
        return event.action

    def _apply_drop(self, paths: list[Path]) -> None:
        if not self.alive or self.busy:
            return
        if self.state == "menu":
            existing = {os.path.normcase(str(p)) for p in self.paths}
            merged = list(self.paths)
            for p in paths:
                key = os.path.normcase(str(p))
                if key not in existing:
                    existing.add(key)
                    merged.append(p)
            paths = merged
        self.show_menu(paths)
        self.top.after(60, self._refocus)

    def show_menu(self, paths: list[Path]):
        if not self.alive:
            return
        if all(p.is_file() for p in paths) and all(
                p.suffix.lower() == ".pdf" or p.suffix.lower() in IMG_EXTS for p in paths):
            paths = sorted(paths, key=natural_key)  # deterministic page order
        self.paths = paths
        summary, self.actions = build_actions(paths)
        self._clear()
        self.state = "menu"
        pad = px(14)
        head = tk.Frame(self.body, bg=BG)
        head.pack(fill="x", padx=pad, pady=(px(12), 0))
        tk.Label(head, text=summary, fg=FG, bg=BG, font=(FONT, 11, "bold"),
                 anchor="w").pack(side="left")
        self._x_button(head).pack(side="right")
        names = ", ".join(p.name for p in paths[:3])
        if len(paths) > 3:
            names += f"  +{len(paths) - 3} more"
        tk.Label(self.body, text=names, fg=SUBTLE, bg=BG, font=(FONT, 9), anchor="w",
                 justify="left", wraplength=px(300)).pack(fill="x", padx=pad, pady=(0, px(8)))
        for a in self.actions:
            self._button(self.body, a.label, lambda a=a: self.choose(a)).pack(
                fill="x", padx=pad, pady=px(2))
        tk.Label(self.body, text="Drop more files to add them  \u00b7  Esc closes",
                 fg=SUBTLE, bg=BG, font=(FONT, 8)).pack(pady=(px(6), px(10)))
        self._bind_dnd_tree()
        self._fit()
        self.top.focus_force()

    def choose(self, action: Action):
        if action.office_options:
            self.show_office_params(action)
        elif action.media_options:
            self.show_media_params(action)
        elif action.video_options:
            self.show_video_params(action)
        elif action.ask_text:
            self.show_text_params(action)
        elif action.ask:
            self.show_params(action)
        else:
            self.start(action, {})

    def _params_buttons(self, run) -> None:
        pad = px(16)
        row = tk.Frame(self.body, bg=BG)
        row.pack(fill="x", padx=pad, pady=(px(8), px(14)))
        self._button(row, "Back", lambda: self.show_menu(self.paths), center=True).pack(
            side="left", expand=True, fill="x", padx=(0, px(4)))
        self._button(row, "Run", run, primary=True, center=True).pack(
            side="left", expand=True, fill="x", padx=(px(4), 0))
        self.top.bind("<Return>", lambda e: run())
        self.top.bind("<KP_Enter>", lambda e: run())
        self._bind_dnd_tree()
        self._fit()

    def show_text_params(self, action: Action):
        key, caption, default, hint = action.ask_text
        self._clear()
        self.state = "params"
        pad = px(16)
        tk.Label(self.body, text=action.label, fg=FG, bg=BG, font=(FONT, 11, "bold"),
                 anchor="w").pack(fill="x", padx=pad, pady=(px(14), px(2)))
        tk.Label(self.body, text=caption, fg=SUBTLE, bg=BG, font=(FONT, 9),
                 anchor="w").pack(fill="x", padx=pad)
        var = tk.StringVar(value=str(default))
        entry = tk.Entry(self.body, textvariable=var, bg=SURFACE, fg=FG,
                         insertbackground=FG, relief="flat", font=(FONT, 12),
                         highlightthickness=0, bd=0)
        entry.pack(fill="x", padx=pad, pady=(px(8), px(2)), ipady=px(6))
        tk.Label(self.body, text=hint, fg=SUBTLE, bg=BG, font=(FONT, 8),
                 anchor="w", justify="left", wraplength=px(300)).pack(
            fill="x", padx=pad, pady=(0, px(4)))
        self._params_buttons(lambda: self.start(action, {key: var.get()}))
        entry.focus_set()
        entry.select_range(0, "end")
        entry.icursor("end")

    def show_office_params(self, action: Action):
        self._clear()
        self.state = "params"
        pad = px(16)
        docs = [Path(p) for p in self.paths
                if Path(p).is_file() and Path(p).suffix.lower() in OFFICE_INPUT_EXTS]
        order = ["pdf", "png", "jpg", "txt", "docx", "csv", "html", "json",
                 "pptx", "xlsx"]
        if office is not None and docs:
            available = [set(office.OfficeConverter.targets_for(p)) for p in docs]
            common = set.intersection(*available) if available else set()
            formats = [t for t in order if t in common]
        else:
            formats = ["pdf", "txt"]
        label_map = getattr(office, "TARGET_LABELS", {}) if office else {}
        labels = {label_map.get(fmt, fmt.upper()): fmt for fmt in formats}
        values = list(labels)
        tk.Label(self.body, text=action.label, fg=FG, bg=BG,
                 font=(FONT, 11, "bold"), anchor="w").pack(
            fill="x", padx=pad, pady=(px(14), px(2)))
        tk.Label(self.body, text="Output format", fg=SUBTLE, bg=BG,
                 font=(FONT, 9), anchor="w").pack(fill="x", padx=pad)
        default = "pdf" if "pdf" in formats else (formats[0] if formats else "")
        if default != "pdf" and "xlsx" in formats and docs and all(
                p.suffix.lower() in (".csv", ".json") for p in docs):
            default = "xlsx"
        var = tk.StringVar(value=next(
            (label for label, fmt in labels.items() if fmt == default), ""))
        combo = ttk.Combobox(self.body, textvariable=var, state="readonly",
                             values=values, width=34, font=(FONT, 9))
        combo.pack(fill="x", padx=pad, pady=(px(6), px(4)))
        tk.Label(self.body,
                 text="Office \u2192 PDF and slide rendering use LibreOffice for "
                      "layout fidelity. CSV/JSON/HTML/text exports and TXT "
                      "conversions use the built-in Python libraries.",
                 fg=SUBTLE, bg=BG, font=(FONT, 8), anchor="w", justify="left",
                 wraplength=px(300)).pack(fill="x", padx=pad, pady=(0, px(4)))
        if office is None:
            tk.Label(self.body,
                     text="office_converter.py is missing from the DropFlow folder.",
                     fg=RED, bg=BG, font=(FONT, 8), anchor="w",
                     wraplength=px(300)).pack(fill="x", padx=pad)

        def run():
            selected = labels.get(var.get())
            if selected:
                self.start(action, {"target": selected})

        self._params_buttons(run)
        combo.focus_set()

    def show_media_params(self, action: Action):
        self._clear()
        self.state = "params"
        pad = px(16)
        videos = _video_files(self.paths)
        all_mp4 = bool(videos) and all(p.suffix.lower() == ".mp4" for p in videos)
        if all_mp4:
            formats = ["mkv", "avi", "mov", "webm", "gif",
                       "mp3", "aac", "wav", "m4a", "flac"]
        else:
            formats = ["mp4", "mp3", "aac", "wav", "m4a", "flac"]
        labels = {}
        label_map = getattr(media, "FORMAT_LABELS", {}) if media else {}
        for fmt in formats:
            label = label_map.get(fmt, fmt.upper())
            labels[label] = fmt
        values = list(labels)
        tk.Label(self.body, text=action.label, fg=FG, bg=BG,
                 font=(FONT, 11, "bold"), anchor="w").pack(
            fill="x", padx=pad, pady=(px(14), px(2)))
        tk.Label(self.body, text="Output format", fg=SUBTLE, bg=BG,
                 font=(FONT, 9), anchor="w").pack(fill="x", padx=pad)
        var = tk.StringVar(value=values[0] if values else "")
        combo = ttk.Combobox(self.body, textvariable=var, state="readonly",
                             values=values, width=34, font=(FONT, 9))
        combo.pack(fill="x", padx=pad, pady=(px(6), px(4)))
        tk.Label(self.body,
                 text="Video conversions try a lossless stream copy first and "
                      "fall back to re-encoding when the container rejects the "
                      "source codecs. Audio targets strip the video track.",
                 fg=SUBTLE, bg=BG, font=(FONT, 8), anchor="w", justify="left",
                 wraplength=px(300)).pack(fill="x", padx=pad, pady=(0, px(4)))
        if media is None:
            tk.Label(self.body, text="media_converter.py is missing from the "
                                     "DropFlow folder.",
                     fg=RED, bg=BG, font=(FONT, 8), anchor="w",
                     wraplength=px(300)).pack(fill="x", padx=pad)

        def run():
            selected = labels.get(var.get())
            if selected:
                self.start(action, {"format": selected})

        self._params_buttons(run)
        combo.focus_set()

    def show_video_params(self, action: Action):
        self._clear()
        self.state = "params"
        pad = px(16)
        tk.Label(self.body, text=action.label, fg=FG, bg=BG, font=(FONT, 11, "bold"),
                 anchor="w").pack(fill="x", padx=pad, pady=(px(14), px(2)))
        tk.Label(self.body, text="Quality (CRF, lower = better and larger)",
                 fg=SUBTLE, bg=BG, font=(FONT, 9), anchor="w").pack(fill="x", padx=pad)
        crf_var = tk.IntVar(value=23)
        crf_label = tk.Label(self.body, text="23", fg=ACCENT, bg=BG,
                             font=(FONT, 14, "bold"))
        crf_label.pack(pady=(px(2), 0))
        crf_var.trace_add("write", lambda *a: crf_label.config(text=str(crf_var.get())))
        tk.Scale(self.body, from_=18, to=32, resolution=1, orient="horizontal",
                 variable=crf_var, showvalue=False, bg=BG, fg=FG, troughcolor=HOVER,
                 activebackground=ACCENT_HI, highlightthickness=0, bd=0,
                 sliderrelief="flat", length=px(260)).pack(padx=pad, pady=(0, px(4)))

        row = tk.Frame(self.body, bg=BG)
        row.pack(fill="x", padx=pad, pady=(px(2), 0))
        tk.Label(row, text="Resolution", fg=SUBTLE, bg=BG,
                 font=(FONT, 9)).pack(side="left")
        res_var = tk.StringVar(value="Keep original")
        res_box = ttk.Combobox(row, textvariable=res_var, state="readonly",
                               values=("Keep original", "2160p", "1440p",
                                       "1080p", "720p", "480p"),
                               width=16, font=(FONT, 9))
        res_box.pack(side="right")

        target_row = tk.Frame(self.body, bg=BG)
        target_row.pack(fill="x", padx=pad, pady=(px(8), 0))
        tk.Label(target_row, text="Target size (MB, optional)", fg=SUBTLE, bg=BG,
                 font=(FONT, 9)).pack(side="left")
        target_var = tk.StringVar(value="")
        tk.Entry(target_row, textvariable=target_var, width=10, bg=SURFACE, fg=FG,
                 insertbackground=FG, relief="flat", font=(FONT, 10),
                 highlightthickness=0, bd=0).pack(side="right", ipady=px(3))

        gpu_ok = _ffmpeg_has_encoder("h264_nvenc")
        gpu_var = tk.BooleanVar(value=False)
        gpu_text = "Use NVIDIA GPU (NVENC)" if gpu_ok else "GPU encoding unavailable"
        gpu_check = tk.Checkbutton(self.body, text=gpu_text, variable=gpu_var,
                                   fg=FG if gpu_ok else SUBTLE, bg=BG,
                                   selectcolor=HOVER, activebackground=BG,
                                   activeforeground=FG, font=(FONT, 9),
                                   highlightthickness=0, bd=0)
        gpu_check.pack(anchor="w", padx=pad, pady=(px(6), 0))
        if not gpu_ok:
            gpu_check.config(state="disabled")

        tk.Label(self.body,
                 text="Leave target empty to use CRF. A target size ignores CRF; "
                      "GPU encodes in one pass.",
                 fg=SUBTLE, bg=BG, font=(FONT, 8), anchor="w",
                 justify="left", wraplength=px(300)).pack(fill="x", padx=pad,
                                                          pady=(px(4), 0))
        err = tk.Label(self.body, text="", fg=RED, bg=BG, font=(FONT, 8),
                       anchor="w", justify="left", wraplength=px(300))
        err.pack(fill="x", padx=pad)

        def run():
            try:
                _parse_target_mb(target_var.get())
            except Exception as exc:
                err.config(text=str(exc))
                return
            self.start(action, {
                "quality": crf_var.get(),
                "resolution": res_var.get(),
                "target_mb": target_var.get().strip(),
                "gpu": bool(gpu_var.get()),
            })

        self._params_buttons(run)

    def show_params(self, action: Action):
        key, caption, lo, hi, default, unit, step = action.ask
        self._clear()
        self.state = "params"
        pad = px(16)
        tk.Label(self.body, text=action.label, fg=FG, bg=BG, font=(FONT, 11, "bold"),
                 anchor="w").pack(fill="x", padx=pad, pady=(px(14), px(2)))
        tk.Label(self.body, text=caption, fg=SUBTLE, bg=BG, font=(FONT, 9),
                 anchor="w").pack(fill="x", padx=pad)
        var = tk.IntVar(value=default)
        val = tk.Label(self.body, text=f"{default}{unit}", fg=ACCENT, bg=BG,
                       font=(FONT, 16, "bold"))
        val.pack(pady=(px(8), 0))
        var.trace_add("write", lambda *a: val.config(text=f"{var.get()}{unit}"))
        tk.Scale(self.body, from_=lo, to=hi, resolution=step, orient="horizontal",
                 variable=var, showvalue=False, bg=BG, fg=FG, troughcolor=HOVER,
                 activebackground=ACCENT_HI, highlightthickness=0, bd=0,
                 sliderrelief="flat", length=px(260)).pack(padx=pad, pady=px(4))
        self._params_buttons(lambda: self.start(action, {key: var.get()}))

    def show_working(self, title: str):
        self._clear()
        self.state = "working"
        pad = px(16)
        tk.Label(self.body, text=title, fg=FG, bg=BG, font=(FONT, 11, "bold"),
                 anchor="w").pack(fill="x", padx=pad, pady=(px(16), px(4)))
        self.status = tk.Label(self.body, text="Starting\u2026", fg=SUBTLE, bg=BG,
                               font=(FONT, 9), anchor="w", justify="left",
                               wraplength=px(290))
        self.status.pack(fill="x", padx=pad)
        pb = ttk.Progressbar(self.body, mode="indeterminate",
                             style="DF.Horizontal.TProgressbar", length=px(290))
        pb.pack(padx=pad, pady=(px(10), px(12)))
        pb.start(12)
        self.cancel_btn = self._button(self.body, "Cancel", self.request_cancel, center=True)
        self.cancel_btn.pack(fill="x", padx=pad, pady=(0, px(10)))
        tk.Label(self.body, text="Esc also cancels", fg=SUBTLE, bg=BG,
                 font=(FONT, 8)).pack(pady=(0, px(10)))
        self._bind_dnd_tree()
        self._fit()

    def show_done(self, res: Result):
        self._clear()
        self.state = "done"
        pad = px(16)
        title = ("!  Done with warnings" if res.failed else "\u2713  Done")
        title_fg = ACCENT if res.failed else GREEN
        tk.Label(self.body, text=title, fg=title_fg, bg=BG, font=(FONT, 12, "bold"),
                 anchor="w").pack(fill="x", padx=pad, pady=(px(14), px(2)))
        tk.Label(self.body, text=res.message, fg=FG, bg=BG, font=(FONT, 9), anchor="w",
                 justify="left", wraplength=px(300)).pack(fill="x", padx=pad)
        if res.failed:
            tk.Label(self.body, text="Some items failed. Open the log for details.",
                     fg=SUBTLE, bg=BG, font=(FONT, 8), anchor="w").pack(
                fill="x", padx=pad, pady=(px(2), 0))
        if res.preview:
            box_wrap = tk.Frame(self.body, bg=BG)
            box_wrap.pack(fill="x", padx=pad, pady=(px(8), 0))
            box = tk.Text(box_wrap, height=7, width=40, bg=SURFACE, fg=FG, relief="flat",
                          wrap="word", font=(FONT, 10), padx=8, pady=6,
                          highlightthickness=0)
            sb = tk.Scrollbar(box_wrap, command=box.yview, bg=HOVER, troughcolor=BG,
                              bd=0, highlightthickness=0)
            box.configure(yscrollcommand=sb.set)
            sb.pack(side="right", fill="y")
            box.pack(side="left", fill="both", expand=True)
            box.tag_configure("rtl", justify="right")
            for line in res.preview[:2000].splitlines():
                box.insert("end", line + "\n", "rtl" if ARABIC_RE.search(line) else "")
            box.config(state="disabled")
        row = tk.Frame(self.body, bg=BG)
        row.pack(fill="x", padx=pad, pady=(px(10), px(4)))
        if res.outputs:
            self._button(row, "Open folder", lambda: reveal(Path(res.outputs[0])),
                         primary=True, center=True).pack(side="left", expand=True,
                                                         fill="x", padx=(0, px(4)))
        self._button(row, "Close", self.close, center=True).pack(
            side="left", expand=True, fill="x", padx=(px(4), 0))
        if res.clipboard or res.failed:
            row2 = tk.Frame(self.body, bg=BG)
            row2.pack(fill="x", padx=pad, pady=(0, px(12)))
            if res.clipboard:
                copy_btn = self._button(row2, "Copy text", lambda: None, center=True)

                def do_copy():
                    self.app.root.clipboard_clear()
                    self.app.root.clipboard_append(res.clipboard)
                    copy_btn.config(text="Copied \u2713")
                copy_btn.bind("<Button-1>", lambda e: self.top.after(1, do_copy))
                copy_btn.pack(side="left", expand=True, fill="x", padx=(0, px(4)))
            if res.failed:
                self._button(row2, "Open log", lambda: open_path(log_path()),
                             center=True).pack(side="left", expand=True, fill="x",
                                               padx=(px(4), 0))
        self._bind_dnd_tree()
        self._fit()
        if not res.clipboard and not res.preview and not res.failed:
            self._timer = self.top.after(20000, self.close)

    def show_cancelled(self):
        self._clear()
        self.state = "cancelled"
        pad = px(16)
        tk.Label(self.body, text="Cancelled", fg=ACCENT, bg=BG, font=(FONT, 12, "bold"),
                 anchor="w").pack(fill="x", padx=pad, pady=(px(14), px(2)))
        tk.Label(self.body, text="The job was stopped. Incomplete output was cleaned up.",
                 fg=FG, bg=BG, font=(FONT, 9), anchor="w", justify="left",
                 wraplength=px(290)).pack(fill="x", padx=pad)
        row = tk.Frame(self.body, bg=BG)
        row.pack(fill="x", padx=pad, pady=(px(10), px(14)))
        self._button(row, "Back", lambda: self.show_menu(self.paths), center=True).pack(
            side="left", expand=True, fill="x", padx=(0, px(4)))
        self._button(row, "Close", self.close, center=True).pack(
            side="left", expand=True, fill="x", padx=(px(4), 0))
        self._bind_dnd_tree()
        self._fit()

    def show_error(self, message: str):
        self._clear()
        self.state = "error"
        self._last_error = message
        pad = px(16)
        tk.Label(self.body, text="Something went wrong", fg=RED, bg=BG,
                 font=(FONT, 12, "bold"), anchor="w").pack(fill="x", padx=pad,
                                                          pady=(px(14), px(2)))
        tk.Label(self.body, text=message[:600], fg=FG, bg=BG, font=(FONT, 9), anchor="w",
                 justify="left", wraplength=px(300)).pack(fill="x", padx=pad)
        tk.Label(self.body, text=f"Full details: {log_path().name}",
                 fg=SUBTLE, bg=BG, font=(FONT, 8), anchor="w").pack(
            fill="x", padx=pad, pady=(px(2), 0))
        row = tk.Frame(self.body, bg=BG)
        row.pack(fill="x", padx=pad, pady=(px(10), px(4)))
        self._button(row, "Back", lambda: self.show_menu(self.paths), center=True).pack(
            side="left", expand=True, fill="x", padx=(0, px(4)))
        self._button(row, "Close", self.close, center=True).pack(
            side="left", expand=True, fill="x", padx=(px(4), 0))
        row2 = tk.Frame(self.body, bg=BG)
        row2.pack(fill="x", padx=pad, pady=(0, px(12)))
        copy_btn = self._button(row2, "Copy details", lambda: None, center=True)

        def do_copy():
            details = (f"{message}\n\nLog: {log_path()}\n\n{log_tail()}")
            self.app.root.clipboard_clear()
            self.app.root.clipboard_append(details)
            copy_btn.config(text="Copied \u2713")

        copy_btn.bind("<Button-1>", lambda e: self.top.after(1, do_copy))
        copy_btn.pack(side="left", expand=True, fill="x", padx=(0, px(4)))
        self._button(row2, "Open log", lambda: open_path(log_path()),
                     center=True).pack(side="left", expand=True, fill="x",
                                       padx=(px(4), 0))
        self._bind_dnd_tree()
        self._fit()

    # ----- work ----------------------------------------------------------- #
    def request_cancel(self):
        if not self.busy or self._job is None:
            return
        self._job.cancel()
        self._close_password_dialog()
        if self.cancel_btn is not None:
            try:
                self.cancel_btn.config(text="Cancelling\u2026", fg=SUBTLE)
            except tk.TclError:
                pass
        if self.status is not None:
            try:
                self.status.config(text="Stopping the current job\u2026")
            except tk.TclError:
                pass

    def start(self, action: Action, opts: dict):
        self.busy = True
        self.show_working(action.label)
        q = self.app.q
        paths = list(self.paths)
        job = self._job = JobContext(
            lambda m: q.put(("progress", self, m)),
            self.app.cfg,
            lambda req: q.put(("password_request", self, req)))
        log.info("Job start: %s (%d path(s))", action.label, len(paths))

        def work():
            try:
                res = action.func(paths, opts, job)
                q.put(("done", self, res))
            except Cancelled:
                q.put(("cancelled", self))
            except Exception as exc:
                log.exception("Action failed: %s", action.label)
                q.put(("error", self, exc))

        threading.Thread(target=work, name="DropFlowJob", daemon=True).start()

    def on_progress(self, text: str):
        if self.status is not None:
            try:
                self.status.config(text=text)
            except tk.TclError:
                pass

    def on_done(self, res: Result):
        self.busy = False
        self._job = None
        log.info("Job done: %s", res.message)
        self.show_done(res)

    def on_cancelled(self):
        self.busy = False
        self._job = None
        log.info("Job cancelled")
        self.show_cancelled()

    def on_error(self, exc: Exception):
        self.busy = False
        self._job = None
        self.show_error(str(exc) or type(exc).__name__)


# --------------------------------------------------------------------------- #
#  UI: password prompt
# --------------------------------------------------------------------------- #
class PasswordDialog:
    """Small dark-themed modal prompt used by the archive worker."""

    def __init__(self, win: "DropWindow", request: PasswordRequest):
        self.win = win
        self.request = request
        self.alive = True
        t = self.top = tk.Toplevel(win.top)
        t.withdraw()
        t.overrideredirect(True)
        t.configure(bg=ACCENT)
        t.attributes("-topmost", True)
        t.protocol("WM_DELETE_WINDOW", self.cancel)
        body = tk.Frame(t, bg=BG)
        body.pack(fill="both", expand=True, padx=1, pady=1)
        pad = px(16)

        tk.Label(body, text="Password required", fg=FG, bg=BG,
                 font=(FONT, 12, "bold"), anchor="w").pack(
            fill="x", padx=pad, pady=(px(14), px(2)))
        name = request.archive
        if len(name) > 44:
            name = name[:41] + "\u2026"
        tk.Label(body, text=name, fg=ACCENT, bg=BG, font=(FONT, 9, "bold"),
                 anchor="w", justify="left", wraplength=px(300)).pack(
            fill="x", padx=pad)
        tk.Label(body, text=f"{request.kind} archive", fg=SUBTLE, bg=BG,
                 font=(FONT, 8), anchor="w").pack(fill="x", padx=pad)
        self.error = tk.Label(body, text=request.error, fg=RED, bg=BG,
                              font=(FONT, 9), anchor="w", justify="left",
                              wraplength=px(300))
        self.error.pack(fill="x", padx=pad, pady=(px(4), 0))
        self.var = tk.StringVar(value="")
        self.entry = tk.Entry(body, textvariable=self.var, show="\u2022",
                              bg=SURFACE, fg=FG, insertbackground=FG,
                              relief="flat", font=(FONT, 12),
                              highlightthickness=0, bd=0)
        self.entry.pack(fill="x", padx=pad, pady=(px(6), px(2)), ipady=px(6))
        self.show_var = tk.BooleanVar(value=False)
        tk.Checkbutton(body, text="Show password", variable=self.show_var,
                       command=self._toggle_show, fg=SUBTLE, bg=BG,
                       selectcolor=HOVER, activebackground=BG,
                       activeforeground=FG, font=(FONT, 8),
                       highlightthickness=0, bd=0).pack(anchor="w", padx=pad)
        row = tk.Frame(body, bg=BG)
        row.pack(fill="x", padx=pad, pady=(px(10), px(14)))
        self._button(row, "Cancel", self.cancel).pack(
            side="left", expand=True, fill="x", padx=(0, px(4)))
        self._button(row, "Unlock", self.submit, primary=True).pack(
            side="left", expand=True, fill="x", padx=(px(4), 0))

        t.bind("<Return>", lambda e: self.submit())
        t.bind("<KP_Enter>", lambda e: self.submit())
        t.bind("<Escape>", lambda e: self.cancel())
        self.entry.bind("<Return>", lambda e: self.submit())
        self.entry.bind("<KP_Enter>", lambda e: self.submit())

        t.update_idletasks()
        w, h = t.winfo_reqwidth(), t.winfo_reqheight()
        try:
            base_x = win.top.winfo_rootx() + max(0, (win.top.winfo_width() - w) // 2)
            base_y = win.top.winfo_rooty() + max(0, (win.top.winfo_height() - h) // 2)
        except tk.TclError:
            base_x, base_y = 200, 200
        vx, vy, vw, vh = virtual_screen_rect()
        x = max(vx + 8, min(base_x, vx + vw - w - 8))
        y = max(vy + 8, min(base_y, vy + vh - h - 48))
        t.geometry(f"{w}x{h}+{x}+{y}")
        t.deiconify()
        t.lift()
        try:
            t.grab_set()
        except tk.TclError:
            pass
        t.after(140, self._focus)
        t.after(80, lambda: self.alive and force_foreground(t, alt_trick=True))

    def _button(self, parent, text, command, primary=False):
        b = tk.Label(parent, text=text, bg=ACCENT if primary else SURFACE,
                     fg=BG if primary else FG, cursor="hand2",
                     font=(FONT, 10, "bold" if primary else "normal"),
                     padx=px(14), pady=px(7), anchor="center")
        b.bind("<Button-1>", lambda e: self.top.after(1, command))
        return b

    def _focus(self):
        if not self.alive:
            return
        try:
            self.entry.focus_set()
            self.entry.select_range(0, "end")
        except tk.TclError:
            pass

    def _toggle_show(self):
        try:
            self.entry.config(show="" if self.show_var.get() else "\u2022")
        except tk.TclError:
            pass

    def submit(self):
        if not self.alive:
            return
        password = self.var.get()
        if not password:
            self.error.config(text="Enter a password.")
            return
        self.request.password = password
        self.request.cancelled = False
        self._finish()

    def cancel(self):
        if not self.alive:
            return
        self.request.password = None
        self.request.cancelled = True
        self._finish()

    def _finish(self):
        if not self.alive:
            return
        self.alive = False
        try:
            self.top.grab_release()
        except tk.TclError:
            pass
        try:
            self.top.destroy()
        except tk.TclError:
            pass
        try:
            self.var.set("")
        except tk.TclError:
            pass
        if self.win.password_dialog is self:
            self.win.password_dialog = None
        self.request.event.set()


# --------------------------------------------------------------------------- #
#  UI: settings window
# --------------------------------------------------------------------------- #
class SettingsWindow:
    def __init__(self, app: "App"):
        self.app = app
        self.alive, self.recording = True, False
        self.pending = app.cfg["hotkey"]
        t = self.top = tk.Toplevel(app.root)
        t.title(f"{APP_NAME} Settings")
        t.configure(bg=BG)
        t.resizable(False, False)
        t.attributes("-topmost", True)
        t.protocol("WM_DELETE_WINDOW", self.cancel)
        t.bind("<Escape>", self._on_escape)
        try:
            from PIL import ImageTk
            self._icon_img = ImageTk.PhotoImage(make_tray_image())
            t.iconphoto(False, self._icon_img)
        except Exception:
            pass

        f = tk.Frame(t, bg=BG, padx=px(22), pady=px(18))
        f.pack()
        tk.Label(f, text="Global shortcut", fg=FG, bg=BG,
                 font=(FONT, 11, "bold")).grid(row=0, column=0, columnspan=2, sticky="w")
        self.var = tk.StringVar(value=pretty_hotkey(self.pending))
        tk.Entry(f, textvariable=self.var, state="readonly", readonlybackground=HOVER,
                 fg=FG, relief="flat", font=(FONT, 12), width=22,
                 justify="center").grid(row=1, column=0, sticky="we", ipady=px(6),
                                        pady=(px(8), px(2)))
        self.rec_btn = self._btn(f, "Record", self.record)
        self.rec_btn.grid(row=1, column=1, padx=(px(8), 0), pady=(px(8), px(2)))
        tk.Label(f, text="Click Record, then press the new combination "
                         "(e.g. Ctrl + Alt + Space). Esc cancels recording.",
                 fg=SUBTLE, bg=BG, font=(FONT, 9), wraplength=px(350),
                 justify="left").grid(row=2, column=0, columnspan=2, sticky="w",
                                      pady=(px(2), 0))
        self.active = tk.Label(f, text="", fg=GREEN, bg=BG, font=(FONT, 9),
                               wraplength=px(350), justify="left")
        self.active.grid(row=3, column=0, columnspan=2, sticky="w")
        self.err = tk.Label(f, text="", fg=RED, bg=BG, font=(FONT, 9),
                            wraplength=px(350), justify="left")
        self.err.grid(row=4, column=0, columnspan=2, sticky="w")
        self.auto = tk.BooleanVar(value=bool(app.cfg.get("autostart")))
        tk.Checkbutton(f, text="Start DropFlow with Windows", variable=self.auto, fg=FG,
                       bg=BG, selectcolor=HOVER, activebackground=BG,
                       activeforeground=FG, font=(FONT, 10), highlightthickness=0,
                       bd=0).grid(row=5, column=0, columnspan=2, sticky="w",
                                  pady=(px(12), px(8)))
        tk.Frame(f, bg=HOVER, height=1).grid(row=6, column=0, columnspan=2,
                                             sticky="we", pady=px(8))
        tk.Label(f, text="Detected helper programs", fg=FG, bg=BG,
                 font=(FONT, 10, "bold")).grid(row=7, column=0, columnspan=2, sticky="w",
                                               pady=(0, px(4)))
        self.tool_labels = {}
        for i, key in enumerate(("soffice", "ffmpeg", "tesseract", "rar"), start=8):
            lab = tk.Label(f, text="Checking\u2026", fg=SUBTLE, bg=BG,
                           font=(FONT, 9), justify="left", wraplength=px(350),
                           anchor="w")
            lab.grid(row=i, column=0, columnspan=2, sticky="w")
            self.tool_labels[key] = lab

        tk.Label(f, text="OCR languages (Tesseract codes, combine with +)", fg=FG, bg=BG,
                 font=(FONT, 10, "bold")).grid(row=12, column=0, columnspan=2,
                                               sticky="w", pady=(px(10), px(2)))
        self.ocr_var = tk.StringVar(value=app.cfg.get("ocr_langs", "ara+eng"))
        self.ocr_box = ttk.Combobox(f, textvariable=self.ocr_var, width=30,
                                   font=(FONT, 9))
        self.ocr_box.grid(row=13, column=0, columnspan=2, sticky="we")
        self.ocr_hint = tk.Label(
            f, text="Examples: ara+eng, eng, ara, fra. Installed languages appear here.",
            fg=SUBTLE, bg=BG, font=(FONT, 8), justify="left", wraplength=px(350),
            anchor="w")
        self.ocr_hint.grid(row=14, column=0, columnspan=2, sticky="w")

        row = tk.Frame(f, bg=BG)
        row.grid(row=15, column=0, columnspan=2, sticky="e", pady=(px(14), 0))
        self._btn(row, "Reset default", self.reset).pack(side="left", padx=px(4))
        self._btn(row, "Open log", self.open_log).pack(side="left", padx=px(4))
        self._btn(row, "Cancel", self.cancel).pack(side="left", padx=px(4))
        self._btn(row, "Save", self.save, primary=True).pack(side="left", padx=(px(4), 0))

        t.update_idletasks()
        w, h = t.winfo_reqwidth(), t.winfo_reqheight()
        sw, sh = t.winfo_screenwidth(), t.winfo_screenheight()
        x = max(8, (sw - w) // 2)
        y = max(8, min((sh - h) // 2, sh - h - 48))
        t.geometry(f"+{x}+{y}")
        dark_titlebar(t)
        self._update_active()
        t.after(30, self._refresh_tools)
        t.after(80, lambda: force_foreground(t, alt_trick=False))

    def _btn(self, parent, text, command, primary=False):
        return tk.Button(parent, text=text, command=command, relief="flat", bd=0,
                         cursor="hand2", bg=ACCENT if primary else HOVER,
                         fg=BG if primary else FG,
                         activebackground=ACCENT_HI if primary else SURFACE,
                         activeforeground=BG if primary else ACCENT,
                         font=(FONT, 10, "bold" if primary else "normal"),
                         padx=px(14), pady=px(6))

    def _update_active(self):
        if not self.alive:
            return
        if getattr(self.app, "hotkey_ok", False):
            self.active.config(text=f"Active now: {pretty_hotkey(self.app.cfg['hotkey'])}",
                               fg=GREEN)
        else:
            err = getattr(self.app, "hotkey_error", "")
            text = "Active shortcut: none"
            if err:
                text += f" \u2014 {err}"
            self.active.config(text=text, fg=RED)

    def _validate_ocr_langs(self) -> tuple[str | None, str]:
        text = self.ocr_var.get().strip()
        parts = [p.lower() for p in re.split(r"[+,;\s]+", text) if p]
        if not parts:
            return None, "Enter at least one Tesseract language code."
        available = set(_available_ocr_languages())
        if available:
            missing = [p for p in parts if p not in available]
            if missing:
                shown = ", ".join(sorted(available)[:18])
                return None, (f"Not installed: {', '.join(missing)}.\n"
                              f"Installed (first 18): {shown}")
        return "+".join(parts), ""

    def _refresh_tools(self):
        if not self.alive:
            return
        soffice = _find_soffice()
        ffmpeg = _find_ffmpeg()
        tesseract = _find_tesseract()
        rar_tool = _find_rar_tool()
        entries = [
            ("soffice", soffice, "LibreOffice \u2014 ready (best Office \u2192 PDF)",
             "LibreOffice \u2014 not found (using the built-in DOCX renderer)"),
            ("ffmpeg", ffmpeg, "FFmpeg \u2014 ready (video compression + audio extraction)",
             "FFmpeg \u2014 not found (video tools disabled)"),
            ("tesseract", tesseract, "Tesseract OCR \u2014 ready",
             "Tesseract OCR \u2014 not found (OCR disabled)"),
            ("rar", rar_tool,
             f"RAR \u2014 ready ({rar_tool[0] if rar_tool else 'external'})",
             "RAR \u2014 no backend (install 7-Zip or UnRAR)"),
        ]
        for key, value, ok_text, missing_text in entries:
            lab = self.tool_labels.get(key)
            if lab is None:
                continue
            try:
                lab.config(text=ok_text if value else missing_text,
                           fg=GREEN if value else SUBTLE)
            except tk.TclError:
                pass
        langs = _available_ocr_languages()
        try:
            self.ocr_box.config(values=("ara+eng", *langs))
            if langs:
                self.ocr_hint.config(
                    text=f"Installed: {', '.join(langs[:16])}"
                         + (" \u2026" if len(langs) > 16 else "")
                         + "  \u00b7  combine with +")
            else:
                self.ocr_hint.config(text="Tesseract is not available, so OCR will be disabled.")
        except tk.TclError:
            pass

    def _on_escape(self, _event=None):
        if self.recording:
            self._stop_recording(restore=True)
            self.var.set(pretty_hotkey(self.pending))
            self.err.config(text="", fg=RED)
            return "break"
        self.cancel()
        return "break"

    def record(self):
        if self.recording:
            return
        self.recording = True
        self.err.config(text="", fg=RED)
        self.var.set("Press your shortcut\u2026")
        self.rec_btn.config(state="disabled", text="Listening\u2026")
        self.active.config(text="Recording \u2014 the current shortcut is paused.", fg=ACCENT)
        self.app.hotkey.unregister()  # so the old shortcut doesn't pop the zone mid-recording
        self.top.focus_force()
        self.top.bind("<KeyPress>", self._capture_key)

    def _capture_key(self, event):
        if not self.recording:
            return None
        keysym = event.keysym
        if keysym in _MODIFIER_KEYSYMS:
            mods = _mods_down()
            self.var.set(pretty_hotkey("+".join(mods)) if mods else "Press your shortcut\u2026")
            return "break"
        mods = _mods_down()
        token = _keysym_token(keysym, event)
        if not token:
            self.err.config(text="That key can't be used. Try a letter, number or F-key.",
                            fg=RED)
            return "break"
        if not mods:
            self.err.config(text="Add at least one modifier (Ctrl / Alt / Shift / Win).",
                            fg=RED)
            return "break"
        hk = "+".join(mods + [token])
        _mods, _vk, err = parse_hotkey(hk)
        if err:
            self.err.config(text=err, fg=RED)
            return "break"
        # Test the shortcut for real before accepting it.
        ok, msg = self.app.hotkey.register(hk)
        if not ok:
            self.err.config(text=f"That shortcut is unavailable:\n{msg}", fg=RED)
            return "break"
        # Restore the old shortcut until Save is pressed.
        self.app.hotkey.unregister()
        restored, restore_err = self.app.hotkey.register(self.app.cfg["hotkey"])
        self.app.hotkey_ok = bool(restored)
        self.app.hotkey_error = "" if restored else restore_err
        self.pending = hk
        self._stop_recording(restore=False)
        self.var.set(pretty_hotkey(hk))
        self.err.config(text="Shortcut is available. Press Save to apply.", fg=GREEN)
        self._update_active()
        return "break"

    def _stop_recording(self, restore: bool):
        if not self.recording:
            return
        self.recording = False
        try:
            self.top.unbind("<KeyPress>")
        except tk.TclError:
            pass
        self.rec_btn.config(state="normal", text="Record")
        if restore:
            ok, _err = self.app.hotkey.register(self.app.cfg["hotkey"])
            self.app.hotkey_ok = bool(ok)
            if ok:
                self.app.hotkey_error = ""
        self._update_active()

    def reset(self):
        if self.recording:
            return
        self.pending = DEFAULT_HOTKEY
        self.var.set(pretty_hotkey(self.pending))
        self.err.config(text="", fg=RED)

    def open_log(self):
        try:
            open_path(log_path())
        except Exception:
            pass

    def save(self):
        if self.recording:
            self.err.config(text="Finish recording first (Esc cancels recording).",
                            fg=RED)
            return
        langs, lang_err = self._validate_ocr_langs()
        if lang_err:
            self.err.config(text=lang_err, fg=RED)
            return
        ok, msg = self.app.apply_settings(self.pending, self.auto.get(), langs)
        if ok:
            self.close(resume=False)
        else:
            self.err.config(text=msg, fg=RED)
            self._update_active()

    def cancel(self):
        self._stop_recording(restore=False)
        self.close(resume=True)

    def close(self, resume: bool):
        if not self.alive:
            return
        self.alive = False
        try:
            self.top.destroy()
        except tk.TclError:
            pass
        self.app.settings = None
        if resume:
            ok, _err = self.app.hotkey.register(self.app.cfg["hotkey"])
            self.app.hotkey_ok = bool(ok)
            if ok:
                self.app.hotkey_error = ""


# --------------------------------------------------------------------------- #
#  App: tray, hotkey, event loop
# --------------------------------------------------------------------------- #
def make_tray_image() -> Image.Image:
    s = 64
    img = Image.new("RGBA", (s, s), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    d.rounded_rectangle((2, 2, s - 3, s - 3), radius=14, fill=BG, outline=ACCENT, width=3)
    d.line((32, 13, 32, 34), fill=ACCENT, width=5)
    d.polygon([(21, 29), (43, 29), (32, 42)], fill=ACCENT)
    d.line((17, 40, 17, 52, 47, 52, 47, 40), fill=ACCENT, width=4)
    return img


class App:
    def __init__(self):
        global SCALE
        self.cfg = load_config()
        self.q: queue.Queue = queue.Queue()
        self.win: DropWindow | None = None
        self.settings: SettingsWindow | None = None
        self._last_toggle = 0.0
        self._quitting = False
        self.hotkey_ok = False
        self.hotkey_error = ""

        self.root = TkinterDnD.Tk()
        self.root.withdraw()
        self.root.title(APP_NAME)
        self.root.report_callback_exception = lambda *a: log.error("Tk callback error",
                                                                  exc_info=a)
        apply_ui_scale(self.root, primary_dpi() / 96.0)

        style = ttk.Style(self.root)
        try:
            style.theme_use("clam")
        except Exception:
            pass
        style.configure("DF.Horizontal.TProgressbar", troughcolor=HOVER, background=ACCENT,
                        bordercolor=BG, lightcolor=ACCENT, darkcolor=ACCENT,
                        thickness=px(6))

        self.icon = pystray.Icon(APP_NAME, make_tray_image(), APP_NAME, pystray.Menu(
            pystray.MenuItem("Open drop zone", lambda i, it: self.q.put(("toggle",)),
                             default=True),
            pystray.MenuItem("Settings", lambda i, it: self.q.put(("settings",))),
            pystray.MenuItem("Open log", lambda i, it: self.q.put(("log",))),
            pystray.Menu.SEPARATOR,
            pystray.MenuItem("Quit", lambda i, it: self.q.put(("quit",))),
        ))

        self.hotkey = HotkeyManager(lambda: self.q.put(("toggle",)))
        self.hotkey.start()
        ok, err = self.hotkey.register(self.cfg["hotkey"])
        if not ok:
            ok2, _err2 = self.hotkey.register(DEFAULT_HOTKEY)
            if ok2:
                self.cfg["hotkey"] = DEFAULT_HOTKEY
                save_config(self.cfg)
                ok = True
            else:
                self.hotkey_error = err
        self.hotkey_ok = ok
        if ok:
            self.hotkey_error = ""
        self.root.after(60, self._poll)
        self.root.after(1400, self._hello)

    # ----- hotkey --------------------------------------------------------- #
    def apply_settings(self, hotkey: str, autostart: bool,
                       ocr_langs: str | None = None) -> tuple[bool, str]:
        ok, err = self.hotkey.register(hotkey)
        if not ok:
            restored, restore_err = self.hotkey.register(self.cfg["hotkey"])
            self.hotkey_ok = bool(restored)
            self.hotkey_error = "" if restored else restore_err
            return False, f"That shortcut can't be used:\n{err}"
        self.cfg["hotkey"] = hotkey
        self.hotkey_ok = True
        self.hotkey_error = ""
        if ocr_langs:
            self.cfg["ocr_langs"] = ocr_langs
        auto_error = ""
        try:
            set_autostart(autostart)
        except Exception as exc:
            auto_error = str(exc)
            log.exception("Autostart change failed")
        else:
            self.cfg["autostart"] = autostart
        save_config(self.cfg)
        if auto_error:
            self.notify("Shortcut saved, but the startup setting failed.")
            return False, ("The shortcut is active, but Windows startup could not be "
                           f"changed:\n{auto_error}")
        self.notify(f"Shortcut set to {pretty_hotkey(hotkey)}")
        return True, ""
    # ----- plumbing ------------------------------------------------------- #
    def notify(self, message: str):
        try:
            self.icon.notify(message[:220], APP_NAME)
        except Exception:
            log.info("Notify: %s", message)

    def _hello(self):
        if self.hotkey_error:
            self.notify("DropFlow is running in the tray, but the shortcut could not be "
                        f"registered: {self.hotkey_error} Open Settings to pick another.")
        else:
            self.notify(f"DropFlow is running in the tray. Press "
                        f"{pretty_hotkey(self.cfg['hotkey'])} to open the drop zone.")

    def _poll(self):
        try:
            while True:
                msg = self.q.get_nowait()
                try:
                    self._dispatch(msg)
                except Exception:
                    log.exception("Dispatch failed for %r", msg[0])
        except queue.Empty:
            pass
        try:
            self.root.after(60, self._poll)
        except tk.TclError:
            pass

    def _dispatch(self, msg: tuple):
        kind = msg[0]
        if kind == "toggle":
            self._toggle()
        elif kind == "settings":
            self._open_settings()
        elif kind == "log":
            open_path(DATA_DIR / "dropflow.log")
        elif kind == "quit":
            self.quit()
        elif kind == "progress":
            if msg[1].alive:
                msg[1].on_progress(msg[2])
        elif kind == "password_request":
            win, request = msg[1], msg[2]
            if win.alive and not win.closing and win.busy:
                win.show_password_dialog(request)
            else:
                request.cancelled = True
                request.event.set()
        elif kind == "done":
            win, res = msg[1], msg[2]
            if win.alive and not win.closing:
                win.on_done(res)
            else:
                self.notify(res.message)
        elif kind == "cancelled":
            win = msg[1]
            if win.alive and not win.closing:
                win.on_cancelled()
        elif kind == "error":
            win, exc = msg[1], msg[2]
            if win.alive and not win.closing:
                win.on_error(exc)
            else:
                self.notify(f"Failed: {exc}")

    def _toggle(self):
        if self._quitting:
            return
        now = time.monotonic()
        if now - self._last_toggle < 0.35:  # swallow auto-repeat
            return
        self._last_toggle = now
        if self.win and self.win.alive and not self.win.closing:
            self.win.close()
            return
        x, y = self.root.winfo_pointerxy()
        apply_ui_scale(self.root, monitor_dpi_at(x, y) / 96.0)
        self.win = DropWindow(self, x, y)

    def _open_settings(self):
        if self._quitting:
            return
        if self.settings and self.settings.alive:
            try:
                self.settings.top.lift()
                force_foreground(self.settings.top, alt_trick=False)
            except tk.TclError:
                pass
            return
        apply_ui_scale(self.root, primary_dpi() / 96.0)
        self.settings = SettingsWindow(self)

    def quit(self):
        if self._quitting:
            return
        self._quitting = True
        if self.win is not None:
            try:
                self.win.shutdown()
            except Exception:
                pass
        if self.settings is not None:
            try:
                self.settings.close(resume=False)
            except Exception:
                pass
        try:
            self.hotkey.stop()
        except Exception:
            pass
        try:
            self.icon.stop()
        except Exception:
            pass
        try:
            self.root.destroy()
        except Exception:
            pass

    def run(self):
        self.icon.run_detached()
        self.root.mainloop()


def main():
    if "--selftest" in sys.argv:
        raise SystemExit(_selftest())
    _relaunch_without_console()
    _ensure_single_instance()
    _setup_logging()
    log.info("DropFlow starting (pid=%s, python=%s)", os.getpid(), sys.executable)
    try:
        App().run()
    except SystemExit:
        raise
    except Exception:
        log.exception("DropFlow crashed")
        _msgbox(f"DropFlow crashed.\n\nSee {DATA_DIR / 'dropflow.log'} for details.",
                flags=0x10)
        raise


def _selftest() -> int:
    """Headless diagnostics used to validate a packaged exe (--selftest)."""
    import traceback

    lines: list[str] = []
    failed = False

    def item(name: str, ok: bool, detail: str = "") -> None:
        nonlocal failed
        lines.append(("PASS " if ok else "FAIL ") + name + (f": {detail}" if detail else ""))
        if not ok:
            failed = True

    tmp = None
    try:
        tmp = _new_temp_dir("dropflow_selftest_")
        item("Python", True, sys.version.split()[0])
        item("Frozen", True, str(getattr(sys, "frozen", False)))

        # Imports that are easy to lose during packaging.
        for mod in ("pystray", "pymupdf", "reportlab", "docx", "pptx", "openpyxl",
                    "arabic_reshaper", "bidi", "pytesseract", "imageio_ffmpeg"):
            try:
                importlib.import_module(mod)
                item(f"import {mod}", True)
            except Exception as exc:
                item(f"import {mod}", False, repr(exc))

        jpg = tmp / "selftest.jpg"
        png = tmp / "selftest.png"
        try:
            from PIL import Image as PILImage, ImageDraw as PILDraw
            im = PILImage.new("RGB", (320, 240), (30, 70, 120))
            PILDraw.Draw(im).text((20, 100), "DropFlow selftest", fill=(255, 255, 255))
            im.save(jpg, "JPEG", quality=92)
            im.close()
            im = PILImage.new("RGBA", (220, 220), (0, 0, 0, 0))
            PILDraw.Draw(im).ellipse((10, 10, 210, 210), fill=(203, 166, 247, 180))
            im.save(png)
            im.close()
            item("Pillow JPEG/PNG", True)
        except Exception as exc:
            item("Pillow JPEG/PNG", False, repr(exc))

        # Tk + tkdnd (the files that are trickiest to bundle).
        try:
            root = TkinterDnD.Tk()
            root.withdraw()
            root.update()
            top = tk.Toplevel(root)
            top.withdraw()
            top.drop_target_register(DND_FILES)
            top.dnd_bind("<<Drop>>", lambda event: event.action)
            root.update()
            root.destroy()
            item("Tk + tkinterdnd2", True)
        except Exception as exc:
            item("Tk + tkinterdnd2", False, repr(exc))

        ctx = JobContext(lambda _msg: None)
        try:
            res = act_images_to_pdf([jpg, png], {}, ctx)
            doc = fitz.open(str(res.outputs[0]))
            pages = doc.page_count
            doc.close()
            item("images -> PDF", pages == 2, f"pages={pages}")
        except Exception as exc:
            item("images -> PDF", False, repr(exc))

        try:
            doc = fitz.open()
            page = doc.new_page(width=595, height=842)
            page.insert_text((72, 72), "DropFlow selftest PDF", fontsize=20)
            pdf = tmp / "selftest.pdf"
            doc.save(str(pdf))
            doc.close()
            res = make_pdf_to_images("PNG")([pdf], {"dpi": 72}, ctx)
            count = len(list(Path(res.outputs[0]).glob("*.png")))
            item("PDF -> PNG", count == 1, f"pages={count}")
        except Exception as exc:
            item("PDF -> PNG", False, repr(exc))

        try:
            from docx import Document
            doc = Document()
            doc.add_paragraph("DropFlow fallback renderer selftest")
            docx_path = tmp / "selftest.docx"
            doc.save(str(docx_path))
            pdf = tmp / "selftest_docx.pdf"
            _render_docx_to_pdf(docx_path, pdf)
            d = fitz.open(str(pdf))
            pages = d.page_count
            d.close()
            item("DOCX fallback renderer", pages >= 1, f"pages={pages}")
        except Exception as exc:
            item("DOCX fallback renderer", False, repr(exc))

        try:
            exe = _find_ffmpeg()
            item("FFmpeg detected", bool(exe), exe or "not found")
            if exe:
                res = _run_hidden([exe, "-version"], timeout=20)
                item("FFmpeg runs", res.returncode == 0,
                     _tail((res.stdout or "") + (res.stderr or ""), 100))
        except Exception as exc:
            item("FFmpeg detected", False, repr(exc))

        try:
            import imageio_ffmpeg
            bundled = imageio_ffmpeg.get_ffmpeg_exe()
            item("Bundled FFmpeg", Path(bundled).is_file(), bundled)
            if Path(bundled).is_file():
                res = _run_hidden([bundled, "-version"], timeout=20)
                item("Bundled FFmpeg runs", res.returncode == 0)
        except Exception as exc:
            item("Bundled FFmpeg", False, repr(exc))

        for name, finder in (("LibreOffice", _find_soffice), ("Tesseract", _find_tesseract)):
            try:
                value = finder()
                lines.append(f"INFO {name}: {value or 'not found'}")
            except Exception as exc:
                lines.append(f"INFO {name}: error {exc!r}")
        try:
            tool = _find_rar_tool()
            lines.append(f"INFO RAR backend: "
                         f"{tool[0] + ' ' + tool[1] if tool else 'not found (optional)'}")
        except Exception as exc:
            lines.append(f"INFO RAR backend: error {exc!r}")
        if media is None:
            lines.append("INFO Media converter: unavailable"
                         + (f" ({_MEDIA_IMPORT_ERROR})" if _MEDIA_IMPORT_ERROR else ""))
        else:
            lines.append("INFO Media converter: ready")
        if office is None:
            lines.append("INFO Office converter: unavailable"
                         + (f" ({_OFFICE_IMPORT_ERROR})" if _OFFICE_IMPORT_ERROR else ""))
        else:
            lines.append("INFO Office converter: ready")
    except Exception:
        item("selftest crash", False, traceback.format_exc())
    finally:
        if tmp is not None:
            discard([tmp])

    report = "\n".join(lines)
    target = Path(tempfile.gettempdir()) / "DropFlow_selftest.txt"
    try:
        target.write_text(report + "\n", encoding="utf-8")
    except Exception:
        pass
    if sys.stdout is not None:
        try:
            print(report)
        except Exception:
            pass
    return 0 if not failed else 1


if __name__ == "__main__":
    main()
