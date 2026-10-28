"""
office_converter.py - document conversion backend for DropFlow.

Supported directions
--------------------
PowerPoint (.pptx / .ppt):
    -> PDF, PNG/JPG per slide, TXT

Excel (.xlsx / .xls):
    -> PDF, CSV (one file per sheet), HTML, JSON, TXT

Text (.txt):
    -> PDF, DOCX, PPTX, XLSX

Other inputs:
    images / PDF     -> PPTX
    CSV / JSON / PDF -> XLSX
    PDF / DOCX / PPTX / XLSX / CSV / JSON / images -> TXT

High-fidelity Office -> PDF uses LibreOffice headless first, then optional
Microsoft Office COM automation.  If neither is available, Excel and DOCX
PDF output falls back to a basic reportlab renderer; PowerPoint -> PDF
requires LibreOffice or Office COM.

Dependencies
------------
Required for the corresponding formats (all are imported lazily):
    python-pptx   PowerPoint read/write
    openpyxl      Excel read/write
    python-docx   Word read/write
    reportlab     PDF/text rendering
    PyMuPDF       PDF text, tables and page images
    Pillow        JPEG conversion and image sizing
Optional:
    pywin32       Microsoft Office COM fallback (Windows)
    pytesseract   OCR when extracting text from images
    LibreOffice   high-fidelity Office -> PDF and legacy .ppt/.xls handling

The module has no Tk dependency.  Callbacks run in the worker thread; a GUI
should marshal them to its event loop (DropFlow does this through JobContext).
"""
from __future__ import annotations

import csv
import html as html_module
import importlib
import importlib.util
import json
import logging
import os
import re
import shutil
import subprocess
import tempfile
import threading
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, Iterable

LOG = logging.getLogger("DropFlow.OfficeConverter")
CREATE_NO_WINDOW = getattr(subprocess, "CREATE_NO_WINDOW", 0)

PPT_EXTS = {".pptx", ".ppt"}
EXCEL_EXTS = {".xlsx", ".xls"}
TEXT_EXTS = {".txt"}
WORD_EXTS = {".docx"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff", ".gif"}
PDF_EXTS = {".pdf"}
DATA_EXTS = {".csv", ".json"}
DOC_INPUT_EXTS = PPT_EXTS | EXCEL_EXTS | TEXT_EXTS | WORD_EXTS | DATA_EXTS
ALL_INPUT_EXTS = DOC_INPUT_EXTS | IMAGE_EXTS | PDF_EXTS

FILE_TARGETS = {"pdf", "txt", "docx", "pptx", "xlsx", "html", "json"}
DIR_TARGETS = {"png", "jpg", "csv"}
SUPPORTED_TARGETS = FILE_TARGETS | DIR_TARGETS

TARGET_LABELS = {
    "pdf": "PDF",
    "png": "PNG images (one per slide/page)",
    "jpg": "JPG images (one per slide/page)",
    "txt": "Plain text",
    "csv": "CSV files",
    "html": "HTML",
    "json": "JSON",
    "docx": "Word document",
    "pptx": "PowerPoint presentation",
    "xlsx": "Excel workbook",
}


class OfficeConversionError(RuntimeError):
    """Any validation or conversion failure."""


class OfficeDependencyError(OfficeConversionError):
    """A required external tool or Python package is missing."""


class OfficeCancelled(OfficeConversionError):
    """The conversion was cancelled through cancel_event."""


@dataclass
class OfficeResult:
    input: Path
    outputs: list[Path]
    target: str
    method: str
    elapsed: float


@dataclass
class OfficeCallbacks:
    on_start: Callable[[Path, str], None] | None = None
    on_progress: Callable[[float, str], None] | None = None
    on_success: Callable[[OfficeResult], None] | None = None
    on_error: Callable[[Exception], None] | None = None


class OfficeHandle:
    """Handle returned by OfficeConverter.convert_async()."""

    def __init__(self, converter: "OfficeConverter", input_path: str | Path,
                 target: str, output_path: str | Path | None,
                 overwrite: bool, callbacks: OfficeCallbacks):
        self.converter = converter
        self.input_path = input_path
        self.target = target
        self.output_path = output_path
        self.overwrite = overwrite
        self.callbacks = callbacks
        self.cancel_event = threading.Event()
        self.result: OfficeResult | None = None
        self.error: Exception | None = None
        self._thread = threading.Thread(target=self._run,
                                        name="OfficeConversion", daemon=True)

    def start(self):
        self._thread.start()
        return self

    def _run(self):
        try:
            self.result = self.converter.convert(
                self.input_path, self.target,
                output_path=self.output_path,
                overwrite=self.overwrite,
                callbacks=self.callbacks,
                cancel_event=self.cancel_event)
        except Exception as exc:  # noqa: BLE001
            self.error = exc

    def cancel(self):
        self.cancel_event.set()

    def join(self, timeout: float | None = None):
        self._thread.join(timeout)
        return self.result

    def is_alive(self) -> bool:
        return self._thread.is_alive()


class OfficeConverter:
    """Reusable converter for PowerPoint, Excel, Word, text and data files."""

    def __init__(self, soffice_path: str | Path | None = None, dpi: int = 150,
                 prefer_com: bool = False, ocr_lang: str = "eng"):
        self.soffice = self._find_soffice(soffice_path)
        self.dpi = max(72, min(400, int(dpi)))
        self.prefer_com = prefer_com
        self.ocr_lang = ocr_lang
        self._temps: list[Path] = []

    def _tempdir(self, prefix: str) -> Path:
        path = Path(tempfile.mkdtemp(prefix=prefix))
        self._temps.append(path)
        return path

    def cleanup(self) -> None:
        for path in self._temps:
            shutil.rmtree(path, ignore_errors=True)
        self._temps.clear()


    # ------------------------------------------------------------------ #
    #  Dependency / process helpers
    # ------------------------------------------------------------------ #
    @staticmethod
    def _find_soffice(provided: str | Path | None) -> str | None:
        candidates: list[str] = []
        if provided:
            candidates.append(str(provided))
        if os.environ.get("SOFFICE_PATH"):
            candidates.append(os.environ["SOFFICE_PATH"])
        for name in ("soffice", "soffice.com"):
            found = shutil.which(name)
            if found:
                candidates.append(found)
        candidates += [
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
            str(Path(os.environ.get("LOCALAPPDATA", "")) /
                "Programs" / "LibreOffice" / "program" / "soffice.exe"),
        ]
        for candidate in candidates:
            try:
                path = Path(candidate)
                if path.is_file():
                    return str(path)
            except Exception:
                pass
        return None

    @staticmethod
    def _require(module: str, package: str):
        try:
            return importlib.import_module(module)
        except ImportError as exc:
            raise OfficeDependencyError(
                f"{package} is required for this conversion.\n\n"
                f"Install it with: pip install {package}") from exc

    def _require_soffice(self) -> str:
        if not self.soffice:
            raise OfficeDependencyError(
                "LibreOffice was not found.\n\n"
                "Install LibreOffice for high-quality Office/PDF conversion, "
                "or install Microsoft Office and pywin32 for the COM fallback.")
        return self.soffice

    def _check(self, cancel_event: threading.Event | None) -> None:
        if cancel_event is not None and cancel_event.is_set():
            raise OfficeCancelled("Conversion cancelled.")

    def _progress(self, callbacks: OfficeCallbacks, pct: float, message: str) -> None:
        if callbacks.on_progress is None:
            return
        try:
            callbacks.on_progress(max(0.0, min(100.0, float(pct))), message)
        except Exception:
            LOG.exception("Office conversion progress callback failed")

    def _callback(self, callback, *args) -> None:
        if callback is None:
            return
        try:
            callback(*args)
        except Exception:
            LOG.exception("Office conversion callback failed")

    def _run_process(self, cmd: list[str], *, timeout: float = 240,
                     cancel_event: threading.Event | None = None) -> tuple[int, str]:
        proc = subprocess.Popen(
            cmd, stdin=subprocess.DEVNULL, stdout=subprocess.PIPE,
            stderr=subprocess.PIPE, text=True, encoding="utf-8",
            errors="replace", creationflags=CREATE_NO_WINDOW)
        started = time.monotonic()
        try:
            while proc.poll() is None:
                if cancel_event is not None and cancel_event.is_set():
                    self._kill_tree(proc)
                    raise OfficeCancelled("Conversion cancelled.")
                if time.monotonic() - started > timeout:
                    self._kill_tree(proc)
                    raise OfficeConversionError("Conversion timed out.")
                time.sleep(0.1)
            out, err = proc.communicate(timeout=10)
            return int(proc.returncode), (err or "") + (out or "")
        finally:
            if proc.poll() is None:
                self._kill_tree(proc)

    @staticmethod
    def _kill_tree(proc: subprocess.Popen) -> None:
        if proc.poll() is not None:
            return
        try:
            subprocess.run(["taskkill", "/PID", str(proc.pid), "/T", "/F"],
                           stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL,
                           stderr=subprocess.DEVNULL, creationflags=CREATE_NO_WINDOW,
                           timeout=8)
        except Exception:
            try:
                proc.kill()
            except Exception:
                pass

    @staticmethod
    def _tail(text: str, limit: int = 700) -> str:
        return " ".join((text or "").split())[-limit:]

    # ------------------------------------------------------------------ #
    #  Public discovery / validation
    # ------------------------------------------------------------------ #
    @staticmethod
    def normalize_target(target: str) -> str:
        value = str(target or "").strip().lower().lstrip(".")
        if value not in SUPPORTED_TARGETS:
            raise OfficeConversionError(
                f"Unsupported target format: {target!r}. "
                f"Supported: {', '.join(sorted(SUPPORTED_TARGETS))}.")
        return value

    @staticmethod
    def targets_for(input_path: str | Path) -> list[str]:
        ext = Path(input_path).suffix.lower()
        if ext in PPT_EXTS:
            return ["pdf", "png", "jpg", "txt"]
        if ext in EXCEL_EXTS:
            return ["pdf", "csv", "html", "json", "txt"] + (
                ["xlsx"] if ext == ".xls" else [])
        if ext in TEXT_EXTS:
            return ["pdf", "docx", "pptx", "xlsx", "txt"]
        if ext in WORD_EXTS:
            return ["pdf", "txt", "pptx"]
        if ext in PDF_EXTS:
            return ["pptx", "xlsx", "txt", "png", "jpg"]
        if ext in IMAGE_EXTS:
            return ["pptx", "pdf", "txt"]
        if ext in DATA_EXTS:
            return ["xlsx", "txt"]
        return []

    def validate(self, input_path: str | Path, target: str,
                 output_path: str | Path | None = None,
                 overwrite: bool = False) -> tuple[Path, str, Path]:
        src = Path(input_path)
        if not src.exists():
            raise OfficeConversionError(f"Input file does not exist: {src}")
        if not src.is_file():
            raise OfficeConversionError(f"Input is not a file: {src}")
        if src.suffix.lower() not in ALL_INPUT_EXTS:
            raise OfficeConversionError(
                f"Unsupported input format: {src.suffix or '(no extension)'}.")
        target = self.normalize_target(target)
        supported = self.targets_for(src)
        if target not in supported:
            raise OfficeConversionError(
                f"Cannot convert {src.suffix.upper()} to {target.upper()}. "
                f"Supported targets: {', '.join(supported)}.")
        if target in DIR_TARGETS:
            out = Path(output_path) if output_path else src.with_name(
                f"{src.stem}_{'slides' if target in ('png', 'jpg') else 'csv'}")
            if out.exists() and not out.is_dir():
                raise OfficeConversionError(f"Output directory path is a file: {out}")
            out.mkdir(parents=True, exist_ok=True)
            if any(out.iterdir()):
                raise OfficeConversionError(
                    f"Output directory is not empty: {out}")
        else:
            out = Path(output_path) if output_path else src.with_name(
                f"{src.stem}_converted.{target}")
            if out.suffix.lower() != f".{target}":
                raise OfficeConversionError(
                    f"Output extension {out.suffix!r} does not match target "
                    f"{target!r} (expected .{target}).")
            if out.exists() and not overwrite:
                raise OfficeConversionError(f"Output file already exists: {out}")
            out.parent.mkdir(parents=True, exist_ok=True)
        return src, target, out

    # ------------------------------------------------------------------ #
    #  Main API
    # ------------------------------------------------------------------ #
    def convert(self, input_path: str | Path, target: str,
                output_path: str | Path | None = None, *,
                overwrite: bool = False,
                callbacks: OfficeCallbacks | None = None,
                cancel_event: threading.Event | None = None) -> OfficeResult:
        callbacks = callbacks or OfficeCallbacks()
        started = time.monotonic()
        outputs: list[Path] = []
        out: Path | None = None
        before: set[Path] = set()
        try:
            src, target, out = self.validate(input_path, target, output_path,
                                             overwrite=overwrite)
            if overwrite and out.is_file():
                out.unlink(missing_ok=True)
            if target in DIR_TARGETS:
                before = set(out.iterdir())
            self._callback(callbacks.on_start, src, target)
            self._check(cancel_event)
            method = self._dispatch(src, target, out, callbacks, cancel_event)
            outputs = ([out] if out.is_file() else
                       sorted(p for p in out.iterdir() if p.is_file())) \
                if target in DIR_TARGETS else [out]
            if not outputs or any(not p.exists() or p.stat().st_size == 0
                                  for p in outputs):
                raise OfficeConversionError(
                    "Conversion finished without creating a valid output.")
            result = OfficeResult(src, outputs, target, method,
                                  time.monotonic() - started)
            self._callback(callbacks.on_success, result)
            return result
        except Exception as exc:  # noqa: BLE001
            if out is not None:
                if target in DIR_TARGETS:
                    for path in out.iterdir():
                        if path not in before:
                            try:
                                path.unlink(missing_ok=True)
                            except Exception:
                                pass
                else:
                    try:
                        out.unlink(missing_ok=True)
                    except Exception:
                        pass
            self._callback(callbacks.on_error, exc)
            if isinstance(exc, OfficeConversionError):
                raise
            raise OfficeConversionError(str(exc)) from exc
        finally:
            self.cleanup()

    def convert_async(self, input_path: str | Path, target: str,
                      output_path: str | Path | None = None, *,
                      overwrite: bool = False,
                      callbacks: OfficeCallbacks | None = None) -> OfficeHandle:
        return OfficeHandle(self, input_path, target, output_path, overwrite,
                            callbacks or OfficeCallbacks()).start()

    def to_pptx_many(self, inputs: Iterable[str | Path], out: str | Path, *,
                     callbacks: OfficeCallbacks | None = None,
                     cancel_event: threading.Event | None = None) -> OfficeResult:
        """Combine images and/or PDF pages into one PowerPoint presentation."""
        callbacks = callbacks or OfficeCallbacks()
        paths = [Path(p) for p in inputs]
        if not paths:
            raise OfficeConversionError("No images or PDFs were supplied.")
        out_path = Path(out)
        if out_path.suffix.lower() != ".pptx":
            raise OfficeConversionError("Combined presentation output must end in .pptx.")
        if out_path.exists():
            raise OfficeConversionError(f"Output file already exists: {out_path}")
        out_path.parent.mkdir(parents=True, exist_ok=True)
        started = time.monotonic()
        try:
            self._callback(callbacks.on_start, paths[0], "pptx")
            images: list[Path] = []
            for i, path in enumerate(paths, 1):
                self._check(cancel_event)
                if not path.is_file():
                    raise OfficeConversionError(f"Input file does not exist: {path}")
                ext = path.suffix.lower()
                if ext in IMAGE_EXTS:
                    images.append(path)
                elif ext in PDF_EXTS:
                    temp = self._tempdir(f"df-pdf-many-{i}-")
                    self._pdf_to_images(path, temp, "png", callbacks, cancel_event,
                                        f"{path.stem}_page", f"Rendering {path.name}")
                    images.extend(sorted(temp.glob("*.png")))
                else:
                    raise OfficeConversionError(
                        f"Cannot combine {path.suffix} into PowerPoint.")
            if not images:
                raise OfficeConversionError("No pages or images were produced.")
            self._images_to_pptx(images, out_path, callbacks, cancel_event)
            result = OfficeResult(paths[0], [out_path], "pptx", "python-pptx",
                                  time.monotonic() - started)
            self._callback(callbacks.on_success, result)
            return result
        except Exception as exc:  # noqa: BLE001
            try:
                out_path.unlink(missing_ok=True)
            except Exception:
                pass
            self._callback(callbacks.on_error, exc)
            if isinstance(exc, OfficeConversionError):
                raise
            raise OfficeConversionError(str(exc)) from exc
        finally:
            self.cleanup()

    def _dispatch(self, src: Path, target: str, out: Path,
                  callbacks: OfficeCallbacks,
                  cancel_event: threading.Event | None) -> str:
        ext = src.suffix.lower()
        if target == "pdf":
            return self._to_pdf(src, out, callbacks, cancel_event)
        if target in ("png", "jpg"):
            self._to_images(src, out, target, callbacks, cancel_event)
            return "libreoffice+pdf" if ext in PPT_EXTS else "pdf"
        if target == "csv":
            self._to_csv(src, out, callbacks, cancel_event)
            return "openpyxl"
        if target == "html":
            self._to_html(src, out, callbacks, cancel_event)
            return "openpyxl"
        if target == "json":
            self._to_json(src, out, callbacks, cancel_event)
            return "openpyxl"
        if target == "txt":
            self._to_txt(src, out, callbacks, cancel_event)
            return "python"
        if target == "docx":
            self._to_docx(src, out, callbacks, cancel_event)
            return "python-docx"
        if target == "pptx":
            self._to_pptx(src, out, callbacks, cancel_event)
            return "python-pptx"
        if target == "xlsx":
            self._to_xlsx(src, out, callbacks, cancel_event)
            return "openpyxl"
        raise OfficeConversionError(f"Unsupported target: {target}")

    # ------------------------------------------------------------------ #
    #  PDF
    # ------------------------------------------------------------------ #
    def _to_pdf(self, src: Path, out: Path, cb: OfficeCallbacks,
                cancel: threading.Event | None) -> str:
        ext = src.suffix.lower()
        if ext in TEXT_EXTS:
            self._text_to_pdf(src, out, cb, cancel)
            return "reportlab"
        if ext in PPT_EXTS | EXCEL_EXTS | WORD_EXTS:
            return self._office_to_pdf(src, out, cb, cancel)
        if ext in PDF_EXTS:
            shutil.copyfile(src, out)
            return "copy"
        if ext in IMAGE_EXTS:
            self._image_to_pdf(src, out, cb, cancel)
            return "reportlab"
        if ext in DATA_EXTS:
            temp = self._tempdir("df-office-") / f"{src.stem}.xlsx"
            try:
                self._to_xlsx(src, temp, cb, cancel)
                return self._office_to_pdf(temp, out, cb, cancel)
            finally:
                shutil.rmtree(temp.parent, ignore_errors=True)
        raise OfficeConversionError(f"No PDF converter for {src.suffix}.")

    def _office_to_pdf(self, src: Path, out: Path, cb: OfficeCallbacks,
                       cancel: threading.Event | None) -> str:
        errors: list[str] = []
        if self.prefer_com:
            try:
                self._com_to_pdf(src, out, cb, cancel)
                return "com"
            except Exception as exc:
                errors.append(f"COM: {self._tail(str(exc), 300)}")
                LOG.warning("Office COM PDF failed for %s: %s", src.name, exc)
        if self.soffice:
            try:
                self._soffice_to_pdf(src, out, cb, cancel)
                return "libreoffice"
            except Exception as exc:
                errors.append(f"LibreOffice: {self._tail(str(exc), 300)}")
                LOG.warning("LibreOffice PDF failed for %s: %s", src.name, exc)
        if not self.prefer_com:
            try:
                self._com_to_pdf(src, out, cb, cancel)
                return "com"
            except Exception as exc:
                errors.append(f"COM: {self._tail(str(exc), 300)}")
                LOG.warning("Office COM PDF failed for %s: %s", src.name, exc)
        ext = src.suffix.lower()
        if ext in EXCEL_EXTS:
            self._xlsx_to_pdf_basic(src, out, cb, cancel)
            return "openpyxl+reportlab"
        if ext in WORD_EXTS:
            self._docx_to_pdf_basic(src, out, cb, cancel)
            return "python-docx+reportlab"
        raise OfficeDependencyError(
            "PowerPoint -> PDF needs LibreOffice or Microsoft PowerPoint COM.\n\n"
            "Install LibreOffice.  Details: " + " | ".join(errors))

    def _soffice_to_pdf(self, src: Path, out: Path, cb: OfficeCallbacks,
                        cancel: threading.Event | None) -> None:
        soffice = self._require_soffice()
        self._check(cancel)
        self._progress(cb, 5, "Rendering PDF with LibreOffice")
        tmpdir = self._tempdir("df-lo-")
        profile = tmpdir / "profile"
        outdir = tmpdir / "out"
        outdir.mkdir()
        try:
            cmd = [
                soffice, "--headless", "--nologo", "--nodefault",
                "--nofirststartwizard",
                f"-env:UserInstallation={profile.resolve().as_uri()}",
                "--convert-to", "pdf", "--outdir", str(outdir), str(src),
            ]
            rc, detail = self._run_process(cmd, timeout=300, cancel_event=cancel)
            produced = outdir / f"{src.stem}.pdf"
            if rc != 0 or not produced.is_file() or produced.stat().st_size == 0:
                raise OfficeConversionError(
                    f"LibreOffice failed for '{src.name}'. {self._tail(detail)}")
            shutil.copyfile(produced, out)
            self._progress(cb, 100, "PDF ready")
        finally:
            shutil.rmtree(tmpdir, ignore_errors=True)

    def _com_to_pdf(self, src: Path, out: Path, cb: OfficeCallbacks,
                    cancel: threading.Event | None) -> None:
        if os.name != "nt":
            raise OfficeDependencyError("Microsoft Office COM is Windows-only.")
        try:
            import pythoncom  # type: ignore
            import win32com.client  # type: ignore
        except Exception as exc:
            raise OfficeDependencyError(
                "pywin32 is not installed. Install it with: pip install pywin32"
            ) from exc
        self._check(cancel)
        self._progress(cb, 10, "Rendering PDF with Microsoft Office")
        ext = src.suffix.lower()
        pythoncom.CoInitialize()
        app = None
        doc = None
        try:
            if ext in PPT_EXTS:
                app = win32com.client.DispatchEx("PowerPoint.Application")
                doc = app.Presentations.Open(str(src.resolve()), WithWindow=False,
                                             ReadOnly=True)
                doc.SaveAs(str(out.resolve()), 32)  # ppSaveAsPDF
            elif ext in EXCEL_EXTS:
                app = win32com.client.DispatchEx("Excel.Application")
                app.Visible = False
                app.DisplayAlerts = False
                doc = app.Workbooks.Open(str(src.resolve()), ReadOnly=True,
                                         UpdateLinks=0)
                doc.ExportAsFixedFormat(0, str(out.resolve()))
            elif ext in WORD_EXTS:
                app = win32com.client.DispatchEx("Word.Application")
                app.Visible = False
                app.DisplayAlerts = 0
                doc = app.Documents.Open(str(src.resolve()), ReadOnly=True,
                                         AddToRecentFiles=False)
                doc.ExportAsFixedFormat(str(out.resolve()), 17)
            else:
                raise OfficeDependencyError("No Office COM PDF path for this file.")
            if not out.exists() or out.stat().st_size == 0:
                raise OfficeConversionError("Microsoft Office produced no PDF.")
            self._progress(cb, 100, "PDF ready")
        finally:
            try:
                if doc is not None:
                    doc.Close(False)
            except Exception:
                pass
            try:
                if app is not None:
                    app.Quit()
            except Exception:
                pass
            try:
                pythoncom.CoUninitialize()
            except Exception:
                pass

    def _text_to_pdf(self, src: Path, out: Path, cb: OfficeCallbacks,
                     cancel: threading.Event | None) -> None:
        reportlab = self._require("reportlab", "reportlab")
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

        text = self._read_text(src)
        font_name = "Helvetica"
        font_file = Path(r"C:\Windows\Fonts\arial.ttf")
        if font_file.is_file():
            try:
                pdfmetrics.registerFont(TTFont("DropFlowSans", str(font_file)))
                font_name = "DropFlowSans"
            except Exception:
                pass
        style = ParagraphStyle("dropflow-text", fontName=font_name, fontSize=10,
                               leading=14, spaceAfter=2)
        self._progress(cb, 5, "Building text PDF")
        doc = SimpleDocTemplate(str(out), pagesize=A4,
                                leftMargin=18 * mm, rightMargin=18 * mm,
                                topMargin=16 * mm, bottomMargin=16 * mm)
        story = []
        lines = text.splitlines() or [""]
        for i, line in enumerate(lines, 1):
            self._check(cancel)
            value = self._shape_text(line)
            story.append(Paragraph(html_module.escape(value) or "&nbsp;", style))
            if i % 100 == 0:
                self._progress(cb, min(90, 5 + i / len(lines) * 85),
                               f"Writing text: {i}/{len(lines)}")
        doc.build(story or [Spacer(1, 1)])
        self._progress(cb, 100, "PDF ready")

    @staticmethod
    def _shape_text(value: str) -> str:
        """Best-effort Arabic shaping for reportlab output."""
        if not re.search(r"[\u0590-\u08FF]", value):
            return value
        try:
            import arabic_reshaper  # type: ignore
            from bidi.algorithm import get_display  # type: ignore
            return get_display(arabic_reshaper.reshape(value))
        except Exception:
            return value

    def _xlsx_to_pdf_basic(self, src: Path, out: Path, cb: OfficeCallbacks,
                           cancel: threading.Event | None) -> None:
        self._require("reportlab", "reportlab")
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import (LongTable, PageBreak, Paragraph,
                                         SimpleDocTemplate, Spacer, TableStyle)
        workbook = self._load_workbook(src)
        styles = getSampleStyleSheet()
        doc = SimpleDocTemplate(str(out), pagesize=landscape(A4),
                                leftMargin=12 * mm, rightMargin=12 * mm,
                                topMargin=12 * mm, bottomMargin=12 * mm)
        story = []
        try:
            for ws in workbook.worksheets:
                self._check(cancel)
                story.append(Paragraph(html_module.escape(ws.title), styles["Heading2"]))
                rows = [[self._cell_text(v) for v in row]
                        for row in ws.iter_rows(values_only=True)]
                if rows:
                    width = max(len(row) for row in rows)
                    if width:
                        rows = [row + [""] * (width - len(row)) for row in rows]
                        table = LongTable(rows, repeatRows=1, hAlign="LEFT")
                        table.setStyle(TableStyle([
                            ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
                            ("FONTSIZE", (0, 0), (-1, -1), 6),
                            ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                            ("VALIGN", (0, 0), (-1, -1), "TOP"),
                            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8e8f0")),
                        ]))
                        story.append(table)
                story.append(PageBreak())
        finally:
            workbook.close()
        doc.build(story or [Spacer(1, 1)])
        self._progress(cb, 100, "PDF ready")

    def _docx_to_pdf_basic(self, src: Path, out: Path, cb: OfficeCallbacks,
                           cancel: threading.Event | None) -> None:
        from docx import Document
        text_parts = []
        doc = Document(str(src))
        for para in doc.paragraphs:
            text_parts.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                text_parts.append(" | ".join(cell.text for cell in row.cells))
        temp = self._tempdir("df-docx-pdf-") / "text.txt"
        try:
            temp.write_text("\n".join(text_parts), encoding="utf-8")
            self._text_to_pdf(temp, out, cb, cancel)
        finally:
            shutil.rmtree(temp.parent, ignore_errors=True)

    def _image_to_pdf(self, src: Path, out: Path, cb: OfficeCallbacks,
                      cancel: threading.Event | None) -> None:
        self._require("reportlab", "reportlab")
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.units import mm
        from reportlab.pdfgen import canvas
        from PIL import Image
        self._check(cancel)
        with Image.open(src) as im:
            width, height = im.size
        page_w, page_h = A4
        scale = min((page_w - 30 * mm) / width, (page_h - 30 * mm) / height)
        draw_w, draw_h = width * scale, height * scale
        x = (page_w - draw_w) / 2
        y = (page_h - draw_h) / 2
        c = canvas.Canvas(str(out), pagesize=A4)
        c.drawImage(str(src), x, y, width=draw_w, height=draw_h,
                    preserveAspectRatio=True, mask="auto")
        c.showPage()
        c.save()
        self._progress(cb, 100, "PDF ready")

    # ------------------------------------------------------------------ #
    #  Images / PowerPoint
    # ------------------------------------------------------------------ #
    def _to_images(self, src: Path, out_dir: Path, fmt: str,
                   cb: OfficeCallbacks,
                   cancel: threading.Event | None) -> None:
        ext = src.suffix.lower()
        tmp_root = self._tempdir("df-images-")
        try:
            if ext in PPT_EXTS:
                pdf = tmp_root / f"{src.stem}.pdf"
                self._to_pdf(src, pdf, cb, cancel)
                prefix = f"{src.stem}_slide"
                self._pdf_to_images(pdf, out_dir, fmt, cb, cancel, prefix,
                                    "Rendering slides")
            elif ext in PDF_EXTS:
                self._pdf_to_images(src, out_dir, fmt, cb, cancel,
                                    f"{src.stem}_page", "Rendering pages")
            else:
                raise OfficeConversionError(
                    f"{src.suffix} -> images is not supported.")
        finally:
            shutil.rmtree(tmp_root, ignore_errors=True)

    def _pdf_to_images(self, pdf: Path, out_dir: Path, fmt: str,
                       cb: OfficeCallbacks, cancel: threading.Event | None,
                       prefix: str, label: str) -> None:
        fitz = self._require("fitz", "PyMuPDF")
        from PIL import Image
        doc = fitz.open(str(pdf))
        try:
            total = doc.page_count
            for i, page in enumerate(doc, 1):
                self._check(cancel)
                pix = page.get_pixmap(dpi=self.dpi, alpha=False)
                target = out_dir / f"{prefix}_{i:03d}.{fmt}"
                if fmt == "png":
                    pix.save(str(target))
                else:
                    image = Image.frombytes("RGB", (pix.width, pix.height),
                                            pix.samples)
                    image.save(str(target), "JPEG", quality=92)
                self._progress(cb, i / max(1, total) * 100,
                               f"{label}: {i}/{total}")
        finally:
            doc.close()

    def _to_pptx(self, src: Path, out: Path, cb: OfficeCallbacks,
                 cancel: threading.Event | None) -> None:
        ext = src.suffix.lower()
        if ext in IMAGE_EXTS:
            self._images_to_pptx([src], out, cb, cancel)
            return
        if ext in PDF_EXTS:
            temp = self._tempdir("df-pdf-pptx-")
            try:
                self._pdf_to_images(src, temp, "png", cb, cancel,
                                    f"{src.stem}_page", "Rendering PDF pages")
                images = sorted(temp.glob("*.png"))
                self._images_to_pptx(images, out, cb, cancel)
            finally:
                shutil.rmtree(temp, ignore_errors=True)
            return
        if ext in WORD_EXTS:
            self._text_to_pptx(self._docx_text(src), out, cb, cancel)
            return
        if ext in TEXT_EXTS:
            self._text_to_pptx(self._read_text(src), out, cb, cancel)
            return
        raise OfficeConversionError(f"{src.suffix} -> PowerPoint is not supported.")

    def _images_to_pptx(self, images: Iterable[Path], out: Path,
                        cb: OfficeCallbacks,
                        cancel: threading.Event | None) -> None:
        pptx = self._require("pptx", "python-pptx")
        from PIL import Image
        from pptx import Presentation
        from pptx.util import Inches, Pt
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        items = list(images)
        if not items:
            raise OfficeConversionError("No images to put into PowerPoint.")
        for i, image_path in enumerate(items, 1):
            self._check(cancel)
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            with Image.open(image_path) as im:
                iw, ih = im.size
            max_w, max_h = 12.4, 6.4
            scale = min(max_w / iw, max_h / ih)
            w, h = iw * scale, ih * scale
            left = (13.333 - w) / 2
            top = (7.5 - h) / 2 - 0.15
            slide.shapes.add_picture(str(image_path), Inches(left), Inches(top),
                                     width=Inches(w), height=Inches(h))
            caption = slide.shapes.add_textbox(Inches(0.4), Inches(7.05),
                                               Inches(12.5), Inches(0.3))
            caption.text_frame.text = image_path.stem
            caption.text_frame.paragraphs[0].font.size = Pt(9)
            self._progress(cb, i / len(items) * 100,
                           f"Adding slide {i}/{len(items)}")
        prs.save(str(out))

    def _text_to_pptx(self, text: str, out: Path, cb: OfficeCallbacks,
                      cancel: threading.Event | None) -> None:
        self._require("pptx", "python-pptx")
        from pptx import Presentation
        from pptx.util import Inches, Pt
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        title_slide = prs.slides.add_slide(prs.slide_layouts[0])
        title_slide.shapes.title.text = out.stem.replace("_", " ").title()
        title_slide.placeholders[1].text = "Converted from text"
        lines = text.splitlines()
        chunks: list[list[str]] = []
        current: list[str] = []
        chars = 0
        for line in lines:
            if len(current) >= 12 or chars + len(line) > 1300:
                chunks.append(current)
                current, chars = [], 0
            current.append(line)
            chars += len(line) + 1
        if current or not chunks:
            chunks.append(current)
        for i, chunk in enumerate(chunks, 1):
            self._check(cancel)
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = f"Part {i}"
            body = slide.placeholders[1].text_frame
            body.text = "\n".join(chunk[:14]) or " "
            for para in body.paragraphs:
                para.font.size = Pt(18)
            self._progress(cb, i / len(chunks) * 100,
                           f"Creating slide {i}/{len(chunks)}")
        prs.save(str(out))

    # ------------------------------------------------------------------ #
    #  Excel / data
    # ------------------------------------------------------------------ #
    def _to_csv(self, src: Path, out_dir: Path, cb: OfficeCallbacks,
                cancel: threading.Event | None) -> None:
        workbook = self._load_workbook(src)
        total = len(workbook.worksheets)
        used: set[str] = set()
        try:
            for i, ws in enumerate(workbook.worksheets, 1):
                self._check(cancel)
                slug = self._sheet_slug(ws.title)
                target = out_dir / f"{src.stem}_{slug}.csv"
                n = 2
                while target.name.lower() in used:
                    target = out_dir / f"{src.stem}_{slug}_{n}.csv"
                    n += 1
                used.add(target.name.lower())
                with open(target, "w", encoding="utf-8-sig", newline="") as fh:
                    writer = csv.writer(fh)
                    for row in ws.iter_rows(values_only=True):
                        writer.writerow([self._cell_text(v) for v in row])
                self._progress(cb, i / total * 100, f"Exporting sheet: {ws.title}")
        finally:
            workbook.close()

    def _to_html(self, src: Path, out: Path, cb: OfficeCallbacks,
                 cancel: threading.Event | None) -> None:
        workbook = self._load_workbook(src)
        parts = [
            "<!doctype html><html><head><meta charset='utf-8'>",
            "<title>" + html_module.escape(src.stem) + "</title>",
            "<style>body{font-family:Segoe UI,Arial,sans-serif;margin:24px;"
            "color:#222}h2{margin-top:28px}table{border-collapse:collapse;"
            "margin:8px 0 18px}td,th{border:1px solid #ccc;padding:5px 8px;"
            "vertical-align:top}th{background:#eee;text-align:left}</style>",
            "</head><body>",
        ]
        try:
            sheets = list(workbook.worksheets)
            for i, ws in enumerate(sheets, 1):
                self._check(cancel)
                parts.append(f"<h2>{html_module.escape(ws.title)}</h2><table>")
                for r, row in enumerate(ws.iter_rows(values_only=True)):
                    tag = "th" if r == 0 else "td"
                    parts.append("<tr>" + "".join(
                        f"<{tag}>{html_module.escape(self._cell_text(v))}</{tag}>"
                        for v in row) + "</tr>")
                parts.append("</table>")
                self._progress(cb, i / len(sheets) * 100, f"Writing sheet: {ws.title}")
        finally:
            workbook.close()
        parts.append("</body></html>")
        out.write_text("\n".join(parts), encoding="utf-8")

    def _to_json(self, src: Path, out: Path, cb: OfficeCallbacks,
                 cancel: threading.Event | None) -> None:
        workbook = self._load_workbook(src)
        payload: dict[str, Any] = {}
        try:
            sheets = list(workbook.worksheets)
            for i, ws in enumerate(sheets, 1):
                self._check(cancel)
                rows = [[self._excel_value(v) for v in row]
                        for row in ws.iter_rows(values_only=True)]
                headers = rows[0] if rows else []
                body = rows[1:] if rows else []
                if headers and all(str(h).strip() for h in headers):
                    records = []
                    for row in body:
                        records.append({str(headers[j]): (row[j] if j < len(row) else "")
                                        for j in range(len(headers))})
                    payload[ws.title] = records
                else:
                    payload[ws.title] = rows
                self._progress(cb, i / len(sheets) * 100, f"Reading sheet: {ws.title}")
        finally:
            workbook.close()
        out.write_text(json.dumps(payload, ensure_ascii=False, indent=2,
                                  default=str), encoding="utf-8")


    def _to_txt(self, src: Path, out: Path, cb: OfficeCallbacks,
                cancel: threading.Event | None) -> None:
        text = self._extract_text(src, cb, cancel)
        out.write_text(text, encoding="utf-8")
        self._progress(cb, 100, "Text ready")

    def _to_docx(self, src: Path, out: Path, cb: OfficeCallbacks,
                 cancel: threading.Event | None) -> None:
        self._require("docx", "python-docx")
        from docx import Document
        text = self._extract_text(src, cb, cancel)
        doc = Document()
        for line in text.splitlines():
            self._check(cancel)
            doc.add_paragraph(line)
        doc.save(str(out))
        self._progress(cb, 100, "Word document ready")

    def _to_xlsx(self, src: Path, out: Path, cb: OfficeCallbacks,
                 cancel: threading.Event | None) -> None:
        ext = src.suffix.lower()
        if ext in EXCEL_EXTS:
            shutil.copyfile(self._load_xlsx_path(src), out)
            return
        if ext in TEXT_EXTS:
            self._text_to_xlsx(src, out)
            return
        if ext == ".csv":
            self._csv_to_xlsx(src, out)
            return
        if ext == ".json":
            self._json_to_xlsx(src, out)
            return
        if ext in PDF_EXTS:
            self._pdf_to_xlsx(src, out, cb, cancel)
            return
        raise OfficeConversionError(f"{src.suffix} -> Excel is not supported.")

    def _text_to_xlsx(self, src: Path, out: Path) -> None:
        lines = self._read_text(src).splitlines()
        if any("\t" in line for line in lines):
            rows = [line.split("\t") for line in lines]
        elif any("," in line for line in lines):
            rows = list(csv.reader(lines))
        else:
            rows = [[line] for line in lines]
        self._write_workbook(out, [("Text", rows)])

    def _csv_to_xlsx(self, src: Path, out: Path) -> None:
        text = self._read_text(src)
        rows = list(csv.reader(text.splitlines()))
        self._write_workbook(out, [(src.stem, rows)])

    def _json_to_xlsx(self, src: Path, out: Path) -> None:
        data = json.loads(self._read_text(src))
        self._write_workbook(out, self._json_sheets(data))

    def _pdf_to_xlsx(self, src: Path, out: Path, cb: OfficeCallbacks,
                     cancel: threading.Event | None) -> None:
        fitz = self._require("fitz", "PyMuPDF")
        doc = fitz.open(str(src))
        sheets: list[tuple[str, list[list[Any]]]] = []
        try:
            for page_no, page in enumerate(doc, 1):
                self._check(cancel)
                found = False
                finder = getattr(page, "find_tables", None)
                if callable(finder):
                    try:
                        tables = finder().tables
                    except Exception:
                        tables = []
                    for t_no, table in enumerate(tables, 1):
                        rows = table.extract()
                        if rows:
                            sheets.append((f"Page {page_no} Table {t_no}", rows))
                            found = True
                if not found:
                    rows = [[line] for line in page.get_text("text").splitlines()]
                    if rows:
                        sheets.append((f"Page {page_no}", rows))
                self._progress(cb, page_no / max(1, doc.page_count) * 100,
                               f"Reading PDF page {page_no}/{doc.page_count}")
        finally:
            doc.close()
        if not sheets:
            raise OfficeConversionError("No table or text data found in the PDF.")
        self._write_workbook(out, sheets)

    def _write_workbook(self, out: Path,
                        sheets: list[tuple[str, list[list[Any]]]]) -> None:
        self._require("openpyxl", "openpyxl")
        from openpyxl import Workbook
        from openpyxl.styles import Font
        wb = Workbook()
        wb.remove(wb.active)
        used: set[str] = set()
        for name, rows in sheets:
            title = self._unique_sheet(name, used)
            ws = wb.create_sheet(title)
            for row in rows:
                ws.append([self._excel_value(v) for v in row])
            if rows:
                for cell in ws[1]:
                    cell.font = Font(bold=True)
                ws.freeze_panes = "A2"
            for column in ws.columns:
                width = min(60, max(10, max(
                    (len(str(c.value)) for c in column if c.value is not None),
                    default=10) + 2))
                ws.column_dimensions[column[0].column_letter].width = width
        if not wb.worksheets:
            wb.create_sheet("Sheet1")
        wb.save(str(out))

    # ------------------------------------------------------------------ #
    #  Text extraction
    # ------------------------------------------------------------------ #
    def _extract_text(self, src: Path, cb: OfficeCallbacks,
                      cancel: threading.Event | None) -> str:
        ext = src.suffix.lower()
        self._progress(cb, 5, f"Reading {src.name}")
        if ext in TEXT_EXTS:
            return self._read_text(src)
        if ext in PPT_EXTS:
            return self._pptx_text(src, cb, cancel)
        if ext in EXCEL_EXTS:
            return self._xlsx_text(src, cb, cancel)
        if ext in WORD_EXTS:
            return self._docx_text(src)
        if ext in PDF_EXTS:
            return self._pdf_text(src, cb, cancel)
        if ext == ".csv":
            return self._read_text(src)
        if ext == ".json":
            return json.dumps(json.loads(self._read_text(src)),
                              ensure_ascii=False, indent=2, default=str)
        if ext in IMAGE_EXTS:
            return self._image_text(src)
        raise OfficeConversionError(f"No text extractor for {src.suffix}.")

    def _pptx_text(self, src: Path, cb: OfficeCallbacks,
                   cancel: threading.Event | None) -> str:
        self._require("pptx", "python-pptx")
        from pptx import Presentation
        path = self._modernize(src, "pptx")
        try:
            prs = Presentation(str(path))
            lines: list[str] = []
            total = len(prs.slides)
            for i, slide in enumerate(prs.slides, 1):
                self._check(cancel)
                lines.append(f"Slide {i}")
                for shape in slide.shapes:
                    if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                        lines.append(shape.text_frame.text)
                    if getattr(shape, "has_table", False):
                        for row in shape.table.rows:
                            lines.append(" | ".join(cell.text for cell in row.cells))
                lines.append("")
                self._progress(cb, i / max(1, total) * 90 + 5,
                               f"Reading slide {i}/{total}")
            return "\n".join(lines).strip()
        finally:
            pass

    def _xlsx_text(self, src: Path, cb: OfficeCallbacks,
                   cancel: threading.Event | None) -> str:
        workbook = self._load_workbook(src)
        lines: list[str] = []
        try:
            sheets = list(workbook.worksheets)
            for i, ws in enumerate(sheets, 1):
                self._check(cancel)
                lines.append(f"Sheet: {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    if any(v is not None for v in row):
                        lines.append("\t".join(self._cell_text(v) for v in row))
                lines.append("")
                self._progress(cb, i / max(1, len(sheets)) * 90 + 5,
                               f"Reading sheet {i}/{len(sheets)}")
        finally:
            workbook.close()
        return "\n".join(lines).strip()

    def _docx_text(self, src: Path) -> str:
        self._require("docx", "python-docx")
        from docx import Document
        doc = Document(str(src))
        lines = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(lines)

    def _pdf_text(self, src: Path, cb: OfficeCallbacks,
                  cancel: threading.Event | None) -> str:
        fitz = self._require("fitz", "PyMuPDF")
        doc = fitz.open(str(src))
        try:
            parts = []
            total = doc.page_count
            for i, page in enumerate(doc, 1):
                self._check(cancel)
                parts.append(f"Page {i}")
                parts.append(page.get_text("text"))
                parts.append("")
                self._progress(cb, i / max(1, total) * 90 + 5,
                               f"Reading page {i}/{total}")
            return "\n".join(parts).strip()
        finally:
            doc.close()

    @staticmethod
    def _find_tesseract() -> str | None:
        candidates = []
        if os.environ.get("TESSERACT_PATH"):
            candidates.append(os.environ["TESSERACT_PATH"])
        found = shutil.which("tesseract")
        if found:
            candidates.append(found)
        candidates += [
            r"C:\Program Files\Tesseract-OCR\tesseract.exe",
            r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe",
        ]
        for candidate in candidates:
            try:
                if Path(candidate).is_file():
                    return str(candidate)
            except Exception:
                pass
        return None

    def _image_text(self, src: Path) -> str:
        try:
            import pytesseract  # type: ignore
            from PIL import Image
        except Exception as exc:
            raise OfficeDependencyError(
                "pytesseract and Pillow are required for image text extraction."
            ) from exc
        tesseract = self._find_tesseract()
        if not tesseract:
            raise OfficeDependencyError(
                "Tesseract was not found.\n\n"
                "Install it from https://github.com/UB-Mannheim/tesseract/wiki "
                "or set TESSERACT_PATH.")
        pytesseract.pytesseract.tesseract_cmd = tesseract
        try:
            with Image.open(src) as im:
                return pytesseract.image_to_string(im, lang=self.ocr_lang)
        except Exception as exc:
            raise OfficeConversionError(f"OCR failed for {src.name}: {exc}") from exc


    # ------------------------------------------------------------------ #
    #  Workbook / legacy helpers
    # ------------------------------------------------------------------ #
    def _load_workbook(self, src: Path):
        self._require("openpyxl", "openpyxl")
        from openpyxl import load_workbook
        path = self._load_xlsx_path(src)
        wb = load_workbook(str(path), data_only=True, read_only=True)
        return wb

    def _load_xlsx_path(self, src: Path) -> Path:
        if src.suffix.lower() == ".xlsx":
            return src
        return self._modernize(src, "xlsx")

    def _modernize(self, src: Path, target_ext: str) -> Path:
        """Convert .ppt/.xls to .pptx/.xlsx in a temporary directory."""
        if src.suffix.lower() == f".{target_ext}":
            return src
        soffice = self._require_soffice()
        tmpdir = self._tempdir("df-modern-")
        profile = tmpdir / "profile"
        outdir = tmpdir / "out"
        outdir.mkdir()
        try:
            cmd = [
                soffice, "--headless", "--nologo", "--nodefault",
                "--nofirststartwizard",
                f"-env:UserInstallation={profile.resolve().as_uri()}",
                "--convert-to", target_ext, "--outdir", str(outdir), str(src),
            ]
            rc, detail = self._run_process(cmd, timeout=180)
            produced = outdir / f"{src.stem}.{target_ext}"
            if rc != 0 or not produced.is_file():
                raise OfficeConversionError(
                    f"LibreOffice could not read '{src.name}'. {self._tail(detail)}")
            return produced
        except Exception:
            shutil.rmtree(tmpdir, ignore_errors=True)
            raise

    @staticmethod
    def _cell_text(value: Any) -> str:
        if value is None:
            return ""
        return str(value)

    @staticmethod
    def _excel_value(value: Any) -> Any:
        if isinstance(value, (str, int, float, bool)) or value is None:
            return value
        return str(value)

    @staticmethod
    def _sheet_slug(name: str) -> str:
        value = re.sub(r"[^A-Za-z0-9._-]+", "_", name).strip("_")
        return value or "sheet"

    @staticmethod
    def _unique_sheet(name: str, used: set[str]) -> str:
        clean = re.sub(r"[\[\]:*?/\\]", "_", str(name or "Sheet"))[:31] or "Sheet"
        candidate = clean
        n = 2
        while candidate.lower() in used:
            suffix = f"_{n}"
            candidate = clean[:31 - len(suffix)] + suffix
            n += 1
        used.add(candidate.lower())
        return candidate

    @staticmethod
    def _read_text(path: Path) -> str:
        for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
            try:
                return path.read_text(encoding=encoding)
            except UnicodeDecodeError:
                continue
        return path.read_text(encoding="utf-8", errors="replace")

    @staticmethod
    def _json_rows(rows: list[list[Any]]) -> list[list[Any]]:
        return [[value if isinstance(value, (str, int, float, bool)) or value is None
                 else json.dumps(value, ensure_ascii=False, default=str)
                 for value in row] for row in rows]

    def _json_sheets(self, data: Any) -> list[tuple[str, list[list[Any]]]]:
        if isinstance(data, list):
            if data and all(isinstance(x, dict) for x in data):
                keys: list[str] = []
                for item in data:
                    for key in item:
                        if key not in keys:
                            keys.append(str(key))
                rows: list[list[Any]] = [keys]
                for item in data:
                    rows.append([item.get(key, "") for key in keys])
                return [("Data", self._json_rows(rows))]
            if data and all(isinstance(x, (list, tuple)) for x in data):
                return [("Data", self._json_rows([list(x) for x in data]))]
            return [("Data", self._json_rows([[x] for x in data]))]
        if isinstance(data, dict):
            list_sheets = {str(k): v for k, v in data.items()
                           if isinstance(v, list) and v and
                           all(isinstance(x, (dict, list, tuple)) for x in v)}
            if list_sheets:
                return [(name, self._json_sheets(value)[0][1])
                        for name, value in list_sheets.items()]
            rows = [["Key", "Value"]]
            for key, value in data.items():
                if isinstance(value, (dict, list)):
                    value = json.dumps(value, ensure_ascii=False, default=str)
                rows.append([str(key), value])
            return [("Data", self._json_rows(rows))]
        return [("Data", [["Value"], [data]])]

    # ------------------------------------------------------------------ #
    #  Convenience methods requested by the integration
    # ------------------------------------------------------------------ #
    def pptx_to_pdf(self, src: str | Path, out: str | Path, **kw) -> OfficeResult:
        return self.convert(src, "pdf", out, **kw)

    def pptx_to_images(self, src: str | Path, out_dir: str | Path,
                       fmt: str = "png", **kw) -> OfficeResult:
        return self.convert(src, fmt, out_dir, **kw)

    def pptx_to_txt(self, src: str | Path, out: str | Path, **kw) -> OfficeResult:
        return self.convert(src, "txt", out, **kw)

    def excel_to_pdf(self, src: str | Path, out: str | Path, **kw) -> OfficeResult:
        return self.convert(src, "pdf", out, **kw)

    def excel_to_csv(self, src: str | Path, out_dir: str | Path, **kw) -> OfficeResult:
        return self.convert(src, "csv", out_dir, **kw)

    def excel_to_html(self, src: str | Path, out: str | Path, **kw) -> OfficeResult:
        return self.convert(src, "html", out, **kw)

    def excel_to_json(self, src: str | Path, out: str | Path, **kw) -> OfficeResult:
        return self.convert(src, "json", out, **kw)

    def excel_to_txt(self, src: str | Path, out: str | Path, **kw) -> OfficeResult:
        return self.convert(src, "txt", out, **kw)

    def txt_to_pdf(self, src: str | Path, out: str | Path, **kw) -> OfficeResult:
        return self.convert(src, "pdf", out, **kw)

    def txt_to_docx(self, src: str | Path, out: str | Path, **kw) -> OfficeResult:
        return self.convert(src, "docx", out, **kw)

    def txt_to_pptx(self, src: str | Path, out: str | Path, **kw) -> OfficeResult:
        return self.convert(src, "pptx", out, **kw)

    def txt_to_xlsx(self, src: str | Path, out: str | Path, **kw) -> OfficeResult:
        return self.convert(src, "xlsx", out, **kw)
